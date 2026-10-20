"""Synthetic language acceptance through real parsers, using the production engine."""

import io
import json
from collections import Counter
from email import policy
from email.message import EmailMessage
from html import escape
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from panda import formats
from panda.detectors import default_detector_engine
from panda.errors import ImageOnlyError
from panda.handlers import register_defaults
from panda.handlers.pdf import sanitise_pdf_file
from panda.leakcheck import scan_output
from panda.pseudonyms import FixedPseudonymiser

PASSAGES = json.loads(Path('tests/fixtures/t10/passages.json').read_text())


@pytest.fixture(scope='module')
def engine():
    return default_detector_engine()


@pytest.fixture(autouse=True)
def registry():
    handlers, extractors = dict(formats.HANDLERS), dict(formats.TEXT_EXTRACTORS)
    register_defaults()
    yield
    formats.HANDLERS.clear()
    formats.HANDLERS.update(handlers)
    formats.TEXT_EXTRACTORS.clear()
    formats.TEXT_EXTRACTORS.update(extractors)


def html_bytes(lines):
    # Splitting every character exercises joined inline detection, including names.
    return ('<html><body>' + ''.join(
        '<p>' + ''.join('<span>' + escape(c) + '</span>' for c in line) + '</p>'
        for line in lines
    ) + '</body></html>').encode()


def docx_bytes(lines):
    doc = Document()
    for line in lines:
        paragraph = doc.add_paragraph()
        for character in line:
            paragraph.add_run(character)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def eml_bytes(lines, nested=False):
    message = EmailMessage()
    message.set_content('\n'.join(lines))
    if nested:
        message.add_alternative(html_bytes(lines).decode(), subtype='html')
        message.add_attachment(html_bytes(lines).decode(), subtype='html',
                               filename='synthetic.html')
        message.add_attachment(docx_bytes(lines), maintype='application',
                               subtype='vnd.openxmlformats-officedocument.wordprocessingml.document',
                               filename='synthetic.docx')
        parent = EmailMessage()
        parent.set_content('Delivery confirmed.')
        parent.add_attachment(message)
        message = parent
    return message.as_bytes(policy=policy.default)


def pdf_bytes(lines):
    """Explicit synthetic ToUnicode map; proves extraction, not real font fidelity."""
    writer = PdfWriter()
    chars = sorted(set(''.join(lines)))
    codes = {char: i + 1 for i, char in enumerate(chars)}
    cmap = DecodedStreamObject()
    mappings = '\n'.join(f'<{code:04x}> <{ord(char):04x}>' for char, code in codes.items())
    cmap.set_data((
        '/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n'
        '/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n'
        '/CMapName /Synthetic def /CMapType 2 def\n'
        '1 begincodespacerange <0000> <ffff> endcodespacerange\n'
        f'{len(codes)} beginbfchar\n{mappings}\nendbfchar\n'
        'endcmap CMapName currentdict /CMap defineresource pop end end'
    ).encode())
    cid = DictionaryObject({
        NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/CIDFontType2'),
        NameObject('/BaseFont'): NameObject('/Synthetic'),
    })
    font = DictionaryObject({
        NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type0'),
        NameObject('/BaseFont'): NameObject('/Synthetic'),
        NameObject('/Encoding'): NameObject('/Identity-H'),
        NameObject('/DescendantFonts'): ArrayObject([writer._add_object(cid)]),
        NameObject('/ToUnicode'): writer._add_object(cmap),
    })
    image = DecodedStreamObject()
    image.set_data(b'\xff\x00\x00')
    image.update({
        NameObject('/Type'): NameObject('/XObject'),
        NameObject('/Subtype'): NameObject('/Image'),
        NameObject('/Width'): NumberObject(1), NameObject('/Height'): NumberObject(1),
        NameObject('/ColorSpace'): NameObject('/DeviceRGB'),
        NameObject('/BitsPerComponent'): NumberObject(8),
    })
    # A mixed page and a genuinely image-only page, not an extraction mock.
    for page_lines in (lines, []):
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)}),
            NameObject('/XObject'): DictionaryObject({
                NameObject('/Im1'): writer._add_object(image),
            }),
        })
        stream = DecodedStreamObject()
        commands = ['q 10 0 0 10 0 0 cm /Im1 Do Q', 'BT /F1 12 Tf 14 TL 60 760 Td']
        for line in page_lines:
            encoded = ''.join(f'{codes[c]:04x}' for c in line)
            commands.append(f'<{encoded}> Tj T*')
        commands.append('ET')
        stream.set_data('\n'.join(commands).encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


@pytest.mark.parametrize('language', PASSAGES)
@pytest.mark.parametrize('extension', ['.html', '.htm', '.docx', '.eml', '.pdf', '.nested.eml'])
def test_real_format_languages_and_nested_attachments(tmp_path, engine, language, extension):
    rows = PASSAGES[language]
    lines = [r['text'] for r in rows] + ['Delivery confirmed.']
    if extension in ('.html', '.htm'):
        raw = html_bytes(lines)
    elif extension == '.docx':
        raw = docx_bytes(lines)
    elif extension == '.pdf':
        raw = pdf_bytes(lines)
        reader = PdfReader(io.BytesIO(raw))
        extracted = reader.pages[0].extract_text()
        assert all(line in extracted for line in lines)
        assert reader.pages[1].extract_text() == ''
    else:
        raw = eml_bytes(lines, nested=extension == '.nested.eml')
    source = tmp_path / ('synthetic' + extension)
    source.write_bytes(raw)
    destination = tmp_path / 'out' / ('file-0001' + source.suffix)
    result = formats.HANDLERS[source.suffix](source, destination, FixedPseudonymiser(), engine)
    if extension == '.pdf':
        destination = destination.with_suffix('.pdf.txt')
        text = destination.read_text()
        assert text.split('\f')[1].strip() == ''
    else:
        text = '\n'.join(unit.text for unit in formats.TEXT_EXTRACTORS[source.suffix](destination))
    expected = Counter((e['kind'], e['value']) for row in rows for e in row['entities'])
    multiplicity = 4 if extension == '.nested.eml' else 1
    actual = Counter((r.detection.kind.value, r.detection.text) for r in result.replacements)
    assert actual == Counter({key: count * multiplicity for key, count in expected.items()})
    assert all(value not in text for _, value in expected)
    assert 'Delivery confirmed.' in text
    assert all(r.detection.locator for r in result.replacements)
    assert source.read_bytes() == raw
    report = scan_output(destination.parent, engine, generated_dummies={
        destination: frozenset(item.surrogate.casefold() for item in result.replacements),
    })
    assert report.clean and not report.unscannable


def test_real_image_only_pdf_is_refused(tmp_path, engine):
    source = tmp_path / 'synthetic.pdf'
    source.write_bytes(pdf_bytes([]))
    destination = tmp_path / 'out.pdf'
    with pytest.raises(ImageOnlyError):
        sanitise_pdf_file(source, destination, FixedPseudonymiser(), engine)
    assert not destination.with_suffix('.pdf.txt').exists()
