import io
import zipfile
from email import message_from_bytes, policy
from email.message import EmailMessage

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from panda import formats
from panda.detectors import default_detector_engine
from panda.handlers import register_defaults
from panda.handlers.docx import extract_docx_text
from panda.handlers.eml import sanitise_eml_file
from panda.leakcheck import scan_output
from panda.model import PIIKind
from panda.pseudonyms import FixedPseudonymiser


@pytest.fixture(autouse=True)
def clean_registry():
    """Register the built-in handlers for attachment recursion, then restore.

    The `.eml` handler recurses into attachments through the shared registry, so it
    must be populated; snapshotting and restoring keeps it from leaking into
    `test_formats`, which asserts the registry starts empty.
    """
    handlers = dict(formats.HANDLERS)
    extractors = dict(formats.TEXT_EXTRACTORS)
    register_defaults()
    yield
    formats.HANDLERS.clear()
    formats.HANDLERS.update(handlers)
    formats.TEXT_EXTRACTORS.clear()
    formats.TEXT_EXTRACTORS.update(extractors)


def _build(headers=None, body=None, html=None, attachments=(), nested=None) -> bytes:
    """Assemble a small .eml. attachments is (filename, maintype, subtype, bytes)."""
    message = EmailMessage()
    for name, value in (headers or {}).items():
        message[name] = value
    if body is not None:
        message.set_content(body)
    if html is not None:
        message.add_alternative(html, subtype="html")
    for filename, maintype, subtype, payload in attachments:
        message.add_attachment(payload, maintype=maintype, subtype=subtype, filename=filename)
    if nested is not None:
        message.add_attachment(nested)
    return message.as_bytes(policy=policy.default)


def _run(tmp_path, raw):
    source = tmp_path / "in.eml"
    destination = tmp_path / "out" / "in.eml"
    source.write_bytes(raw)
    result = sanitise_eml_file(
        source, destination, FixedPseudonymiser(), default_detector_engine()
    )
    return destination.read_bytes(), result


def test_replaces_pii_in_plain_body(tmp_path):
    raw = _build(
        headers={"Subject": "hi"},
        body="Call me on 0412 345 678 or email real.person@mail.org",
    )
    out, result = _run(tmp_path, raw)
    assert b"0412 345 678" not in out and b"real.person@mail.org" not in out
    assert b"0400000000" in out and b"john.doe@example.com" in out
    kinds = {r.detection.kind for r in result.replacements}
    assert PIIKind.PHONE in kinds and PIIKind.EMAIL in kinds


def test_html_part_removes_attributes(tmp_path):
    """HTML bodies share the closed HTML review dialect."""
    raw = _build(
        headers={"Subject": "hi"},
        body="plain",
        html="<html><body><a href='mailto:real.person@mail.org'>x</a></body></html>",
    )
    out, _ = _run(tmp_path, raw)
    assert b"real.person@mail.org" not in out
    assert b"mailto:" not in out


def test_headers_are_sanitised(tmp_path):
    raw = _build(
        headers={
            "From": '"Margaret Whitfield" <m.whitfield@client.com>',
            "To": "advisor@firm.com",
            "Subject": "call 0412 345 678",
            "Message-ID": "<trace123@mail.client.com>",
        },
        body="body",
    )
    out, _ = _run(tmp_path, raw)
    text = out.decode()
    assert "Margaret Whitfield" not in text and "m.whitfield@client.com" not in text
    assert "advisor@firm.com" not in text
    # The Message-ID's email-like token is replaced too, or the rescan would flag it.
    assert "trace123@mail.client.com" not in text
    assert "0412 345 678" not in text


def test_image_attachment_is_dropped(tmp_path):
    raw = _build(
        headers={"Subject": "hi"},
        body="body",
        attachments=[("signature.jpg", "image", "jpeg", b"\xff\xd8\xff-bogus-jpeg-bytes")],
    )
    out, _ = _run(tmp_path, raw)
    assert b"bogus-jpeg-bytes" not in out
    assert b"signature.jpg" not in out


def test_attachment_with_a_handler_is_sanitised_and_kept(tmp_path):
    raw = _build(
        headers={"Subject": "hi"},
        body="body",
        attachments=[
            ("notes.html", "application", "octet-stream", b"<html>ring 0298761234</html>")
        ],
    )
    out, _ = _run(tmp_path, raw)
    # Kept under a structural filename but its content sanitised.
    assert b"attachment.html" in out
    parsed = message_from_bytes(out, policy=policy.default)
    payloads = [p.get_payload(decode=True) for p in parsed.iter_parts() if p.get_filename()]
    assert payloads and b"0298761234" not in payloads[0]
    assert b"0400000000" in payloads[0]


def test_docx_attachment_is_sanitised_and_kept(tmp_path):
    """A .docx attachment is routed through the registry handler and re-attached."""
    doc = Document()
    doc.add_paragraph("Call 0298761234 or email real.person@mail.org")
    doc.core_properties.author = "synthetic.author@example.org"
    doc.part.relate_to("mailto:synthetic.link@example.org", RT.HYPERLINK, True)
    buffer = io.BytesIO()
    doc.save(buffer)
    raw = _build(
        headers={"Subject": "hi"},
        body="body",
        attachments=[
            (
                "notes.docx",
                "application",
                "vnd.openxmlformats-officedocument.wordprocessingml.document",
                buffer.getvalue(),
            )
        ],
    )
    out, _ = _run(tmp_path, raw)
    assert b"attachment.docx" in out
    parsed = message_from_bytes(out, policy=policy.default)
    payload = next(
        p.get_payload(decode=True) for p in parsed.iter_parts() if p.get_filename()
    )
    temp = tmp_path / "attachment.docx"
    temp.write_bytes(payload)
    with zipfile.ZipFile(io.BytesIO(payload)) as package:
        assert b"synthetic." not in b"".join(package.read(n) for n in package.namelist())
    text = extract_docx_text(temp)
    assert "0298761234" not in text and "real.person@mail.org" not in text
    assert "0400000000" in text and "john.doe@example.com" in text


def test_pdf_attachment_is_dropped_but_named(tmp_path):
    """A PDF attachment: content gone, but the sanitised filename kept for reference."""
    raw = _build(
        headers={"Subject": "hi"},
        body="body",
        attachments=[
            ("statement.pdf", "application", "pdf", b"%PDF-1.4 real.person@mail.org bytes")
        ],
    )
    out, _ = _run(tmp_path, raw)
    # The PDF bytes are gone, but a safely named placeholder remains.
    assert b"%PDF-1.4" not in out
    assert b"attachment.pdf" in out
    parsed = message_from_bytes(out, policy=policy.default)
    placeholder = next(p for p in parsed.iter_parts() if p.get_filename() == "attachment.pdf")
    assert placeholder.get_content_type() == "text/plain"
    assert "omitted" in placeholder.get_content()


def test_pdf_attachment_with_pii_filename_is_sanitised(tmp_path):
    """A client name in the attachment filename is replaced in the kept placeholder."""
    raw = _build(
        headers={"Subject": "hi"},
        body="body",
        attachments=[("Margaret Whitfield statement.pdf", "application", "pdf", b"%PDF-1.4")],
    )
    out, _ = _run(tmp_path, raw)
    assert b"Margaret Whitfield" not in out
    assert b"attachment.pdf" in out


def test_nested_message_is_sanitised_in_full(tmp_path):
    inner = EmailMessage()
    inner["From"] = "someone@else.com"
    inner["Subject"] = "fwd"
    inner.set_content("old number 0413 111 222")
    raw = _build(headers={"Subject": "hi"}, body="body", nested=inner)
    out, _ = _run(tmp_path, raw)
    text = out.decode()
    assert "someone@else.com" not in text
    assert "0413 111 222" not in text


def test_nested_message_that_is_a_bare_image_is_emptied(tmp_path):
    """A forwarded email that is only an image must not survive inside the wrapper."""
    inner = EmailMessage()
    inner["Subject"] = "scan"
    inner.set_content(b"\xff\xd8\xff-bogus-jpeg", maintype="image", subtype="jpeg")
    raw = _build(headers={"Subject": "hi"}, body="body", nested=inner)
    out, _ = _run(tmp_path, raw)
    # Assert on structure, not raw bytes: a kept image would be base64, so its bytes
    # would not appear literally even if it had survived.
    parsed = message_from_bytes(out, policy=policy.default)
    assert "image/jpeg" not in {p.get_content_type() for p in parsed.walk()}


def test_output_passes_the_leak_rescan(tmp_path):
    """The real acceptance test: nothing the handler emits survives a rescan."""
    inner = EmailMessage()
    inner["From"] = "someone@else.com"
    inner.set_content("old number 0413 111 222")
    raw = _build(
        headers={
            "From": '"Margaret Whitfield" <m.whitfield@client.com>',
            "To": "advisor@firm.com",
            "Subject": "Account 5540912",
            "Message-ID": "<trace123@mail.client.com>",
        },
        body="Call 0412 345 678 or email m.whitfield@client.com",
        html="<html><body><a href='mailto:m.whitfield@client.com'>x</a></body></html>",
        attachments=[
            ("notes.html", "application", "octet-stream", b"<html>ring 0298761234</html>"),
            ("photo.jpg", "image", "jpeg", b"\xff\xd8\xff-bogus"),
        ],
        nested=inner,
    )
    _, result = _run(tmp_path, raw)
    report = scan_output(tmp_path / "out", default_detector_engine(), generated_dummies={
        tmp_path / "out" / "in.eml": frozenset(
            item.surrogate.casefold() for item in result.replacements
        ),
    })
    assert report.clean, [(leak.kind, leak.detector) for leak in report.leaks]


def test_rescan_decodes_pii_inside_a_docx_attachment(tmp_path):
    """Raw base64 must never make an unsanitised embedded document look clean."""
    doc = Document()
    # Isolate decoded-text checking from the default template's custom XML,
    # which is independently refused by the DOCX package policy.
    for rid, relationship in list(doc.part.rels.items()):
        if relationship.reltype == RT.CUSTOM_XML:
            del doc.part.rels[rid]
    doc.add_paragraph("Contact Margaret Whitfield")
    doc.core_properties.author = "synthetic.author@example.org"
    doc.part.relate_to("mailto:synthetic.link@example.org", RT.HYPERLINK, True)
    buffer = io.BytesIO()
    doc.save(buffer)
    message = EmailMessage()
    message["Subject"] = "synthetic"
    message.set_content("safe body")
    message.add_attachment(
        buffer.getvalue(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="attachment.docx",
    )
    path = tmp_path / "unsafe.eml"
    path.write_bytes(message.as_bytes(policy=policy.default))

    report = scan_output(tmp_path, default_detector_engine())
    assert not report.clean
    # Retained metadata/links are now refused even without recognisable PII.
    assert report.unscannable == [path]


def test_output_still_parses_as_a_multipart_email(tmp_path):
    raw = _build(headers={"Subject": "hi"}, body="body", html="<p>hi</p>")
    out, _ = _run(tmp_path, raw)
    parsed = message_from_bytes(out, policy=policy.default)
    assert parsed.is_multipart()
    assert {p.get_content_type() for p in parsed.iter_parts()} == {"text/plain", "text/html"}


def test_no_pii_leaves_the_message_intact(tmp_path):
    raw = _build(headers={"Subject": "meeting notes"}, body="nothing to redact here")
    out, result = _run(tmp_path, raw)
    assert result.replacements == []
    assert b"nothing to redact here" in out
    assert b"meeting notes" in out


@pytest.mark.parametrize('surface', (
    'settings', 'styles', 'fontTable', 'numbering', 'bookmark', 'control',
    'attribute', 'extension', 'namespace', 'orphan_story', 'unknown_part', 'image_alias',
    'relationship_id', 'part_name', 'xml_comment', 'xml_tail', 'stray_text',
))
def test_docx_attachment_retained_surface_audit(tmp_path, surface):
    import hashlib

    from panda.detectors import DetectorEngine
    from test_docx import _audit_document

    attachment = _audit_document(surface)
    raw = _build(body='Visible body', attachments=[(
        'notes.docx', 'application',
        'vnd.openxmlformats-officedocument.wordprocessingml.document', attachment,
    )])
    source = tmp_path / 'source.eml'
    source.write_bytes(raw)
    before = hashlib.sha256(raw).digest()
    destination = tmp_path / 'out' / 'safe.eml'
    engine = DetectorEngine()
    # Independently refuse the unsanitised MIME payload before exercising the handler.
    assert scan_output(tmp_path, engine).unscannable == [source]
    sanitise_eml_file(source, destination, FixedPseudonymiser(), engine)
    assert hashlib.sha256(source.read_bytes()).digest() == before
    message = message_from_bytes(destination.read_bytes(), policy=policy.default)
    payload = next(p.get_payload(decode=True) for p in message.iter_attachments())
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for entry in archive.infolist():
            assert b'synthetic.' not in archive.read(entry)
    assert Document(io.BytesIO(payload)).paragraphs[0].text.startswith('Visible review text')
    assert scan_output(destination.parent, engine).clean


@pytest.mark.parametrize('attachment', [False, True])
def test_html_review_policy_in_mail(tmp_path, attachment):
    from panda.detectors import DetectorEngine
    from panda.handlers.html import extract_html_markup_units

    markup = ('<p>before real.<b>person</b>@mail.org</p>'
              '<svg><text>SECRET</text></svg><img src="data:image/png;base64,SECRET">'
              '<p style="background:url(SECRET)">after</p>')
    raw = _build(body='safe', **(
        {'attachments': [('notes.html', 'text', 'html', markup.encode())]}
        if attachment else {'html': markup}
    ))
    out, _ = _run(tmp_path, raw)
    assert (tmp_path / 'in.eml').read_bytes() == raw
    message = message_from_bytes(out, policy=policy.default)
    payloads = [p.get_payload(decode=True) for p in message.walk()
                if p.get_content_type() == 'text/html']
    assert len(payloads) == 1
    payload = payloads[0]
    assert b'SECRET' not in payload and b'real.' not in payload
    assert [u.text for u in extract_html_markup_units(payload)] == [
        'before john.doe@example.com', 'after']
    assert scan_output(tmp_path / 'out', DetectorEngine()).clean


@pytest.mark.parametrize('attachment', [False, True])
def test_mail_checker_rejects_forbidden_html(tmp_path, attachment):
    from panda.detectors import DetectorEngine

    markup = '<html><body><p>safe</p><img src="data:image/png;base64,AAAA"></body></html>'
    raw = _build(body='safe', **(
        {'attachments': [('notes.html', 'text', 'html', markup.encode())]}
        if attachment else {'html': markup}
    ))
    (tmp_path / 'unsafe.eml').write_bytes(raw)
    report = scan_output(tmp_path, DetectorEngine())
    assert not report.clean
    assert report.unscannable == [tmp_path / 'unsafe.eml']


@pytest.mark.parametrize('surface', [
    'boundary', 'preamble', 'epilogue', 'parameter', 'content_id',
    'child_header', 'container_header', 'duplicate_header', 'encoded_header',
    'nested', 'dropped', 'disposition',
])
def test_mime_metadata_review(tmp_path, surface):
    from panda.detectors import DetectorEngine
    from panda.handlers.eml import extract_eml_units

    secret = 'synthetic.person@example.org'
    message = EmailMessage()
    message.set_content('Safe reading order first')
    message.add_alternative('Safe reading order second')
    child = list(message.iter_parts())[0]
    if surface == 'boundary':
        message.set_boundary(secret)
    elif surface in {'preamble', 'epilogue'}:
        setattr(message, surface, secret)
    elif surface == 'parameter':
        message.set_param('arbitrary', secret)
        child.set_param('arbitrary', secret)
    elif surface == 'content_id':
        message['Content-ID'] = '<' + secret + '>'
        child['Content-ID'] = '<' + secret + '>'
    elif surface in {'child_header', 'container_header', 'duplicate_header'}:
        target = message if surface == 'container_header' else child
        target['X-Private'] = secret
        if surface == 'duplicate_header':
            target['X-Private'] = secret
    elif surface == 'encoded_header':
        child['X-Private'] = '=?utf-8?b?c3ludGhldGljLnBlcnNvbkBleGFtcGxlLm9yZw==?='
    elif surface == 'nested':
        nested = EmailMessage()
        nested['Subject'] = secret
        nested.set_content(secret)
        message.add_attachment(nested)
        list(message.iter_parts())[-1]['Content-ID'] = '<' + secret + '>'
    elif surface == 'dropped':
        message.add_attachment(b'not inspected', maintype='image', subtype='png',
                               filename=secret + '.png')
        list(message.iter_parts())[-1]['X-Private'] = secret
    elif surface == 'disposition':
        child.add_header('Content-Disposition', 'inline', custom=secret)
    raw = message.as_bytes()
    source = tmp_path / 'source.eml'
    destination = tmp_path / 'safe.eml'
    source.write_bytes(raw)
    sanitise_eml_file(source, destination, FixedPseudonymiser(), DetectorEngine())
    assert source.read_bytes() == raw
    output = destination.read_bytes()
    assert secret.encode() not in output
    assert secret not in '\n'.join(unit.text for unit in extract_eml_units(destination))
    parsed = message_from_bytes(output, policy=policy.default)
    assert all(not p.defects for p in parsed.walk())
    text = '\n'.join(p.get_content() for p in parsed.walk()
                     if p.get_content_type() == 'text/plain')
    assert text.index('first') < text.index('second')


@pytest.mark.parametrize('raw', [
    b'Content-Type: multipart/mixed; boundary=missing\n\nno delimiter',
    b'Content-Type: text/plain\nContent-Type: text/html\n\nambiguous',
    b'Content-Transfer-Encoding: base64\n\n%%%invalid',
    b'Content-Transfer-Encoding: base64\n\nYQ',
    b'Content-Transfer-Encoding: quoted-printable\n\nbad=GG',
    b'Content-Transfer-Encoding: unknown\n\ntext',
    b'Content-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: 8bit\n\n\xff',
    b'Content-Type: text/plain; charset=unknown-codec\n\ntext',
    b'not a header\n\ntext',
])
def test_malformed_mime_refused_by_handler_and_checker(tmp_path, raw):
    from panda.detectors import DetectorEngine
    from panda.handlers.eml import extract_eml_units

    source = tmp_path / 'source.eml'
    destination = tmp_path / 'safe.eml'
    source.write_bytes(raw)
    with pytest.raises(ValueError):
        sanitise_eml_file(source, destination, FixedPseudonymiser(), DetectorEngine())
    assert not destination.exists()
    assert source.read_bytes() == raw
    with pytest.raises(ValueError):
        extract_eml_units(source)


@pytest.mark.parametrize('surface', [
    'boundary', 'preamble', 'epilogue', 'parameter', 'content_id',
    'header', 'disposition', 'type', 'version',
])
def test_checker_independently_refuses_unverified_mime(tmp_path, surface):
    from panda.handlers.eml import extract_eml_units

    # Construct independently; no call to the sanitiser creates this fixture.
    part = EmailMessage()
    part.set_content('Safe body', cte='8bit')
    if surface == 'boundary':
        part.make_mixed()
        part.set_boundary('uncontrolled')
    elif surface in {'preamble', 'epilogue'}:
        part.make_mixed()
        part.set_boundary('panda-' + 'a' * 32)
        setattr(part, surface, 'uncontrolled')
    elif surface == 'parameter':
        part.set_param('uncontrolled', 'opaque')
    elif surface == 'content_id':
        part['Content-ID'] = '<opaque>'
    elif surface == 'header':
        part['X-Uncontrolled'] = 'opaque'
    elif surface == 'disposition':
        part.add_header('Content-Disposition', 'attachment', filename='opaque.txt')
    elif surface == 'type':
        part.set_type('text/unknown')
    elif surface == 'version':
        part.replace_header('MIME-Version', '2.0')
    path = tmp_path / 'unsafe.eml'
    path.write_bytes(part.as_bytes())
    with pytest.raises(ValueError):
        extract_eml_units(path)


def test_mime_utf8_and_nested_defect(tmp_path):
    from panda.detectors import DetectorEngine
    from panda.handlers.eml import extract_eml_units

    source = tmp_path / 'source.eml'
    destination = tmp_path / 'safe.eml'
    message = EmailMessage()
    message.set_content('安全 review text', cte='8bit')
    source.write_bytes(message.as_bytes())
    sanitise_eml_file(source, destination, FixedPseudonymiser(), DetectorEngine())
    assert any('安全 review text' in unit.text for unit in extract_eml_units(destination))
    destination.unlink()
    source.write_bytes(b'Content-Type: message/rfc822\n\n'
                       b'Content-Transfer-Encoding: base64\n\n%%%')
    with pytest.raises(ValueError):
        sanitise_eml_file(source, destination, FixedPseudonymiser(), DetectorEngine())
    assert not destination.exists()
    with pytest.raises(ValueError):
        extract_eml_units(source)
