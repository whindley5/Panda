import io
import struct
import zipfile
import zlib

import pytest
from docx import Document

from panda import formats
from panda.detectors import DetectorEngine, default_detector_engine
from panda.handlers import register_defaults
from panda.handlers.docx import extract_docx_text, sanitise_docx_file
from panda.leakcheck import scan_output
from panda.model import PIIKind
from panda.pseudonyms import FixedPseudonymiser


@pytest.fixture(autouse=True)
def clean_registry():
    """Register the extractor so the leak check can read .docx back, then restore."""
    handlers = dict(formats.HANDLERS)
    extractors = dict(formats.TEXT_EXTRACTORS)
    register_defaults()
    yield
    formats.HANDLERS.clear()
    formats.HANDLERS.update(handlers)
    formats.TEXT_EXTRACTORS.clear()
    formats.TEXT_EXTRACTORS.update(extractors)


def _png() -> bytes:
    """A minimal valid 1x1 PNG, enough for python-docx to embed as a picture."""
    def chunk(tag, data):
        body = tag + data
        crc = struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
        return struct.pack(">I", len(data)) + body + crc

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
        + chunk(b"IEND", b"")
    )


def _inject_comment(data: bytes, text: str) -> bytes:
    """Add a word/comments.xml part carrying `text` to a .docx built without one.

    python-docx 1.1.2 has no comment API, so the fixture is assembled at the zip level:
    the part, its content-type override and a document relationship, which is all the
    handler needs to find and sanitise it.
    """
    comments = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:comment w:id="1" w:author="A"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:comment>'
        "</w:comments>"
    ).encode()
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as sink:
        for item in source.infolist():
            blob = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(
                    b"</Types>",
                    b'<Override PartName="/word/comments.xml" ContentType='
                    b'"application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>',
                )
            if item.filename == "word/_rels/document.xml.rels":
                blob = blob.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdCmt" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/comments" Target="comments.xml"/>'
                    b"</Relationships>",
                )
            sink.writestr(item, blob)
        sink.writestr("word/comments.xml", comments)
    return out.getvalue()


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _run(tmp_path, raw, engine=None):
    source = tmp_path / "in.docx"
    destination = tmp_path / "out" / "in.docx"
    source.write_bytes(raw)
    result = sanitise_docx_file(
        source, destination, FixedPseudonymiser(), engine or default_detector_engine()
    )
    return destination, result


def test_replaces_pii_in_a_body_paragraph(tmp_path):
    doc = Document()
    doc.add_paragraph("Email real.person@mail.org or call 0298765432")
    dest, result = _run(tmp_path, _save(doc))
    text = extract_docx_text(dest)
    assert "real.person@mail.org" not in text and "0298765432" not in text
    assert "john.doe@example.com" in text
    kinds = {r.detection.kind for r in result.replacements}
    assert PIIKind.EMAIL in kinds and PIIKind.PHONE in kinds


def test_value_split_across_runs_is_caught(tmp_path):
    """The core reason to sanitise per paragraph, not per run: Word splits values."""
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Call 0412 ")
    paragraph.add_run("345 ")
    paragraph.add_run("678 today")
    dest, result = _run(tmp_path, _save(doc))
    assert [r.detection.kind for r in result.replacements] == [PIIKind.PHONE]
    assert "0400000000" in extract_docx_text(dest)


def test_table_cell_is_sanitised(tmp_path):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Account 5540912 here"
    dest, _ = _run(tmp_path, _save(doc))
    assert "5540912" not in extract_docx_text(dest)


def test_header_and_footer_are_sanitised(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Header advisor@firm.com"
    doc.sections[0].footer.paragraphs[0].text = "Footer 0298765432"
    doc.add_paragraph("body")
    dest, _ = _run(tmp_path, _save(doc))
    text = extract_docx_text(dest)
    assert "advisor@firm.com" not in text and "0298765432" not in text


def test_comment_is_sanitised(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    raw = _inject_comment(_save(doc), "reach me at real.person@mail.org")
    dest, _ = _run(tmp_path, raw)
    assert "real.person@mail.org" not in extract_docx_text(dest)


def test_images_and_thumbnail_are_dropped(tmp_path):
    doc = Document()
    doc.add_paragraph("body")
    doc.add_picture(io.BytesIO(_png()))
    dest, _ = _run(tmp_path, _save(doc))
    names = zipfile.ZipFile(dest).namelist()
    assert not [n for n in names if "media" in n or "thumbnail" in n]


def test_output_passes_the_leak_rescan(tmp_path):
    """The acceptance test: nothing the handler emits survives the docx-aware rescan."""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "advisor@firm.com"
    para = doc.add_paragraph()
    para.add_run("Call 0412 ")
    para.add_run("345 678, client Margaret Whitfield")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Account 5540912"
    doc.add_picture(io.BytesIO(_png()))
    raw = _inject_comment(_save(doc), "note real.person@mail.org")
    engine = default_detector_engine()
    dest, result = _run(tmp_path, raw)
    report = scan_output(dest.parent, engine, generated_dummies={
        dest: frozenset(item.surrogate.casefold() for item in result.replacements),
    })
    assert report.clean, [(leak.kind, leak.detector) for leak in report.leaks]
    assert not report.unscannable


def test_output_reopens_as_a_valid_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Email real.person@mail.org")
    dest, _ = _run(tmp_path, _save(doc))
    reopened = Document(str(dest))
    assert "john.doe@example.com" in reopened.paragraphs[0].text


def test_no_pii_leaves_the_document_unchanged(tmp_path):
    doc = Document()
    run = doc.add_paragraph().add_run("nothing to redact")
    run.bold = True
    dest, result = _run(tmp_path, _save(doc))
    assert result.replacements == []
    reopened = Document(str(dest))
    # Formatting is preserved where no value had to be removed.
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.paragraphs[0].text == "nothing to redact"


def _verification_base_document():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    # Isolate metadata/link verification from the independently forbidden custom
    # XML bibliography shipped in python-docx's default template.
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.CUSTOM_XML:
            del doc.part.rels[rid]
    return doc


def _private_document():
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = _verification_base_document()
    doc.core_properties.author = "synthetic.author@example.org"
    doc.core_properties.last_modified_by = "synthetic.editor@example.org"
    for part in (doc.part, doc.sections[0].header.part):
        rid = part.relate_to("mailto:synthetic.link%40example.org", RT.HYPERLINK, True)
        paragraph = (
            doc.add_paragraph() if part is doc.part else doc.sections[0].header.paragraphs[0]
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Contact synthetic.visible@example.org"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
    doc.part.relate_to("mailto:synthetic.orphan@example.org", RT.HYPERLINK, True)
    return doc


def test_metadata_and_hyperlink_targets_are_removed_from_package(tmp_path):
    raw = _save(_private_document())
    dest, _ = _run(tmp_path, raw)
    with zipfile.ZipFile(dest) as package:
        assert not any(name.startswith("docProps/") for name in package.namelist())
        content = b"\n".join(package.read(name) for name in package.namelist())
        assert b"synthetic." not in content
        assert b"/relationships/hyperlink" not in content
    reopened = Document(str(dest))
    assert "Contact john.doe@example.com" in reopened.paragraphs[0].text
    assert "Contact john.doe@example.com" in reopened.sections[0].header.paragraphs[0].text
    assert (tmp_path / "in.docx").read_bytes() == raw
    assert scan_output(dest.parent, default_detector_engine()).clean


def test_rescan_checks_metadata_and_encoded_relationship_targets(tmp_path):
    path = tmp_path / "unsafe.docx"
    path.write_bytes(_save(_private_document()))
    report = scan_output(tmp_path, default_detector_engine())
    assert not report.clean
    # The closed package policy now refuses these surfaces regardless of PII.
    assert report.unscannable == [path]
    assert "synthetic." not in repr(report)


def test_rescan_checks_orphaned_properties_and_relationships(tmp_path):
    buffer = io.BytesIO(_save(_verification_base_document()))
    with zipfile.ZipFile(buffer, "a") as package:
        package.writestr(
            "docProps/custom.xml",
            '<Properties><value>hidden.person@example.org</value></Properties>',
        )
        package.writestr(
            "word/_rels/orphan.xml.rels",
            '<Relationships><Relationship Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/hyperlink" '
            'Target="mailto:hidden.link%40example.org"/></Relationships>',
        )
    (tmp_path / "unsafe.docx").write_bytes(buffer.getvalue())
    report = scan_output(tmp_path, default_detector_engine())
    assert not report.clean
    assert report.unscannable == [tmp_path / "unsafe.docx"]


def test_custom_and_extended_properties_are_removed(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    doc = Document()
    package = doc.part.package
    custom = Part(
        PackURI('/docProps/custom.xml'),
        'application/vnd.openxmlformats-officedocument.custom-properties+xml',
        b'<Properties><value>synthetic.custom@example.org</value></Properties>',
        package,
    )
    package.relate_to(custom, RT.CUSTOM_PROPERTIES)
    for part in package.iter_parts():
        if str(part.partname) == '/docProps/app.xml':
            part._blob = (
                b'<Properties><Company>synthetic.company@example.org</Company></Properties>'
            )
    dest, _ = _run(tmp_path, _save(doc))
    with zipfile.ZipFile(dest) as archive:
        assert not any(name.startswith('docProps/') for name in archive.namelist())
    assert scan_output(dest.parent, default_detector_engine()).clean


def test_malformed_orphan_metadata_fails_verification(tmp_path):
    buffer = io.BytesIO(_save(Document()))
    with zipfile.ZipFile(buffer, 'a') as archive:
        archive.writestr('docProps/custom.xml', b'<broken')
    path = tmp_path / 'unsafe.docx'
    path.write_bytes(buffer.getvalue())
    report = scan_output(tmp_path, default_detector_engine())
    assert not report.clean
    assert report.unscannable == [path]


def _history_and_fields_document():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph("Review ")
    for kind, text_tag in (("del", "delText"), ("moveFrom", "t"), ("ins", "t"), ("moveTo", "t")):
        revision = OxmlElement(f"w:{kind}")
        revision.set(qn("w:id"), "1")
        revision.set(qn("w:author"), "synthetic.author@example.org")
        revision.set(qn("w:date"), "2026-01-01T00:00:00Z")
        run = OxmlElement("w:r")
        text = OxmlElement(f"w:{text_tag}")
        text.text = f"synthetic.{kind}@example.org "
        run.append(text)
        revision.append(run)
        paragraph._p.append(revision)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'HYPERLINK "mailto:synthetic.simple@example.org"')
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "synthetic.result@example.org"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    for fragment in ('HYPERLINK "mailto:synthetic.', 'split@example.org"'):
        run = OxmlElement("w:r")
        instruction = OxmlElement("w:instrText")
        instruction.text = fragment
        run.append(instruction)
        paragraph._p.append(run)
    return doc


def test_history_and_fields_removed_with_visible_results_sanitised(tmp_path):
    raw = _save(_history_and_fields_document())
    engine = DetectorEngine()
    dest, result = _run(tmp_path, raw, engine)
    with zipfile.ZipFile(dest) as archive:
        content = b"\n".join(archive.read(name) for name in archive.namelist())
    assert b"synthetic." not in content
    assert b"instrText" not in content
    assert b"fldSimple" not in content
    assert b"w:author" not in content
    assert b"w:date" not in content
    assert len([r for r in result.replacements if r.detection.kind == PIIKind.EMAIL]) == 3
    assert "john.doe@example.com" in extract_docx_text(dest)
    Document(str(dest))
    assert (tmp_path / "in.docx").read_bytes() == raw
    assert scan_output(dest.parent, engine).clean


@pytest.mark.parametrize("surface", [
    "del", "moveFrom", "delText", "instrText", "delInstrText", "fldChar", "fldSimple",
    "author", "initials", "date",
])
def test_verifier_refuses_retained_history_or_fields_without_known_pii(tmp_path, surface):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph("ordinary text")
    if surface in {"author", "initials", "date"}:
        paragraph._p.set(qn(f"w:{surface}"), "unrecognised value")
    else:
        paragraph._p.append(OxmlElement(f"w:{surface}"))
    path = tmp_path / "unsafe.docx"
    path.write_bytes(_save(doc))
    report = scan_output(tmp_path, default_detector_engine())
    assert not report.clean
    assert report.unscannable == [path]


@pytest.mark.parametrize("story", ["header", "comment"])
def test_history_policy_applies_outside_document_body(tmp_path, story):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    if story == "header":
        paragraph = doc.sections[0].header.paragraphs[0]
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "synthetic.header@example.org")
        paragraph._p.append(field)
        raw = _save(doc)
    else:
        raw = _inject_comment(_save(doc), "synthetic.comment@example.org")
        # The generic comment part's author attribute must be stripped too.
    engine = DetectorEngine()
    dest, _ = _run(tmp_path, raw, engine)
    with zipfile.ZipFile(dest) as archive:
        content = b"\n".join(archive.read(name) for name in archive.namelist())
    assert b"synthetic." not in content
    assert b"w:author" not in content
    assert scan_output(dest.parent, engine).clean


@pytest.mark.parametrize("relationship", [
    "customXml", "oleObject", "package", "aFChunk", "control", "attachedTemplate", "subDocument",
])
def test_unsupported_payload_removed_through_all_incoming_relationships(tmp_path, relationship):
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    doc = Document()
    doc.add_paragraph("Review text")
    payload = Part(
        PackURI('/word/opaque.xml'), 'application/xml',
        b'<private>synthetic.payload@example.org</private>', doc.part.package,
    )
    doc.part.relate_to(
        payload, f'http://schemas.openxmlformats.org/officeDocument/2006/relationships/{relationship}'
    )
    # An unexpected alias must not keep the same payload reachable after the
    # recognised incoming relationship is removed.
    doc.sections[0].header.part.relate_to(payload, 'urn:synthetic:alias')
    raw = _save(doc)
    engine = DetectorEngine()
    dest, _ = _run(tmp_path, raw, engine)
    with zipfile.ZipFile(dest) as archive:
        assert 'word/opaque.xml' not in archive.namelist()
        assert not any(name.startswith('customXml/') for name in archive.namelist())
        content = b'\n'.join(archive.read(name) for name in archive.namelist())
    assert b'synthetic.payload' not in content
    assert b'urn:synthetic:alias' not in content
    assert (tmp_path / 'in.docx').read_bytes() == raw
    assert Document(str(dest)).paragraphs[0].text == 'Review text'
    assert scan_output(dest.parent, engine).clean


def test_external_template_and_custom_binding_references_removed(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph('Review text')
    settings = doc.part._settings_part
    rid = settings.relate_to(
        'file:///synthetic.template@example.org/template.dotm',
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate',
        True,
    )
    reference = OxmlElement('w:attachedTemplate')
    reference.set(qn('r:id'), rid)
    settings.element.append(reference)
    binding = OxmlElement('w:dataBinding')
    binding.set(qn('w:xpath'), 'synthetic.binding@example.org')
    doc.paragraphs[0]._p.append(binding)
    dest, _ = _run(tmp_path, _save(doc), DetectorEngine())
    with zipfile.ZipFile(dest) as archive:
        content = b'\n'.join(archive.read(name) for name in archive.namelist())
    assert b'synthetic.' not in content
    assert b'attachedTemplate' not in content
    assert b'dataBinding' not in content


@pytest.mark.parametrize('name,content_type', [
    ('customXml/orphan.xml', 'application/xml'),
    ('word/embeddings/orphan.bin', 'application/octet-stream'),
    ('word/activeX/orphan.bin', 'application/octet-stream'),
    ('word/opaque.bin', 'application/vnd.openxmlformats-officedocument.oleObject'),
    ('word/opaque.xml', 'application/vnd.openxmlformats-officedocument.customXmlProperties+xml'),
    ('word/opaque.bin', 'application/vnd.ms-office.vbaProject'),
])
def test_verifier_refuses_orphan_payloads_without_detectable_pii(tmp_path, name, content_type):
    doc = _verification_base_document()
    source = io.BytesIO(_save(doc))
    output = io.BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, 'w') as sink:
        for entry in archive.infolist():
            data = archive.read(entry)
            if entry.filename == '[Content_Types].xml':
                data = data.replace(
                    b'</Types>',
                    f'<Override PartName="/{name}" ContentType="{content_type}"/></Types>'.encode(),
                )
            sink.writestr(entry, data)
        sink.writestr(name, b'unrecognised private content')
    path = tmp_path / 'unsafe.docx'
    path.write_bytes(output.getvalue())
    report = scan_output(tmp_path, DetectorEngine())
    assert not report.clean
    assert report.unscannable == [path]


@pytest.mark.parametrize('relationship,external', [
    ('oleObject', False), ('attachedTemplate', True),
])
def test_verifier_refuses_orphan_unsupported_relationships(tmp_path, relationship, external):
    buffer = io.BytesIO(_save(_verification_base_document()))
    with zipfile.ZipFile(buffer, 'a') as archive:
        mode = ' TargetMode="External"' if external else ''
        archive.writestr(
            'word/_rels/orphan.xml.rels',
            '<Relationships><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            f'{relationship}" Target="opaque"{mode}/></Relationships>',
        )
    path = tmp_path / 'unsafe.docx'
    path.write_bytes(buffer.getvalue())
    assert scan_output(tmp_path, DetectorEngine()).unscannable == [path]


_AUDIT_SURFACES = (
    'settings', 'styles', 'fontTable', 'numbering', 'bookmark', 'control',
    'attribute', 'extension', 'namespace', 'orphan_story', 'unknown_part', 'image_alias',
    'relationship_id', 'part_name', 'xml_comment', 'xml_tail', 'stray_text',
)


def _audit_document(surface):
    """Entirely synthetic adversarial containers, including unreachable ZIP entries."""
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    doc = Document()
    paragraph = doc.add_paragraph('Visible review text')
    doc.sections[0].header.paragraphs[0].text = 'Visible header'
    private = 'synthetic.secret@example.org'
    if surface in {'settings', 'styles', 'fontTable', 'numbering'}:
        part = next(p for p in doc.part.package.iter_parts()
                    if str(p.partname) == f'/word/{surface}.xml')
        root = etree.fromstring(part.blob)
        root.set('private', private)
        if hasattr(part, 'element'):
            part.element.set('private', private)
        else:
            part._blob = etree.tostring(root)
    elif surface in {'bookmark', 'control'}:
        node = OxmlElement('w:bookmarkStart' if surface == 'bookmark' else 'w:sdt')
        node.set(qn('w:name'), private)
        if surface == 'control':
            content = OxmlElement('w:sdtContent')
            run = OxmlElement('w:r')
            text = OxmlElement('w:t')
            text.text = ' Visible control synthetic.visible@example.org'
            run.append(text)
            content.append(run)
            node.append(content)
        paragraph._p.append(node)
    elif surface == 'attribute':
        paragraph._p.set(qn('w:rsidR'), private)
    elif surface == 'extension':
        etree.SubElement(paragraph._p, '{urn:private}payload').text = private
    elif surface == 'namespace':
        paragraph._p.append(etree.Element(qn('w:r'), nsmap={'private': 'urn:' + private}))
    elif surface == 'xml_comment':
        paragraph._p.append(etree.Comment(private))
    elif surface == 'xml_tail':
        paragraph._p[0].tail = private
    elif surface == 'stray_text':
        etree.SubElement(doc.element.body, qn('w:t')).text = private
    elif surface == 'image_alias':
        doc.add_picture(io.BytesIO(_png()))
        image = next(p for p in doc.part.package.iter_parts() if p.content_type == 'image/png')
        image._blob += private.encode()
        doc.sections[0].header.part.relate_to(image, 'urn:alias')
    elif surface == 'unknown_part':
        part = Part(PackURI('/word/unknown.xml'), 'application/xml',
                    f'<private>{private}</private>'.encode(), doc.part.package)
        doc.part.relate_to(part, 'urn:alias')
    elif surface == 'relationship_id':
        rels = doc.part.rels
        rid = next(rid for rid, rel in rels.items() if rel.reltype.endswith('/header'))
        rel = rels.pop(rid)
        rels.add_relationship(rel.reltype, rel.target_part, private)
        ref = next(doc.element.iter(qn('w:headerReference')))
        ref.set(qn('r:id'), private)
    elif surface == 'part_name':
        doc.sections[0].header.part._partname = PackURI('/word/' + private + '.xml')
    raw = _save(doc)
    if surface == 'orphan_story':
        buffer = io.BytesIO(raw)
        with zipfile.ZipFile(buffer, 'a') as archive:
            archive.writestr('word/orphan.xml',
                            f'<w:hdr xmlns:w="{qn("w:p")[1:].split("}")[0]}">'
                            f'<w:p><w:r><w:t>{private}</w:t></w:r></w:p></w:hdr>')
        raw = buffer.getvalue()
    return raw


@pytest.mark.parametrize('surface', _AUDIT_SURFACES)
def test_complete_retained_surface_audit(tmp_path, surface):
    import hashlib

    raw = _audit_document(surface)
    before = hashlib.sha256(raw).digest()
    destination, _ = _run(tmp_path, raw, DetectorEngine())
    assert hashlib.sha256((tmp_path / 'in.docx').read_bytes()).digest() == before
    with zipfile.ZipFile(destination) as archive:
        for entry in archive.infolist():
            assert b'synthetic.' not in archive.read(entry)
            assert 'synthetic.' not in entry.filename
    reopened = Document(destination)
    assert reopened.paragraphs[0].text.startswith('Visible review text')
    assert reopened.sections[0].header.paragraphs[0].text == 'Visible header'
    if surface == 'control':
        assert 'Visible control john.doe@example.com' in reopened.paragraphs[0].text
    assert scan_output(destination.parent, DetectorEngine()).clean


@pytest.mark.parametrize('surface', _AUDIT_SURFACES)
def test_independent_verifier_refuses_unaudited_source_surfaces(tmp_path, surface):
    path = tmp_path / 'unsafe.docx'
    path.write_bytes(_audit_document(surface))
    report = scan_output(tmp_path, DetectorEngine())
    assert not report.clean
    assert report.unscannable == [path]
    assert 'synthetic.' not in repr(report)


@pytest.mark.parametrize('surface', (
    'attribute', 'extension', 'namespace', 'comment', 'tail', 'stray_text',
    'nested_paragraph', 'orphan', 'alias', 'duplicate', 'zip_comment', 'doctype',
    'content_type', 'relationship', 'settings', 'styles', 'fontTable', 'numbering',
))
def test_verifier_rejects_tampering_with_otherwise_safe_package(tmp_path, surface):
    from docx.oxml.ns import qn
    from lxml import etree

    doc = Document()
    doc.add_paragraph('Visible review text')
    destination, _ = _run(tmp_path, _save(doc), DetectorEngine())
    with zipfile.ZipFile(destination) as archive:
        entries = {n: archive.read(n) for n in archive.namelist()}
    root = etree.fromstring(entries['word/document.xml'])
    paragraph = next(root.iter(qn('w:p')))
    # Unknown values intentionally contain no recognised PII: structural policy
    # must reject them independently of detector coverage.
    if surface == 'attribute':
        paragraph.set(qn('w:rsidR'), 'opaque')
    elif surface == 'extension':
        etree.SubElement(paragraph, '{urn:opaque}data').text = 'opaque'
    elif surface == 'namespace':
        paragraph.append(etree.Element(qn('w:r'), nsmap={'extra': 'urn:opaque'}))
    elif surface == 'comment':
        paragraph.append(etree.Comment('opaque'))
    elif surface == 'tail':
        paragraph[0].tail = 'opaque'
    elif surface == 'stray_text':
        etree.SubElement(root[0], qn('w:t')).text = 'opaque'
    elif surface == 'nested_paragraph':
        paragraph.append(etree.Element(qn('w:p')))
    elif surface in {'orphan', 'alias', 'settings', 'styles', 'fontTable', 'numbering'}:
        entries[f'word/{surface}.xml'] = b'<opaque/>'
    elif surface == 'content_type':
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(
            b'document.main+xml', b'opaque+xml')
    elif surface == 'relationship':
        entries['_rels/.rels'] = entries['_rels/.rels'].replace(b'rId0', b'opaque')
    entries['word/document.xml'] = etree.tostring(
        root, xml_declaration=True, encoding='UTF-8', standalone=True)
    if surface == 'doctype':
        entries['word/document.xml'] = entries['word/document.xml'].replace(
            b'<w:document', b'<!DOCTYPE document [<!ENTITY hidden "opaque">]><w:document', 1)
    with zipfile.ZipFile(destination, 'w') as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
        if surface == 'duplicate':
            with pytest.warns(UserWarning, match='Duplicate name'):
                archive.writestr('word/document.xml', entries['word/document.xml'])
        if surface == 'zip_comment':
            archive.comment = b'opaque'
    assert scan_output(destination.parent, DetectorEngine()).unscannable == [destination]
