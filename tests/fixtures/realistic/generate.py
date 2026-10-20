"""Deterministic synthetic corpus for end-to-end PII-survival evaluation.

Everything here is invented. Written independently of the detector code: content is
modelled on how brokerage customer-service chats/emails look, not on what detectors
match. Run: python tests/fixtures/realistic/generate.py
"""
from __future__ import annotations

import html
import json
import shutil
import sys
from datetime import datetime, timedelta
from email.headerregistry import Address
from email.message import EmailMessage
from pathlib import Path

HERE = Path(__file__).resolve().parent
CORPUS = HERE / "corpus"
LABELS = HERE / "labels.json"


def cn_id(area: str, birth: str, seq: str) -> str:
    """Checksum-valid (ISO 7064 MOD 11-2) fictional 18-digit Chinese ID number."""
    body = area + birth + seq
    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    total = sum(int(c) * w for c, w in zip(body, weights, strict=False))
    return body + "10X98765432"[total % 11]


def build_pdf(lines: list[str], per_page: int = 42) -> bytes:
    """Minimal Latin-only base-14 Helvetica PDF (same approach as tests/test_pdf.py)."""
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[""]]
    objects: dict[int, bytes] = {
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    }
    kids, nxt = [], 4
    for pl in pages:
        page_obj, content_obj = nxt, nxt + 1
        nxt += 2
        kids.append(page_obj)
        parts = ["BT", "/F1 11 Tf", "14 TL", "1 0 0 1 40 760 Tm"]
        for i, line in enumerate(pl):
            esc = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            parts.append(f"({esc}) Tj" if i == 0 else f"T* ({esc}) Tj")
        parts.append("ET")
        content = "\n".join(parts).encode("latin-1")
        objects[page_obj] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content_obj
        )
        objects[content_obj] = b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream"
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    out = bytearray(b"%PDF-1.4\n")
    offs = {}
    for n in sorted(objects):
        offs[n] = len(out)
        out += b"%d 0 obj\n" % n + objects[n] + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for n in sorted(objects):
        out += b"%010d 00000 n \n" % offs[n]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def stamp(day: str, start: str, n: int) -> list[str]:
    t = datetime.fromisoformat(f"{day} {start}")
    out = []
    for i in range(n):
        out.append(t.strftime("%Y-%m-%d %H:%M:%S"))
        t += timedelta(seconds=23 + (i * 17) % 41)
    return out


def chat(day, start, msgs):
    ts = stamp(day, start, len(msgs))
    return [(r, h, ts[i], tx) for i, (r, h, tx) in enumerate(msgs)]


def chat_lines(msgs, glue=False):
    lines = []
    for i, (role, h, ts, text) in enumerate(msgs):
        hdr = f"[{role}]{h} {ts}"
        if glue and i > 0 and i % 2 == 1:
            lines[-1] += hdr  # PDF extraction glues a header onto the previous line
        else:
            lines.append(hdr)
        lines.extend(text.split("\n"))
    return lines


DOCS: list[dict] = []


def add(name, fmt, lang, pii, keep, plain, payload):
    plain_text = plain if isinstance(plain, str) else "\n".join(plain)
    for s in (pii + keep) if plain_text else []:
        assert s in plain_text, f"{name}: label not in source: {s!r}"
    DOCS.append(dict(name=name, fmt=fmt, lang=lang, pii=pii, keep=keep, payload=payload))


# --------------------------------------------------------------------------- PDF chats
def pdf_chats():
    m = chat("2026-05-02", "10:14:03", [
        ("Customer", "MarkT", "Hi, can someone call me back? I cannot log in to my trading account."),
        ("agent", "AU-0031", "Sure Mark, may I have your full name and best number please"),
        ("Customer", "MarkT", "Mark Tanner, 0412 555 019. Email is mark.tanner77@example.com"),
        ("agent", "AU-0031", "Thanks Tanner, can you confirm your date of birth?"),
        ("Customer", "MarkT", "14/03/1984"),
        ("agent", "AU-0031", "Thanks Mark~ Which account number is it?"),
        ("Customer", "MarkT", "Account 20481937. I live at 42 Wattle Crescent, Parramatta NSW 2150"),
        ("AI", "CraneAI", "Summary: Mr Tanner requested a callback about a locked login."),
        ("agent", "AU-0031", "Mr Tanner, we have reset it. Anything else, Mark?"),
        ("Customer", "MarkT", "The error appeared at 07:13:16\nPlease check the trade dated 4 May 2026 as well"),
        ("Customer", "MarkT", "Also my referral code from a friend is not working, the code is MTAN-8842"),
        ("agent", "AU-0031", "Mark, the Crane team will follow up by email."),
        ("agent", "system", "Session closed by system"),
    ])
    lines = ["Live chat transcript - export 2026-05-02"] + chat_lines(m, glue=True)
    add("pdf_en_camelcase_mark", "pdf", "en",
        ["MarkT", "Mark Tanner", "Tanner", "Mark", "0412 555 019", "mark.tanner77@example.com",
         "14/03/1984", "20481937", "42 Wattle Crescent, Parramatta NSW 2150", "MTAN-8842"],
        ["AU-0031", "07:13:16", "Please check the trade", "4 May 2026", "referral code from a friend",
         "Crane team", "system"], lines, build_pdf(lines))

    m = chat("2026-05-03", "14:02:11", [
        ("Customer", "Priya Nair", "hello i sent money yesterday and it hasnt shown up"),
        ("agent", "AU-0044", "Hi Priya, sorry to hear that. Can I get the reference and your mobile?"),
        ("Customer", "Priya Nair", "ref is NAIR-2026-0417, mobile +61 430 555 218"),
        ("agent", "AU-0044", "Thank you Ms Nair. Is priya.nair@example.net still your email?"),
        ("Customer", "Priya Nair", "yes. Born 9 November 1992 if you need it"),
        ("agent", "AU-0044", "Priya, could you also confirm the address on file? Nair is not matching"),
        ("Customer", "Priya Nair", "Unit 5/18 Corella Street, Newtown NSW 2042"),
        ("agent", "AU-0044", "Thanks. Priya? I have escalated to the team."),
    ])
    lines = ["Live chat transcript - export 2026-05-03"] + chat_lines(m, glue=True)
    add("pdf_en_fullname_handle_priya", "pdf", "en",
        ["Priya Nair", "Priya", "Nair", "NAIR-2026-0417", "+61 430 555 218", "priya.nair@example.net",
         "9 November 1992", "Unit 5/18 Corella Street, Newtown NSW 2042"],
        ["AU-0044", "team"], lines, build_pdf(lines))

    m = chat("2026-05-04", "09:31:40", [
        ("Customer", "0412555019", "Hi this is Sam, I need to update my phone number"),
        ("agent", "AU-0027", "Hello Sam, what is the new number?"),
        ("Customer", "0412555019", "0498 555 302. and can you use sam.holloway@example.org"),
        ("agent", "AU-0027", "Done. Holloway, is the postal address still 7 Ironbark Lane, Kew VIC 3101?"),
        ("Customer", "0412555019", "yes. Sam Holloway is right, thanks"),
        ("agent", "AU-0027", "Thanks Sam, anything else?"),
        ("Customer", "0412555019", "no all good. Payment due 15/05/2026 right?"),
    ])
    lines = ["Live chat transcript - export 2026-05-04"] + chat_lines(m)
    add("pdf_en_phone_handle_sam", "pdf", "en",
        ["0412555019", "Sam", "0498 555 302", "sam.holloway@example.org", "Holloway", "Sam Holloway",
         "7 Ironbark Lane, Kew VIC 3101"],
        ["AU-0027", "15/05/2026"], lines, build_pdf(lines))

    m = chat("2026-05-05", "18:45:02", [
        ("Customer", "jenny_wu88", "hi there, why was my withdrawal declined"),
        ("agent", "AU-0031", "Hello, may I have your name please?"),
        ("Customer", "jenny_wu88", "Jenny Wu"),
        ("agent", "AU-0031", "Thanks Jenny. Ms Wu, can you confirm date of birth and the last 4 digits of your account?"),
        ("Customer", "jenny_wu88", "DOB 22-07-1988, account 00913377"),
        ("agent", "AU-0031", "Jenny, the bank name on the withdrawal did not match. Wu Jenny vs Jenny Wu"),
        ("Customer", "jenny_wu88", "oh my card is under my maiden name Jenny Lau. call me 0421 555 764"),
        ("agent", "AU-0031", "Understood Jenny~ we will call you back on 0421 555 764"),
        ("agent", "system", "Chat ended"),
    ])
    lines = ["Live chat transcript - export 2026-05-05"] + chat_lines(m, glue=True)
    add("pdf_en_lowercase_jenny", "pdf", "en",
        ["jenny_wu88", "Jenny Wu", "Jenny", "Ms Wu", "22-07-1988", "00913377", "Jenny Lau",
         "0421 555 764", "Wu Jenny"],
        ["AU-0031", "system"], lines, build_pdf(lines))


# --------------------------------------------------------------------------- HTML chats
def html_chat(title, msgs, style):
    rows = []
    for role, h, ts, text in msgs:
        body = "<br>".join(html.escape(t) for t in text.split("\n"))
        if style == "spans":
            rows.append(f'<div class="msg"><span class="role">[{role}]</span>'
                        f'<span class="handle">{html.escape(h)}</span> '
                        f'<span class="ts">{ts}</span><p>{body}</p></div>')
        else:
            rows.append(f'<p><b>[{role}]{html.escape(h)} {ts}</b><br>{body}</p>')
    return (f'<!DOCTYPE html>\n<html><head><meta charset="utf-8"><title>{html.escape(title)}</title>'
            f'</head><body><h1>{html.escape(title)}</h1>\n' + "\n".join(rows) + "\n</body></html>\n")


def html_chats():
    m = chat("2026-05-06", "11:20:15", [
        ("Customer", "8f3a91c2", "Hello my name is David Kowalski, I want to close my account"),
        ("agent", "AU-0058", "Hi David, sorry to see you go. Can you verify your DOB?"),
        ("Customer", "8f3a91c2", "03/12/1975"),
        ("agent", "AU-0058", "Thanks Kowalski. Please confirm the mobile number on file"),
        ("Customer", "8f3a91c2", "0411 555 846 and dave.kowalski@example.com"),
        ("agent", "AU-0058", "Thank you David. And your home address? Dave, this is just to confirm"),
        ("Customer", "8f3a91c2", "3 Rosella Court, Glen Waverley VIC 3150, account 30177264"),
        ("AI", "Assistant", "Reminder: closing an account takes 5 business days."),
        ("agent", "AU-0058", "David？ is there a reason you are leaving?"),
        ("Customer", "8f3a91c2", "just moving overseas, Dave K is fine"),
    ])
    add("html_en_idhandle_david", "html", "en",
        ["David Kowalski", "David", "Kowalski", "Dave", "03/12/1975", "0411 555 846",
         "dave.kowalski@example.com", "3 Rosella Court, Glen Waverley VIC 3150", "30177264", "Dave K",
         "8f3a91c2"],
        # A customer's platform ID identifies them as surely as a name does.
        ["AU-0058"], "", html_chat("Chat export", m, "spans").encode())
    DOCS[-1]["payload"] = html_chat("Chat export", m, "spans").encode()
    _fix_plain(DOCS[-1], DOCS[-1]["payload"].decode())

    m = chat("2026-05-07", "16:05:59", [
        ("Customer", "SarahJ", "Hi, my friend gave me a referral code from a friend: SJ-FR-771. how do i enter it"),
        ("agent", "AU-0031", "Hi Sarah, you can enter it in Settings. What is your registered email?"),
        ("Customer", "SarahJ", "sarah.jennings@example.org"),
        ("agent", "AU-0031", "Thanks Jennings. Can I also get your phone?"),
        ("Customer", "SarahJ", "(03) 9555 0188"),
        ("agent", "AU-0031", "Sarah, for security what is your date of birth?"),
        ("Customer", "SarahJ", "1 February 1990"),
        ("agent", "AU-0031", "Perfect. Sarah Jennings, your code has been applied. The Crane team says thanks."),
    ])
    add("html_en_camelcase_sarah", "html", "en",
        ["SarahJ", "SJ-FR-771", "Sarah", "sarah.jennings@example.org", "Jennings", "(03) 9555 0188",
         "1 February 1990", "Sarah Jennings"],
        ["AU-0031", "referral code from a friend", "Crane team"], "", b"")
    b = html_chat("Chat export", m, "bold").encode()
    DOCS[-1]["payload"] = b
    _fix_plain(DOCS[-1], b.decode())

    # Chinese chats
    id1 = cn_id("110105", "19850312", "002")
    m = chat("2026-05-08", "10:02:30", [
        ("Customer", "Meiling_W", "你好，我想问一下入金怎么还没到账"),
        ("agent", "AU-0112", "您好，请问怎么称呼？"),
        ("Customer", "Meiling_W", "我叫王美玲，手机13812345678"),
        ("agent", "AU-0112", "好的王女士，请问您的身份证号码是多少？"),
        ("Customer", "Meiling_W", id1),
        ("agent", "AU-0112", "美玲，您的邮箱是？"),
        ("Customer", "Meiling_W", "meiling.wang@example.com，生日是1985年3月12日"),
        ("agent", "AU-0112", "谢谢美玲~ 我们的客服团队会在24小时内处理。王女士还有其他问题吗？"),
        ("Customer", "Meiling_W", "没有了，谢谢。系统提示7月1日到账，请帮忙看一下"),
        ("agent", "system", "会话已结束"),
    ])
    add("html_zh_meiling_wang", "html", "zh",
        ["Meiling_W", "王美玲", "美玲", "王女士", "13812345678", id1, "meiling.wang@example.com",
         "1985年3月12日"],
        ["AU-0112", "客服", "system"], "", b"")
    b = html_chat("聊天记录", m, "spans").encode()
    DOCS[-1]["payload"] = b
    _fix_plain(DOCS[-1], b.decode())

    m = chat("2026-05-09", "13:40:00", [
        ("Customer", "guoqiang1979", "客服你好，我要修改银行卡"),
        ("agent", "AU-0067", "李先生您好，请问新的银行卡号是多少？"),
        ("Customer", "guoqiang1979", "6222 0212 3456 7890 123，户名李国强"),
        ("agent", "AU-0067", "好的国强先生，请确认联系电话"),
        ("Customer", "guoqiang1979", "+86 139 8765 4321，邮箱 li.guoqiang@example.net"),
        ("agent", "AU-0067", "收到。地址还是上海市浦东新区世纪大道100号吗？"),
        ("Customer", "guoqiang1979", "对，是的。李先生我的出生日期1979年6月15日，你们核对一下"),
        ("agent", "AU-0067", "国强，已经提交了，请等待审核？"),
    ])
    add("html_zh_lowercase_handle_li", "html", "zh",
        ["guoqiang1979", "李先生", "6222 0212 3456 7890 123", "李国强", "国强", "+86 139 8765 4321",
         "li.guoqiang@example.net", "上海市浦东新区世纪大道100号", "1979年6月15日"],
        ["AU-0067", "客服"], "", b"")
    b = html_chat("聊天记录", m, "bold").encode()
    DOCS[-1]["payload"] = b
    _fix_plain(DOCS[-1], b.decode())

    m = chat("2026-05-10", "20:12:44", [
        ("Customer", "KevinC", "Hi 我是Kevin，想问一下出金手续费"),
        ("agent", "AU-0031", "Kevin您好，请问全名是？"),
        ("Customer", "KevinC", "Kevin Chen，中文名陈凯文"),
        ("agent", "AU-0031", "谢谢陈先生。请提供手机号"),
        ("Customer", "KevinC", "0455 123 456"),
        ("agent", "AU-0031", "凯文，我们这边显示地址是Unit 12, 88 Victoria Avenue, Chatswood NSW 2067，对吗？"),
        ("Customer", "KevinC", "对的。邮箱kevin.chen@example.com"),
        ("agent", "AU-0031", "好的Kevin，手续费是每笔15澳元。还有问题吗？"),
        ("AI", "CraneAI", "本次对话摘要：陈先生咨询出金手续费。"),
    ])
    add("html_zh_mixed_kevin", "html", "zh",
        ["KevinC", "Kevin", "Kevin Chen", "陈凯文", "陈先生", "0455 123 456", "凯文",
         "Unit 12, 88 Victoria Avenue, Chatswood NSW 2067", "kevin.chen@example.com"],
        ["AU-0031", "CraneAI"], "", b"")
    b = html_chat("聊天记录", m, "spans").encode()
    DOCS[-1]["payload"] = b
    _fix_plain(DOCS[-1], b.decode())

    id2 = cn_id("310115", "19910824", "045")
    m = chat("2026-05-11", "08:55:10", [
        ("Customer", "13912345678", "你好，账户被冻结了"),
        ("agent", "AU-0112", "您好，请问您贵姓？"),
        ("Customer", "13912345678", "姓赵，赵雪梅。身份证" + id2),
        ("agent", "AU-0112", "赵女士，请稍等，我帮您查一下"),
        ("Customer", "13912345678", "好的雪梅不急"),
        ("agent", "AU-0112", "雪梅～ 请问您的账户号码是70312298？"),
        ("Customer", "13912345678", "是的，我现住北京市朝阳区建国路88号3号楼1201室"),
        ("agent", "AU-0112", "谢谢，赵女士。已经解冻。"),
    ])
    add("html_zh_phone_handle_zhao", "html", "zh",
        ["13912345678", "赵雪梅", "赵女士", "雪梅", id2, "70312298", "北京市朝阳区建国路88号3号楼1201室"],
        ["AU-0112"], "", b"")
    b = html_chat("聊天记录", m, "bold").encode()
    DOCS[-1]["payload"] = b
    _fix_plain(DOCS[-1], b.decode())


def _fix_plain(doc, plain):
    for s in doc["pii"] + doc["keep"]:
        assert s in plain, f"{doc['name']}: label not in source: {s!r}"


# --------------------------------------------------------------------------- DOCX chats / memos
def make_docx(path, paras=(), header=None, footer=None, tables=()):
    from docx import Document
    d = Document()
    sec = d.sections[0]
    if header:
        sec.header.paragraphs[0].text = header
    if footer:
        sec.footer.paragraphs[0].text = footer
    for p in paras:
        d.add_paragraph(p)
    for rows in tables:
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                t.cell(i, j).text = c
    d.save(str(path))


def docx_doc(name, lang, pii, keep, paras=(), header=None, footer=None, tables=()):
    plain = "\n".join(list(paras) + [header or "", footer or ""] +
                      [c for t in tables for r in t for c in r])
    add(name, "docx", lang, pii, keep, plain,
        lambda p, a=(paras, header, footer, tables): make_docx(p, *a[:1], header=a[1], footer=a[2], tables=a[3]))


def docx_docs():
    m = chat("2026-05-12", "12:00:05", [
        ("Customer", "XiaoMing_L", "你好，我想开通期权交易权限"),
        ("agent", "AU-0067", "您好，请问贵姓？"),
        ("Customer", "XiaoMing_L", "刘小明，电话 0433 555 917"),
        ("agent", "AU-0067", "刘先生您好，请提供邮箱和出生日期"),
        ("Customer", "XiaoMing_L", "xiaoming.liu@example.org 1993-11-02"),
        ("agent", "AU-0067", "小明，申请已提交。"),
        ("Customer", "XiaoMing_L", "谢谢！我的推荐码是 XM-2291，朋友给的 referral code from a friend"),
    ])
    lines = chat_lines(m)
    docx_doc("docx_zh_chat_xiaoming", "zh",
             ["XiaoMing_L", "刘小明", "0433 555 917", "刘先生", "xiaoming.liu@example.org", "1993-11-02",
              "小明", "XM-2291"],
             ["AU-0067", "referral code from a friend"], paras=lines)

    m = chat("2026-05-13", "15:15:15", [
        ("Customer", "Tom Ridley", "Hi I am Tom, did my deposit clear?"),
        ("agent", "AU-0044", "Hi Tom, what is your account number?"),
        ("Customer", "Tom Ridley", "40518826"),
        ("agent", "AU-0044", "Mr Ridley yes it cleared at 06:41:09\nPlease keep the receipt."),
        ("Customer", "Tom Ridley", "great thanks, call me on 07 3555 0129 if anything else"),
    ])
    docx_doc("docx_en_chat_tom", "en",
             ["Tom Ridley", "Tom", "40518826", "Mr Ridley", "07 3555 0129"],
             ["AU-0044", "06:41:09", "Please keep the receipt."], paras=chat_lines(m))

    docx_doc("docx_en_memo_amelia", "en",
             ["Amelia Foster", "amelia.foster@example.com", "0402 555 613", "Amelia", "Foster",
              "28/09/1987", "Ms Foster"],
             ["Crane", "AU-0031", "team", "2 May 2026", "15 June 2026"],
             header="File note - Client: Amelia Foster",
             footer="Prepared for Amelia Foster - Confidential - Crane internal",
             paras=["MEMO", "Date: 2 May 2026",
                    "Ms Foster called about her margin call. Amelia was upset but calm by the end.",
                    "Follow-up due 15 June 2026 by AU-0031 and the compliance team.",
                    "Identity check passed with DOB 28/09/1987."],
             tables=[[["Field", "Value"], ["Name", "Amelia Foster"],
                      ["Email", "amelia.foster@example.com"], ["Mobile", "0402 555 613"]]])

    docx_doc("docx_en_letter_rahman", "en",
             ["Mohammed Rahman", "Mr Rahman", "17 Banksia Road, Box Hill VIC 3128", "Mohammed",
              "Rahman", "0411 555 774", "Moe"],
             ["Crane", "team", "12 May 2026"],
             paras=["Crane Brokerage Client Services", "12 May 2026",
                    "Mohammed Rahman", "17 Banksia Road, Box Hill VIC 3128", "",
                    "Dear Mr Rahman,",
                    "Thank you for contacting us. As discussed, Mohammed, your request has been noted.",
                    "If you have questions please call us or reply. We have your number as 0411 555 774.",
                    "Our team knows you as Moe from earlier correspondence.",
                    "Yours sincerely,", "Client Services team"])

    docx_doc("docx_en_filenote_grech", "en",
             ["Fiona Grech", "Ms Grech", "Fiona", "Grech", "fgrech@example.net", "0419 555 302",
              "1 April 1979", "062-000", "12345678"],
             ["AU-0058", "system", "Crane", "4 May 2026"],
             header="Case notes", footer="Page 1",
             paras=["Case opened 4 May 2026 by AU-0058.",
                    "Ms Grech reports the system rejected her transfer.",
                    "Fiona says she was born on 1 April 1979 and will send ID."],
             tables=[[["Holder", "BSB", "Account", "Contact"],
                      ["Fiona Grech", "062-000", "12345678", "fgrech@example.net / 0419 555 302"]],
                     [["Note", "Owner"], ["Crane escalation", "AU-0058"]]])

    idc = cn_id("440106", "19880720", "038")
    docx_doc("docx_zh_letter_sun", "zh",
             ["孙晓燕", "孙女士", "晓燕", idc, "13566778899", "北京市海淀区中关村大街27号5单元802室",
              "sun.xiaoyan@example.org"],
             ["客服", "AU-0112", "Crane"],
             header="致客户函件 - 孙晓燕",
             footer="Crane 客服中心",
             paras=["尊敬的孙女士：", "您好！感谢您联系Crane客服中心。",
                    "根据您的申请，晓燕女士名下账户已完成实名认证。",
                    "我们留存的身份证号为" + idc + "，手机号13566778899。",
                    "通讯地址：北京市海淀区中关村大街27号5单元802室，电子邮箱 sun.xiaoyan@example.org。",
                    "如有疑问请联系客服 AU-0112。"])


# --------------------------------------------------------------------------- EML
def addr(name, email):
    u, d = email.split("@")
    return Address(name, u, d)


def make_eml(path, frm, to, cc, subject, text, html_alt=None, date="Sat, 02 May 2026 10:14:03 +1000"):
    m = EmailMessage()
    m["From"] = addr(*frm)
    m["To"] = [addr(*t) for t in to]
    if cc:
        m["Cc"] = [addr(*c) for c in cc]
    m["Subject"] = subject
    m["Date"] = date
    m["Message-ID"] = f"<{abs(hash(subject)) % 10**8 if False else sum(map(ord, subject))}@example.com>"
    m.set_content(text)
    if html_alt:
        m.add_alternative(html_alt, subtype="html")
    Path(path).write_bytes(m.as_bytes())


def eml_doc(name, lang, pii, keep, frm, to, cc, subject, text, html_alt=None):
    plain = "\n".join([frm[0], frm[1], subject, text] + [x for t in to + (cc or []) for x in t])
    add(name, "eml", lang, pii, keep, plain,
        lambda p: make_eml(p, frm, to, cc, subject, text, html_alt))


def eml_docs():
    eml_doc("eml_en_thread_olivia", "en",
            ["Olivia Bennett", "olivia.bennett@example.com", "Harold Nguyen", "harold.nguyen@example.net",
             "Grace Liu", "grace.liu@example.com", "Olivia", "Harold", "+61 412 555 730", "Bennett",
             "Nguyen"],
            ["Crane", "team", "3 May 2026", "AU-0031"],
            ("Olivia Bennett", "olivia.bennett@example.com"),
            [("Harold Nguyen", "harold.nguyen@example.net")],
            [("Grace Liu", "grace.liu@example.com")],
            "Re: Withdrawal query - Olivia Bennett",
            "Hi Harold,\n\nThanks for the update. Can Grace be copied on the outcome? She handles our books.\n\n"
            "Regards,\nOlivia\n\n--\nOlivia Bennett | Senior Analyst\nM: +61 412 555 730\n\n"
            "> On 3 May 2026, Harold Nguyen wrote:\n> Hi Ms Bennett,\n> The Crane team has reviewed it.\n"
            "> Case handled by AU-0031.\n> Nguyen\n")

    html_alt = ("<html><body><p>Dear Dr Okafor,</p><p>Your statement is ready. Please call Samuel on "
                "(02) 9555 0143 or reply here.</p><p>Kind regards,<br><b>Priscilla Tan</b><br>"
                "Crane Client Services</p></body></html>")
    eml_doc("eml_en_html_alt_okafor", "en",
            ["Samuel Okafor", "samuel.okafor@example.org", "Priscilla Tan", "priscilla.tan@example.net",
             "Dr Okafor", "Samuel", "(02) 9555 0143", "Priscilla"],
            ["Crane", "Client Services"],
            ("Priscilla Tan", "priscilla.tan@example.net"),
            [("Samuel Okafor", "samuel.okafor@example.org")], None,
            "Your statement is ready",
            "Dear Dr Okafor,\n\nYour statement is ready. Please call Samuel on (02) 9555 0143 or reply here.\n\n"
            "Kind regards,\nPriscilla Tan\nCrane Client Services\n", html_alt)

    eml_doc("eml_en_surname_first_wu", "en",
            ["Wu, Jenny", "jenny.wu@example.org", "Jenny", "Ms Wu", "Peter Adeyemi", "peter.adeyemi@example.com",
             "Peter", "0421 555 764"],
            ["AU-0031", "system", "team"],
            ("Wu, Jenny", "jenny.wu@example.org"),
            [("Peter Adeyemi", "peter.adeyemi@example.com")],
            [("Wu, Jenny (personal)", "jenny.wu.personal@example.net")],
            "Card mismatch",
            "Hi Peter,\n\nJenny here. Please ring me on 0421 555 764 about the mismatch.\n\nThanks\nJenny\n"
            "\nOn Tue, 5 May 2026 at 6:45 pm Peter wrote:\n> Ms Wu, our system flagged it. AU-0031 is on the case.\n"
            "> The team will call.\n")

    eml_doc("eml_en_forwarded_chat_ridley", "en",
            ["Tom Ridley", "tom.ridley@example.org", "Tom", "DOB: 05/07/1991", "5 July 1991", "Ridley",
             "Rachel Zhou", "rachel.zhou@example.com", "Rachel", "MarkT"],
            ["AU-0044", "Crane", "system"],
            ("Rachel Zhou", "rachel.zhou@example.com"),
            [("Compliance", "compliance@example.org")], None,
            "Fwd: chat transcript for review",
            "Hi,\n\nPlease review the transcript below. Rachel\n\n---------- Forwarded message ---------\n"
            "From: Tom Ridley <tom.ridley@example.org>\nDate: Sat, 2 May 2026\nSubject: help\n\n"
            "[Customer]MarkT 2026-05-02 10:14:03\nHi, my name is Tom Ridley, born 5 July 1991\n"
            "[agent]AU-0044 2026-05-02 10:15:10\nThanks Tom. DOB: 05/07/1991 noted. Mr Ridley, the Crane system logged it.\n",
            None)

    idw = cn_id("510104", "19720418", "011")
    eml_doc("eml_zh_wang_jianguo", "zh",
            ["王建国", "wang.jianguo@example.com", "建国", "王先生", idw, "13711223344",
             "成都市锦江区红星路三段99号", "李婷", "li.ting@example.net"],
            ["客服", "AU-0067"],
            ("王建国", "wang.jianguo@example.com"),
            [("客服中心", "cs@example.net")], [("李婷", "li.ting@example.net")],
            "关于账户实名认证 - 王建国",
            "您好，\n\n我是王建国，想确认实名认证进度。身份证号：" + idw + "，手机13711223344。\n"
            "地址：成都市锦江区红星路三段99号。\n请抄送给我的助理李婷。\n\n此致\n建国\n\n"
            "> 王先生您好，您的申请已收到，客服 AU-0067 会跟进。\n")

    zh_html = ("<html><body><p>张女士您好：</p><p>您的账户已更新。丽华，请留意短信。</p>"
               "<p>客服中心</p></body></html>")
    eml_doc("eml_zh_html_alt_zhang", "zh",
            ["张丽华", "zhang.lihua@example.org", "张女士", "丽华", "13712340987", "+86 137 1234 0987",
             "陈志远", "chen.zhiyuan@example.com", "志远"],
            ["客服"],
            ("陈志远", "chen.zhiyuan@example.com"),
            [("张丽华", "zhang.lihua@example.org")], None,
            "账户更新通知",
            "张女士您好：\n\n您的账户已更新。丽华，请留意短信（13712340987 / +86 137 1234 0987）。\n\n"
            "客服中心 志远\n", zh_html)

    eml_doc("eml_en_signature_lastname_only", "en",
            ["Beatrice Castellano", "b.castellano@example.com", "Castellano", "Beatrice", "0392 555 017",
             "Nathan Pike", "nathan.pike@example.org", "Pike", "Nathan"],
            ["Crane", "AU-0058", "team"],
            ("Nathan Pike", "nathan.pike@example.org"),
            [("Beatrice Castellano", "b.castellano@example.com")], None,
            "Your enquiry",
            "Dear Ms Castellano,\n\nThanks for reaching out. Beatrice, I have passed your enquiry to the Crane team "
            "(ref AU-0058). Castellano is spelled correctly on your profile.\n\nCheers,\nNathan\n\n"
            "Nathan Pike\nClient Services\nP: 0392 555 017\n")


# --------------------------------------------------------------------------- HTML email
def html_email_docs():
    b = (b"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Email</title></head><body>"
         b"<table><tr><td><b>From:</b></td><td>Lucas Ferreira &lt;lucas.ferreira@example.com&gt;</td></tr>"
         b"<tr><td><b>To:</b></td><td>Client Services &lt;support@example.org&gt;</td></tr>"
         b"<tr><td><b>Subject:</b></td><td>Update my details</td></tr></table>"
         b"<p>Hello team,</p><p>My name is Lucas Ferreira. Please update my mobile to 0417 555 660 and my "
         b"address to 9 Jacaranda Drive, Toowong QLD 4066.</p><p>My old surname was Ferro so Ferreira "
         b"might be misspelt. Lucas</p><p>Ref AU-0027, effective 20 May 2026.</p></body></html>")
    add("html_email_en_lucas", "html", "en",
        ["Lucas Ferreira", "lucas.ferreira@example.com", "0417 555 660",
         "9 Jacaranda Drive, Toowong QLD 4066", "Ferreira", "Lucas"],
        ["AU-0027", "20 May 2026", "team"], b.decode(), b)

    b = ("<!DOCTYPE html><html><head><meta charset='utf-8'><title>邮件</title></head><body>"
         "<p>发件人：周文博 &lt;zhou.wenbo@example.net&gt;</p><p>收件人：客服中心</p>"
         "<p>您好，我是周文博。请把我的手机号改成15012345678。</p>"
         "<p>周先生的账户号是 55023819。文博 谢谢！</p><p>出生日期：1990年12月1日</p></body></html>").encode()
    add("html_email_zh_zhou", "html", "zh",
        ["周文博", "zhou.wenbo@example.net", "15012345678", "周先生", "55023819", "文博", "1990年12月1日"],
        ["客服"], b.decode(), b)


# --------------------------------------------------------------------------- main
EXT = {"pdf": "pdf", "html": "html", "docx": "docx", "eml": "eml"}


def main():
    if CORPUS.exists():
        shutil.rmtree(CORPUS)
    CORPUS.mkdir(parents=True)
    DOCS.clear()
    pdf_chats()
    html_chats()
    docx_docs()
    eml_docs()
    html_email_docs()
    labels = {}
    for d in DOCS:
        fn = f"{d['name']}.{EXT[d['fmt']]}"
        p = CORPUS / fn
        if callable(d["payload"]):
            d["payload"](p)
        else:
            p.write_bytes(d["payload"])
        labels[fn] = {"format": d["fmt"], "lang": d["lang"], "pii": d["pii"], "must_keep": d["keep"]}
    LABELS.write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {len(labels)} documents to {CORPUS}", file=sys.stderr)


if __name__ == "__main__":
    main()
