#!/usr/bin/env python
"""Deterministic blind-corpus generator. All people, numbers and addresses are invented."""
import base64
import json
import re
from email.header import Header
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: F401
import datetime

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "corpus"
OUT.mkdir(parents=True, exist_ok=True)
for old in OUT.iterdir():
    old.unlink()

LABELS = {}


def reg(name, fmt, lang, text, pii, keep):
    for s in pii + keep:
        assert s in text, f"{name}: label not in source: {s!r}"
    assert len(set(pii)) == len(pii), f"{name}: duplicate pii"
    assert not (set(pii) & set(keep)), f"{name}: pii/keep overlap"
    LABELS[name] = {"format": fmt, "lang": lang, "pii": pii, "must_keep": keep}


# ---------------- identifier helpers ----------------
def cn_id(area, birth, seq):
    b = area + birth + seq
    assert len(b) == 17
    w = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    return b + "10X98765432"[sum(int(c) * k for c, k in zip(b, w)) % 11]


def hk_id(letter, digits):
    s = 36 * 9 + (ord(letter) - 55) * 8 + sum(int(d) * (7 - i) for i, d in enumerate(digits))
    c = (11 - s % 11) % 11
    return f"{letter}{digits}({'A' if c == 10 else c})"


def luhn(prefix, n=16):
    body = prefix
    while len(body) < n - 1:
        body += str((len(body) * 7 + 3) % 10)
    tot = 0
    for i, ch in enumerate(reversed(body)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            d = d - 9 if d > 9 else d
        tot += d
    return body + str((10 - tot % 10) % 10)


def grp4(s):
    return " ".join(s[i:i + 4] for i in range(0, len(s), 4))


# ---------------- format writers ----------------
def write_pdf(name, lines, per_page=70):
    for ln in lines:
        ln.encode("cp1252")
        assert len(ln) <= 200, f"{name}: line too long: {ln!r}"
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)]
    objs = []

    def esc(s):
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    n = len(pages)
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for i, pg in enumerate(pages):
        cs = ["BT", "/F1 7.5 Tf", "11 TL", "36 805 Td"]
        for ln in pg:
            cs.append(f"({esc(ln)}) Tj T*")
        cs.append("ET")
        stream = "\n".join(cs).encode("cp1252")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> "
            f"/Contents {5 + 2 * i} 0 R >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offs = []
    for i, o in enumerate(objs):
        offs.append(len(out))
        out += f"{i + 1} 0 obj\n".encode() + o + b"\nendobj\n"
    xr = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xr}\n%%EOF\n".encode()
    (OUT / name).write_bytes(bytes(out))
    return "\n".join(lines)


def build_docx(blocks, header=None, footer=None):
    """blocks: ('h', text) | ('p', text) | ('t', rows). Returns (bytes, text)."""
    d = Document()
    d.core_properties.created = datetime.datetime(2026, 8, 1)
    d.core_properties.modified = datetime.datetime(2026, 8, 1)
    d.core_properties.author = "Crane Brokers Operations"
    texts = []
    if header:
        d.sections[0].header.paragraphs[0].text = header
        texts.append(header)
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
        texts.append(footer)
    for kind, val in blocks:
        if kind == "h":
            d.add_heading(val, level=1)
            texts.append(val)
        elif kind == "p":
            d.add_paragraph(val)
            texts.append(val)
        else:
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
                    texts.append(cell)
    import io
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue(), "\n".join(texts)


def write_docx(name, blocks, header=None, footer=None):
    data, text = build_docx(blocks, header, footer)
    (OUT / name).write_bytes(data)
    return text


def write_html(name, html):
    (OUT / name).write_text(html, encoding="utf-8")
    return html


def enc_hdr(k, v):
    if v.isascii():
        return v
    if k in ("From", "To", "Cc", "Reply-To"):
        parts = []
        for p in v.split(", "):
            m = re.match(r"(.*) <(.*)>$", p)
            if m and not m.group(1).isascii():
                parts.append(f"{Header(m.group(1), 'utf-8').encode()} <{m.group(2)}>")
            else:
                parts.append(p)
        return ", ".join(parts)
    return Header(v, "utf-8").encode()


def write_eml(name, headers, body, attachments=(), b64_body=False):
    """attachments: (filename, mimetype, bytes, logical_text). Returns logical text."""
    logical = [f"{k}: {v}" for k, v in headers]
    hl = [f"{k}: {enc_hdr(k, v)}" for k, v in headers]
    hl.append("MIME-Version: 1.0")
    if b64_body:
        benc = base64.encodebytes(body.encode("utf-8")).decode().replace("\n", "\r\n" if False else "\n")
        cte, bpart = "base64", benc
    else:
        cte, bpart = "8bit", body
    if not attachments:
        hl += ["Content-Type: text/plain; charset=utf-8", f"Content-Transfer-Encoding: {cte}", "", bpart]
    else:
        bd = "----=_crane_blind_" + re.sub(r"\W", "", name)
        hl += [f'Content-Type: multipart/mixed; boundary="{bd}"', "", "This is a multi-part message in MIME format.",
               f"--{bd}", "Content-Type: text/plain; charset=utf-8", f"Content-Transfer-Encoding: {cte}", "", bpart]
        for fn, mt, data, _ in attachments:
            b = base64.encodebytes(data).decode()
            hl += [f"--{bd}", f'Content-Type: {mt}; name="{fn}"', "Content-Transfer-Encoding: base64",
                   f'Content-Disposition: attachment; filename="{fn}"', "", b]
        hl.append(f"--{bd}--")
    (OUT / name).write_bytes(("\n".join(hl) + "\n").encode("utf-8"))
    return "\n".join(logical) + "\n\n" + body + "\n" + "\n".join(a[3] for a in attachments)


def page(title, body, lang="en", extra_head=""):
    return (f'<!DOCTYPE html>\n<html lang="{lang}">\n<head>\n<meta charset="utf-8">\n<title>{title}</title>\n'
            f'{extra_head}</head>\n<body>\n{body}\n</body>\n</html>\n')


# =====================================================================
# PDFs (English only)
# =====================================================================
acct1 = "CB-20458817"
t = write_pdf("doc01_chat_en.pdf", [
    "Crane Brokers - Live Chat Export",
    "Session 88213041   Queue: Withdrawals   Exported by AU-0117",
    "[System] Customer joined the chat 2026-08-03 14:20:41",
    "[Bot]CraneAssist 2026-08-03 14:20:42 Hi! I can help with deposits, withdrawals and orders. Type your issue.",
    "[Customer]<kenji_trades88> 2026-08-03 14:21:05 my withdrawal is stuck since monday!!",
    "[agent]AU-0210 2026-08-03 14:21:30 Hello, thanks for contacting Crane. Can I confirm your full name and DOB?",
    "[Customer]<kenji_trades88> 2026-08-03 14:22:10 Kenji Watanabe-Smith, 14/02/1987 and account CB-20458817",
    "[agent]AU-0210 2026-08-03 14:22:41 Thanks Kenji. I can see withdrawal WD-55127 for AUD 4,250.00 is pending.[Customer]<kenji_trades88> 2026-08-03 14:23:02 its been 3 days",
    "[agent]AU-0210 2026-08-03 14:23:40 The bank rejected it. Is the BSB 062-141 and account 4471 8820 still right?",
    "[Customer]<kenji_trades88> 2026-08-03 14:24:15 no i changed banks. new one is BSB 733-021 acc 10099347",
    "[agent]AU-0210 2026-08-03 14:24:50 Understood. Please email the form to us from kenji.ws@example.com, ticket TKT-774102.",
    "[Customer]<kenji_trades88> 2026-08-03 14:25:20 will do. call me on 0412 908 331 if theres a problem[System] Agent AU-0210 sent a file",
    "[Bot]CraneAssist 2026-08-03 14:25:21 Was this chat helpful? Reply 1-5.",
    "[Customer]<kenji_trades88> 2026-08-03 14:25:33 4. also sold 50 AAPL at 187.42 on 31 July, is that settled?",
    "[agent]AU-0210 2026-08-03 14:26:01 Yes, T+1 settlement completed 2026-08-03 with Crane Trader fees of USD 0.99.",
    "[System] Chat ended 2026-08-03 14:26:30",
])
reg("doc01_chat_en.pdf", "pdf", "en", t,
    ["kenji_trades88", "Kenji Watanabe-Smith", "Kenji", "14/02/1987", acct1, "062-141", "4471 8820", "733-021",
     "10099347", "kenji.ws@example.com", "0412 908 331"],
    ["Crane Brokers", "AU-0210", "AU-0117", "WD-55127", "AUD 4,250.00", "TKT-774102", "AAPL", "187.42", "50 AAPL",
     "Crane Trader", "2026-08-03 14:22:10", "CraneAssist", "T+1", "USD 0.99"])

t = write_pdf("doc02_chat_timeonly_en.pdf", [
    "Crane - Support Chat Transcript (time-only format)",
    "Date: 12 Aug 2026  Agent: AU-0342",
    "[10:02:11] AU-0342: Welcome to Crane Brokers support. Who am I speaking with?",
    "[10:02:40] Customer: this is Dimitrios Papadakis but everyone calls me Jim",
    "[10:03:05] AU-0342: Hi Mr Papadakis, could you verify your passport number and date of birth?",
    "[10:03:37] Customer: N4482019 (AU passport) born 9 November 1961",
    "[10:04:02] AU-0342: Thanks Jim. What's the issue today?",
    "[10:04:31] Customer: I wanted to buy 200 BHP but the order got rejected at 09:58, limit 44.10",
    "[10:05:00] AU-0342: The order was rejected because of insufficient settled funds, order ref ORD-903311.",
    "[10:05:20] Customer: ok my wife Eleni Papadaki has a joint account, can she top it up? her mobile is +61 3 9555 0142",
    "[10:05:55] AU-0342: She can deposit via PayID. Please do not share card details in chat.",
    "[10:06:10] Customer: fine. My licence is VIC 0387 5521 if you need id. Email jim.papadakis@example.net",
    "[10:06:44] AU-0342: Noted. Anything else, Jim?[10:06:58] Customer: no thanks, cheers",
    "[10:07:02] System: chat closed by agent",
])
reg("doc02_chat_timeonly_en.pdf", "pdf", "en", t,
    ["Dimitrios Papadakis", "Jim", "Mr Papadakis", "N4482019", "9 November 1961", "Eleni Papadaki",
     "+61 3 9555 0142", "0387 5521", "jim.papadakis@example.net"],
    ["Crane Brokers", "AU-0342", "200 BHP", "44.10", "09:58", "ORD-903311", "PayID", "10:02:11"])

card = grp4(luhn("453201"))
t = write_pdf("doc03_statement_letter_en.pdf", [
    "CRANE BROKERS (AUSTRALIA) - CLIENT ACCOUNT SUMMARY",
    "Statement period: 1 July 2026 to 31 July 2026",
    "",
    "Ms Sarah-Jane Whitcombe",
    "Unit 4/17 Marrickville Road",
    "Marrickville NSW 2204",
    "",
    "Dear SJ,",
    "",
    "Your account number 7730-1194-SJW is in good standing. Below is a summary of activity.",
    "",
    "Date        Type   Ticker  Qty   Price     Amount",
    "2026-07-03  BUY    NVDA    10    121.80    1,218.00",
    "2026-07-11  SELL   TSLA    5     248.15    1,240.75",
    "2026-07-21  BUY    CBA     30    152.60    4,578.00",
    "",
    "Registered TFN on file: 812 345 677 (ending 677). ABN for your trust: 51 824 753 556.",
    "Linked bank account: BSB 012-003 account 3391 0287. Card used for fees: " + card + ".",
    "If any detail is wrong please reply to sj.whitcombe@example.org or ring 02 9550 7718.",
    "Date of birth held: 22-Jan-1974. Your adviser is Marcus Oyelaran-Hughes.",
    "",
    "Yours sincerely,",
    "Client Services, Crane Brokers",
    "Ref: STM-2026-07-88120",
])
reg("doc03_statement_letter_en.pdf", "pdf", "en", t,
    ["Sarah-Jane Whitcombe", "Ms Sarah-Jane Whitcombe", "Unit 4/17 Marrickville Road", "Marrickville NSW 2204", "SJ",
     "7730-1194-SJW", "812 345 677", "51 824 753 556", "012-003", "3391 0287", card,
     "sj.whitcombe@example.org", "02 9550 7718", "22-Jan-1974", "Marcus Oyelaran-Hughes"],
    ["CRANE BROKERS (AUSTRALIA)", "NVDA", "TSLA", "CBA", "121.80", "248.15", "152.60", "2026-07-03",
     "STM-2026-07-88120", "Client Services"])

t = write_pdf("doc04_complaint_letter_en.pdf", [
    "Mr Rajesh Venkataraman",
    "PO Box 1187",
    "Parramatta NSW 2124",
    "18 August 2026",
    "",
    "The Complaints Officer",
    "Crane Brokers",
    "",
    "RE: Formal complaint, ticket TKT-781266",
    "",
    "To whom it may concern,",
    "",
    "I write to complain about the handling of my account. On 5 August I placed a market order to sell",
    "120 shares of WOW at about 34.20 and it filled at 33.05. I was not warned about the spread.",
    "My client number is CB-30981245 and my Medicare card is 2953 66481 2.",
    "Please note R. Venkataraman is also how my super fund shows me, and my wife Priya calls the",
    "support line on 0433 216 090 from time to time. Do not discuss my account with her without consent.",
    "My date of birth is 04/07/1969 and my passport is PA3319075.",
    "",
    "I expect a reply within 30 days as per RG 271. Contact rajesh.v@example.com.",
    "",
    "Yours faithfully,",
    "Rajesh Venkataraman",
    "Raj",
])
reg("doc04_complaint_letter_en.pdf", "pdf", "en", t,
    ["Mr Rajesh Venkataraman", "Rajesh Venkataraman", "PO Box 1187", "Parramatta NSW 2124", "CB-30981245",
     "2953 66481 2", "R. Venkataraman", "Priya", "0433 216 090", "04/07/1969", "PA3319075",
     "rajesh.v@example.com", "Raj"],
    ["Crane Brokers", "TKT-781266", "120 shares of WOW", "34.20", "33.05", "RG 271", "Complaints Officer"])

t = write_pdf("doc05_aml_review_note_en.pdf", [
    "CRANE BROKERS - COMPLIANCE CASE NOTE (INTERNAL)",
    "Case: AML-2026-0419   Reviewer: CO-07   Opened 2026-08-09",
    "",
    "Subject: Fatima Al-Mansour (also appears as F. Almansour and Fatema Al Mansoor in older records)",
    "Occupation stated: dental practice owner.  Residential: 12 Wattle Grove Rd, Kangaroo Valley NSW 2577",
    "",
    "Trigger: three inbound transfers from a third party totalling AUD 48,500 between 1 and 6 August.",
    "Remitter name on the bank feed: Yusuf Al-Mansour (brother).  Reference text: 'family loan'.",
    "Customer contacted via WhatsApp +971 50 123 4567 and by mobile 0455 700 123.",
    "Customer says: 'call me Fati, I only use fatima_dds on Telegram'.",
    "ID sighted: Australian passport N9027613, driver licence 12 998 220 (SA), DOB 30 June 1980.",
    "Bank account for withdrawals: IBAN GB51 EXMP 4020 1177 3305 62.",
    "",
    "Assessment: source of funds plausible. No SMR lodged. Enhanced monitoring for 90 days.",
    "Products held: Crane Trader margin, US equities, ASX options.",
    "Next review CO-07 on 2026-11-09.",
])
reg("doc05_aml_review_note_en.pdf", "pdf", "en", t,
    ["Fatima Al-Mansour", "F. Almansour", "Fatema Al Mansoor", "12 Wattle Grove Rd, Kangaroo Valley NSW 2577",
     "Yusuf Al-Mansour", "+971 50 123 4567", "0455 700 123", "Fati", "fatima_dds", "N9027613", "12 998 220",
     "30 June 1980", "GB51 EXMP 4020 1177 3305 62"],
    ["CRANE BROKERS", "AML-2026-0419", "CO-07", "AUD 48,500", "Crane Trader", "ASX options", "2026-08-09",
     "Enhanced monitoring"])

t = write_pdf("doc06_chat_system_speakers_en.pdf", [
    "Crane Brokers Support - Chat Log 55120977",
    "[System] Visitor 1 connected (web, AU) 2026-09-01 08:14:03",
    "[Bot]CraneAssist 2026-09-01 08:14:04 You are chatting with an automated assistant.",
    "[Customer]<@aroha.terangi> 2026-09-01 08:14:40 kia ora, can I transfer my NZ shares in?",
    "[Bot]CraneAssist 2026-09-01 08:14:41 Transfers in are handled by an agent. Connecting you now.",
    "[System] Agent AU-0451 joined 2026-09-01 08:15:10",
    "[agent]AU-0451 2026-09-01 08:15:30 Hi, you're through to Crane. May I have your name?",
    "[Customer]<@aroha.terangi> 2026-09-01 08:15:52 Aroha Te Rangi. Aroha is fine",
    "[agent]AU-0451 2026-09-01 08:16:20 Thanks Aroha. Which broker holds the shares?",
    "[Customer]<@aroha.terangi> 2026-09-01 08:16:55 Sharesies acct, holder no. 1234567 CSN A5521098",
    "[Customer]<@aroha.terangi> 2026-09-01 08:17:20 my number is +64 21 555 0199 and IRD 049-091-850",
    "[agent]AU-0451 2026-09-01 08:17:48 Great. Form TRF-2 is attached to ticket TKT-790015.[System] Attachment sent",
    "[Customer]<@aroha.terangi> 2026-09-01 08:18:30 sweet. I live at 14 Tui Street Ponsonby Auckland 1011",
    "[System] Chat closed 2026-09-01 08:19:00",
])
reg("doc06_chat_system_speakers_en.pdf", "pdf", "en", t,
    ["@aroha.terangi", "Aroha Te Rangi", "Aroha", "A5521098", "1234567", "+64 21 555 0199", "049-091-850",
     "14 Tui Street Ponsonby Auckland 1011"],
    ["Crane Brokers", "AU-0451", "TKT-790015", "TRF-2", "CraneAssist", "Sharesies", "2026-09-01 08:14:40"])

# =====================================================================
# HTML
# =====================================================================
cnid1 = cn_id("310104", "19850316", "021")
h = page("Web form submission", f"""<h1>Crane Brokers - Account Enquiry Form</h1>
<p class="meta">Received 2026-08-14 09:41 via cranebrokers web form. Ref: ENQ-66120.</p>
<table>
<tr><th>Field</th><th>Value</th></tr>
<tr><td>Full name</td><td>Tran Minh Quan</td></tr>
<tr><td>Preferred name</td><td>Quan</td></tr>
<tr><td>Email</td><td><a href="mailto:quan.tran@example.com">quan.tran@example.com</a></td></tr>
<tr><td>Mobile</td><td>(04) 0871 2255</td></tr>
<tr><td>DOB</td><td>March 3rd, 1979</td></tr>
<tr><td>Address</td><td>Lot 12, Old Coach Road, Yandina QLD 4561</td></tr>
<tr><td>ABN</td><td>73 119 004 852</td></tr>
<tr><td>Enquiry type</td><td>Opening a company account</td></tr>
</table>
<p>Message: Hi, my accountant Ms Bui Thu Ha (thuha@example.net) told me to ask about a Crane Trader
company account for Quan Tran Family Trust. I hold 500 shares of ANZ elsewhere.</p>
<p><small>Do not reply to this address. Case handled by AU-0288.</small></p>""")
t = write_html("doc07_webform_en.html", h)
reg("doc07_webform_en.html", "html", "en", t,
    ["Tran Minh Quan", "Quan", "quan.tran@example.com", "(04) 0871 2255", "March 3rd, 1979",
     "Lot 12, Old Coach Road, Yandina QLD 4561", "73 119 004 852", "Ms Bui Thu Ha", "thuha@example.net",
     "Quan Tran Family Trust"],
    ["Crane Brokers", "ENQ-66120", "Crane Trader", "500 shares of ANZ", "AU-0288", "company account"])

h = page("CRM notes", """<div class="crm">
<h2>CRM Contact Notes - Customer 20913377</h2>
<ul>
<li><b>2026-07-02 11:05 AU-0199</b>: Customer Vanessa Lo-Brennan (VLB) called re dividend on VAS. Verified with DOB 1990-05-27
and postcode 3121.</li>
<li><b>2026-07-09 15:30 AU-0230</b>: VLB emailed from v.lobrennan@example.org; says husband Declan Brennan may ring - he is NOT authorised.</li>
<li><b>2026-07-16 09:12 AU-0199</b>: Mobile updated 0488 341 906 (was 0419 200 774). Also gave Line ID vanlo_88 and QQ 88213640.</li>
<li><b>2026-07-30 16:45 AU-0141</b>: Ms Lo-Brennan disputes fee of AUD 9.90 on order ORD-11902 (sell 75 VAS @ 98.34).
Escalated to team lead, ticket TKT-770233.</li>
</ul>
<p>Address on file: 3/88 Bridge Road, Richmond VIC 3121. Driver licence: 0945 1237 (VIC). Tax file no. 456 789 123.</p>
</div>""")
t = write_html("doc08_crm_notes_en.html", h)
reg("doc08_crm_notes_en.html", "html", "en", t,
    ["Vanessa Lo-Brennan", "VLB", "1990-05-27", "v.lobrennan@example.org", "Declan Brennan", "0488 341 906",
     "0419 200 774", "vanlo_88", "88213640", "Ms Lo-Brennan", "3/88 Bridge Road, Richmond VIC 3121", "0945 1237",
     "456 789 123", "20913377"],
    ["AU-0199", "AU-0230", "AU-0141", "VAS", "AUD 9.90", "ORD-11902", "75 VAS @ 98.34", "TKT-770233",
     "2026-07-02 11:05", "dividend"])

h = page("CRM备注", f"""<h2>客户跟进记录 (CRM)</h2>
<p>客户：陈雅婷（英文名 Tina Chen，拼音 Chen Yating）</p>
<p>2026-08-05 10:12 AU-0355 电话回访。陈女士反映入金未到账，金额 AUD 12,000。她的微信号 tina_cyt2020，手机 +86 138 0013 8000。</p>
<p>2026-08-05 10:20 AU-0355 核对身份证号 {cn_id("440305", "19920718", "004")}，护照 E12345678（中国护照），出生日期 1992年7月18日。</p>
<p>2026-08-06 09:00 AU-0102 银行转账凭证显示付款人为陈雅婷，卡号 {grp4(luhn("622202", 19))[:0]}6222 0245 0100 3345 678，开户行工商银行深圳南山支行。</p>
<p>通信地址：广东省深圳市南山区科技园南路18号阳光大厦12楼1208室，邮编 518057。</p>
<p>备注：Tina 称丈夫林志远也会联系我们，需要客户书面授权。工单号 TKT-795501，产品：白鹤证券 Crane Trader 美股账户。</p>""", "zh")
cn_a = cn_id("440305", "19920718", "004")
t = write_html("doc09_crm_notes_zh.html", h)
reg("doc09_crm_notes_zh.html", "html", "zh", t,
    ["陈雅婷", "Tina Chen", "Chen Yating", "陈女士", "tina_cyt2020", "+86 138 0013 8000", cn_a, "E12345678",
     "1992年7月18日", "6222 0245 0100 3345 678", "广东省深圳市南山区科技园南路18号阳光大厦12楼1208室", "518057",
     "Tina", "林志远"],
    ["AU-0355", "AU-0102", "AUD 12,000", "TKT-795501", "白鹤证券", "Crane Trader", "美股账户", "入金"])

h = page("Chat viewer export", """<div class="chat">
<div class="msg"><span class="hdr">[Customer]&lt;wangjg_1968&gt; 2026-08-20 20:31:05</span> 你好，我想查一下我的美股账户，上周买的 100 股 AAPL 怎么没显示</div>
<div class="msg"><span class="hdr">[agent]AU-0266 2026-08-20 20:31:40</span> 您好，欢迎联系白鹤证券。请问怎么称呼您？</div>
<div class="msg"><span class="hdr">[Customer]&lt;wangjg_1968&gt; 2026-08-20 20:32:12</span> 我姓王，王建国，英文名 Jack Wang。Please speak English if easier</div>
<div class="msg"><span class="hdr">[agent]AU-0266 2026-08-20 20:32:50</span> 好的王先生。请提供您的账户号和注册手机号后四位。</div>
<div class="msg"><span class="hdr">[Customer]&lt;wangjg_1968&gt; 2026-08-20 20:33:30</span> 账户 CB-88140276，手机 0421 667 305。my WeChat is jackwang_syd and WhatsApp same number</div>
<div class="msg"><span class="hdr">[Bot]CraneAssist 2026-08-20 20:33:31</span> 提示：请勿在聊天中发送银行卡完整号码。</div>
<div class="msg"><span class="hdr">[agent]AU-0266 2026-08-20 20:34:10</span> 老王，您的 AAPL 订单 ORD-441209 已于 2026-08-19 22:47 成交，价格 USD 201.55，T+1 结算。</div>
<div class="msg"><span class="hdr">[Customer]&lt;wangjg_1968&gt; 2026-08-20 20:34:55</span> 谢谢。我太太Mei Wang的邮箱是 mei.wang@example.com，以后也发给她一份</div>
</div>""", "mixed")
t = write_html("doc10_chat_mixed.html", h)
t = t.replace("&lt;", "<").replace("&gt;", ">")  # labels are checked against decoded text
reg("doc10_chat_mixed.html", "html", "mixed", t,
    ["wangjg_1968", "王建国", "Jack Wang", "王先生", "老王", "CB-88140276", "0421 667 305", "jackwang_syd",
     "Mei Wang", "mei.wang@example.com", "我姓王"],
    ["AU-0266", "CraneAssist", "100 股 AAPL", "白鹤证券", "ORD-441209", "USD 201.55", "T+1", "2026-08-20 20:31:05"])
LABELS["doc10_chat_mixed.html"]["pii"].remove("我姓王")
LABELS["doc10_chat_mixed.html"]["pii"].append("wangjg_1968")  if False else None

h = page("客户表单提交", """<form>
<h1>开户咨询表 - 白鹤证券</h1>
<dl>
<dt>姓名</dt><dd>张伟</dd>
<dt>英文名</dt><dd>Wei Zhang (David)</dd>
<dt>手机</dt><dd>139-1010-6677</dd>
<dt>微信</dt><dd>wxid_zw_2231</dd>
<dt>邮箱</dt><dd>zhangwei.sz@example.net</dd>
<dt>出生日期</dt><dd>1978/11/02</dd>
<dt>护照号码</dt><dd>G28801932</dd>
<dt>住址</dt><dd>上海市浦东新区世纪大道100号环球金融中心45层</dd>
<dt>咨询产品</dt><dd>港股通、美股期权</dd>
</dl>
<p>补充说明：我之前在别的券商持有 2000 股 700.HK。请张律师（Lawyer Cheung Wing-kit, wk.cheung@example.org）联系我。提交时间 2026-09-02 13:05。</p>
</form>""", "zh")
t = write_html("doc11_webform_zh.html", h)
reg("doc11_webform_zh.html", "html", "zh", t,
    ["张伟", "Wei Zhang", "David", "139-1010-6677", "wxid_zw_2231", "zhangwei.sz@example.net", "1978/11/02",
     "G28801932", "上海市浦东新区世纪大道100号环球金融中心45层", "张律师", "Cheung Wing-kit", "wk.cheung@example.org"],
    ["白鹤证券", "港股通", "美股期权", "700.HK", "2000 股", "2026-09-02 13:05", "开户咨询表"])

h = page("Support ticket export", """<article id="tkt-802214">
<h2>Ticket TKT-802214 - Withdrawal delay (HK bank)</h2>
<p>Opened 2026-09-10 by AU-0387. Priority: high. Status: pending customer.</p>
<section>
<h3>Requester</h3>
<p>Chan Siu Ming (Simon), 陈小明 - HKID <span>K123456(7)</span>; alt ID <span>__HKID__</span></p>
<p>Phone +852 9123 4567 / +852 2810 3344. Email simon.chan@example.com.</p>
<p>Flat 5B, Block 3, Kwun Tong Garden, 88 Hoi Yuen Road, Kwun Tong, Kowloon, Hong Kong</p>
<p>Passport (HK SAR): KJ0294517. DOB 8 Feb 1983.</p>
</section>
<section>
<h3>Thread</h3>
<blockquote>Mr Chan says wire of HKD 150,000 sent 6 Sept to Crane AU account has not shown. Bank ref HSBC-FT-99120.
Called Simon on WhatsApp; he asked us to also update his wife Winnie Lam's contact: winnie.lam@example.net.</blockquote>
<p>Action: check with treasury team. Do not disclose to Ms Lam.</p>
</section></article>""".replace("__HKID__", hk_id("Y", "812345")), "mixed")
hk1 = hk_id("Y", "812345")
t = write_html("doc12_ticket_export_hk.html", h)
reg("doc12_ticket_export_hk.html", "html", "mixed", t,
    ["Chan Siu Ming", "Simon", "陈小明", "K123456(7)", hk1, "+852 9123 4567", "+852 2810 3344",
     "simon.chan@example.com", "Flat 5B, Block 3, Kwun Tong Garden, 88 Hoi Yuen Road, Kwun Tong, Kowloon, Hong Kong",
     "KJ0294517", "8 Feb 1983", "Mr Chan", "Winnie Lam's", "winnie.lam@example.net", "Ms Lam"],
    ["TKT-802214", "AU-0387", "HKD 150,000", "HSBC-FT-99120", "Crane AU", "treasury team", "Withdrawal delay"])
# "Winnie Lam's" is a surface form present verbatim; also add bare name
LABELS["doc12_ticket_export_hk.html"]["pii"].append("Winnie Lam")

h = page("CRM notes - estate", """<h2>Account 61120983 - notes</h2>
<p><i>2026-08-24 AU-0222</i> Caller identified as Ngozi Adeyemi-Clarke, daughter of the account holder, the late Emmanuel Adeyemi
(DOB 12/12/1948, died 2026-07-30). Caller mobile 0466 902 118, email ngozi.ac@example.com.</p>
<p><i>2026-08-24 AU-0222</i> Advised that a certified death certificate and grant of probate are required. Solicitor:
Bernadette Kowalczyk of Kowalczyk &amp; Reid Legal, 9 Queen St, Brisbane City QLD 4000, tel (07) 3555 0100.</p>
<p><i>2026-08-27 AU-0311</i> Account restricted. Holdings: 300 shares FMG, 1,200 shares of TLS. No trades permitted.
Ms Adeyemi-Clarke prefers to be called Ngozi. Ticket TKT-806651.</p>""")
t = write_html("doc13_crm_deceased_en.html", h).replace("&amp;", "&")
reg("doc13_crm_deceased_en.html", "html", "en", t,
    ["Ngozi Adeyemi-Clarke", "Emmanuel Adeyemi", "12/12/1948", "0466 902 118", "ngozi.ac@example.com",
     "Bernadette Kowalczyk", "Kowalczyk & Reid Legal", "9 Queen St, Brisbane City QLD 4000", "(07) 3555 0100",
     "Ms Adeyemi-Clarke", "Ngozi", "61120983"],
    ["AU-0222", "AU-0311", "300 shares FMG", "1,200 shares of TLS", "TKT-806651", "grant of probate", "Account restricted"])

h = page("Web form - complaint", """<h1>Feedback and complaints</h1>
<form><p><label>Name</label> <input value="Grace Lamb"></p>
<p>Name: Grace Lamb<br>Phone: 1300 000 000 (business)<br>Mobile: +61 (0)412 555 321<br>Email: grace.lamb@example.org</p>
<p>Comments: Will Price at your Sydney desk told me on 2026-09-05 that my transfer of 40 units of VGS would be complete
in a week. Mark Jones on the phone said something different. I'm May Lim's neighbour and she suggested Crane, so I do not
appreciate this. My address: 22 Hope St, Bull Creek WA 6149. TFN 123 456 782. Date of birth: 1/9/85.</p></form>""")
t = write_html("doc14_webform_wordnames_en.html", h)
reg("doc14_webform_wordnames_en.html", "html", "en", t,
    ["Grace Lamb", "grace.lamb@example.org", "+61 (0)412 555 321", "Will Price", "Mark Jones", "May Lim",
     "22 Hope St, Bull Creek WA 6149", "123 456 782", "1/9/85"],
    ["Crane", "40 units of VGS", "2026-09-05", "1300 000 000", "Feedback and complaints", "Sydney desk"])

# =====================================================================
# EML
# =====================================================================
quoted_thread = """On Tue, 4 Aug 2026 at 16:02, Crane Brokers Support <support@example.com> wrote:
> Dear Mr Kowalski,
>
> Thank you for contacting Crane. Your ticket number is TKT-771908.
> We have reviewed order ORD-70021 (BUY 300 SOL at 8.72) and confirm it executed at 10:14:33.
>
> On 3 Aug 2026, Tomasz Kowalski <tomasz.kowalski@example.com> wrote:
> > I did not authorise that trade. My account is CB-52007731 and my number is 0402 118 664.
> > Regards, Tom
> > --
> > Tomasz "Tom" Kowalski | 7 Banksia Ct, Toowong QLD 4066 | DOB 21.03.1972
"""
body = f"""Hi team,

Escalating this to compliance. Mr Kowalski has now rung three times and his sister Agnieszka Nowak-Kowalska
(agnieszka.nk@example.net, cc'd) says she will go to AFCA. Please treat as a formal dispute.

Customer's Medicare 4123 45670 1, licence QLD 06731254, TFN 987 654 321.

{quoted_thread}
Thanks,
AU-0288
Customer Operations, Crane Brokers
"""
t = write_eml("doc15_escalation_quoted_en.eml", [
    ("From", "AU-0288 <ops.au0288@example.com>"), ("To", "Compliance <compliance@example.com>"),
    ("Cc", "Agnieszka Nowak-Kowalska <agnieszka.nk@example.net>"),
    ("Date", "Wed, 05 Aug 2026 09:12:00 +1000"), ("Message-ID", "<esc-771908@example.com>"),
    ("Subject", "FW: RE: Unauthorised trade complaint TKT-771908")], body)
reg("doc15_escalation_quoted_en.eml", "eml", "en", t,
    ["Mr Kowalski", "Agnieszka Nowak-Kowalska", "agnieszka.nk@example.net", "4123 45670 1", "06731254", "987 654 321",
     "Dear Mr Kowalski", "Tomasz Kowalski", "tomasz.kowalski@example.com", "CB-52007731", "0402 118 664", "Tom",
     "Tomasz \"Tom\" Kowalski", "7 Banksia Ct, Toowong QLD 4066", "21.03.1972"],
    ["TKT-771908", "ORD-70021", "300 SOL", "8.72", "10:14:33", "AU-0288", "Crane Brokers", "AFCA", "Compliance"])
LABELS["doc15_escalation_quoted_en.eml"]["pii"].remove("Dear Mr Kowalski")

# forwarded chat
body = """Hi Compliance,

Forwarding a chat the customer sent us as a screenshot transcript. Note the account holder is not the person typing.

---------- Forwarded message ---------
[Customer]<lily_zhou_sydney> 2026-08-11 19:02:11 hi, my father 周国强 can't log in. he's 71
[agent]AU-0120 2026-08-11 19:02:50 Hello, we can only speak to the account holder.
[Customer]<lily_zhou_sydney> 2026-08-11 19:03:20 I'm his daughter Lily Zhou (周丽丽). he is with me now. Zhou Guoqiang, DOB 1955-04-09
[agent]AU-0120 2026-08-11 19:04:05 Please have Mr Zhou call 1800 555 019 with his account CB-40219966.
[Customer]<lily_zhou_sydney> 2026-08-11 19:04:31 ok. his phone is 0433 771 208 and email g.zhou@example.net

Regards,
AU-0120
"""
t = write_eml("doc16_forwarded_chat_mixed.eml", [
    ("From", "AU-0120 <au0120@example.com>"), ("To", "Compliance <compliance@example.com>"),
    ("Date", "Tue, 11 Aug 2026 19:20:00 +1000"), ("Subject", "Fwd: login help - third party caller")], body)
reg("doc16_forwarded_chat_mixed.eml", "eml", "mixed", t,
    ["lily_zhou_sydney", "周国强", "Lily Zhou", "周丽丽", "Zhou Guoqiang", "1955-04-09", "Mr Zhou", "CB-40219966",
     "0433 771 208", "g.zhou@example.net"],
    ["AU-0120", "1800 555 019", "2026-08-11 19:02:11", "Compliance"])

# email with docx attachment (KYC) 
cnid2 = cn_id("510107", "19880225", "013")
att_blocks = [("h", "Certified ID summary"),
              ("t", [["Name", "Luo Xiaoyu (罗晓雨)"], ["ID number", cnid2], ["Passport", "EA9922017"],
                     ["Address", "Room 1502, Building 7, Jinjiang District, Chengdu, Sichuan 610021"]]),
              ("p", "Certified by JP 145522 on 1 September 2026. Crane account application AP-90211.")]
adata, atext = build_docx(att_blocks)
body = f"""Hello Crane Onboarding,

Please find the certified ID summary for my client Ms Luo Xiaoyu, who wishes to open a Crane Trader account.
I am her accountant, Andrew Cheung, of Cheung & Partners Accountants (ABN 22 604 883 175).
You can reach me on 0413 002 918 or andrew@example.org. Her personal mobile is +86 186 0028 1102 and WeChat is xiaoyu_luo88.

Kind regards,
Andrew Cheung
CPA
"""
t = write_eml("doc17_accountant_docx_attach.eml", [
    ("From", "Andrew Cheung <andrew@example.org>"), ("To", "Crane Onboarding <onboarding@example.com>"),
    ("Date", "Tue, 01 Sep 2026 11:45:00 +1000"), ("Subject", "Certified ID - new account AP-90211")],
    body, [("id_summary.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", adata, atext)])
reg("doc17_accountant_docx_attach.eml", "eml", "en", t,
    ["Ms Luo Xiaoyu", "Luo Xiaoyu", "罗晓雨", cnid2, "EA9922017",
     "Room 1502, Building 7, Jinjiang District, Chengdu, Sichuan 610021", "JP 145522", "Andrew Cheung",
     "Cheung & Partners Accountants", "22 604 883 175", "0413 002 918", "andrew@example.org", "+86 186 0028 1102",
     "xiaoyu_luo88"],
    ["Crane Onboarding", "Crane Trader", "AP-90211", "Certified ID summary", "CPA"])
LABELS["doc17_accountant_docx_attach.eml"]["pii"].remove("Cheung & Partners Accountants")  # firm name, not personal

# email with html attachment
ahtml = page("CRM export", """<h3>Customer note</h3><p>Customer Olumide Balogun (Ola) 0477 250 981, ola.balogun@example.net,
acct CB-61099120, lives at 18 Cassia Way, Baulkham Hills NSW 2153. DOB 15 March 1985.
Requested 25 shares of MSFT sell at 415.30.</p>""")
body = """Hi Quality Assurance,

Attached is the CRM extract for the call review. The customer sounded distressed. See AU-0173's handling.

Cheers,
AU-0090
"""
t = write_eml("doc18_html_attach_en.eml", [
    ("From", "AU-0090 <au0090@example.com>"), ("To", "QA <qa@example.com>"),
    ("Date", "Fri, 04 Sep 2026 14:00:00 +1000"), ("Subject", "Call review - CRM extract")],
    body, [("crm_extract.html", "text/html", ahtml.encode("utf-8"), ahtml)])
reg("doc18_html_attach_en.eml", "eml", "en", t,
    ["Olumide Balogun", "Ola", "0477 250 981", "ola.balogun@example.net", "CB-61099120",
     "18 Cassia Way, Baulkham Hills NSW 2153", "15 March 1985"],
    ["AU-0090", "AU-0173", "25 shares of MSFT", "415.30", "Call review"])

body = """尊敬的白鹤证券客服团队：

我是黄美玲，账户号 CB-70334518。上周五（2026年8月28日）我在手机App上卖出了 500 股 9988.HK，价格 HKD 96.40，
但是资金还没有到账。请尽快处理。

我的身份证号 __ID__，手机 138-2200-4519，微信 meiling_h777，QQ 邮箱 meiling.huang@example.com。
通讯地址：北京市朝阳区建国路88号SOHO现代城B座2206。

黄女士 (Huang Meiling)
2026年9月1日
"""
cn_b = cn_id("110105", "19790806", "022")
body = body.replace("__ID__", cn_b)
t = write_eml("doc19_complaint_zh.eml", [
    ("From", "黄美玲 <meiling.huang@example.com>"), ("To", "客服 <cs@example.com>"),
    ("Date", "Tue, 01 Sep 2026 08:30:00 +0800"), ("Subject", "资金未到账 - 请尽快处理")],
    body, b64_body=True)
reg("doc19_complaint_zh.eml", "eml", "zh", t,
    ["黄美玲", "CB-70334518", cn_b, "138-2200-4519", "meiling_h777", "meiling.huang@example.com",
     "北京市朝阳区建国路88号SOHO现代城B座2206", "黄女士", "Huang Meiling"],
    ["白鹤证券", "500 股 9988.HK", "HKD 96.40", "2026年8月28日", "资金未到账"])

body = """Dear Sir/Madam,

Re: Our client Mr Anthony (Tony) Vassiliadis - preservation of records

We act for Mr Vassiliadis (born 3 May 1966) in relation to his dispute with Crane Brokers concerning trade ORD-88123.
Please preserve all call recordings and chat logs. Our file reference is MK:2026:0341.

Communications may be directed to the writer, Melissa Kwan, Senior Associate, on (02) 8555 0177 or
melissa.kwan@example.org. Our client's address for service is 5/9 Beach Rd, Bondi NSW 2026. His account is
CB-15003948 and his mobile is 0400 731 220.

Yours faithfully,
Melissa Kwan
Harbourline Legal
"""
t = write_eml("doc20_legal_letter_en.eml", [
    ("From", "Melissa Kwan <melissa.kwan@example.org>"), ("To", "Legal <legal@example.com>"),
    ("Cc", "Tony V <tony.v@example.net>"), ("Date", "Mon, 07 Sep 2026 10:05:00 +1000"),
    ("Subject", "Preservation notice - our ref MK:2026:0341")], body)
reg("doc20_legal_letter_en.eml", "eml", "en", t,
    ["Anthony (Tony) Vassiliadis", "Tony", "Mr Vassiliadis", "3 May 1966", "Melissa Kwan", "(02) 8555 0177",
     "melissa.kwan@example.org", "5/9 Beach Rd, Bondi NSW 2026", "CB-15003948", "0400 731 220", "Tony V",
     "tony.v@example.net", "Mr Anthony (Tony) Vassiliadis"],
    ["Crane Brokers", "ORD-88123", "MK:2026:0341", "Preservation notice", "Senior Associate"])
LABELS["doc20_legal_letter_en.eml"]["pii"].remove("Mr Anthony (Tony) Vassiliadis")
LABELS["doc20_legal_letter_en.eml"]["pii"].remove("Anthony (Tony) Vassiliadis")
LABELS["doc20_legal_letter_en.eml"]["pii"] += ["Anthony", "Vassiliadis"]

body = """Hi Crane client services,

I'm a licensed financial adviser (AFSL 401234, adviser no. 998812) writing for my client Hyun-woo Baek, who
holds account CB-33091277 with you. He has asked that you send confirmations to me as well as him.

Client details for your records:
 - Legal name: Baek Hyun-woo (Park is his mother's surname; he sometimes signs J. Park)
 - DOB: 1983-10-19
 - Mobile: 0450 663 917
 - KakaoTalk: hwbaek83
 - Email: hw.baek@example.com

Please confirm you have noted my authority. I've attached nothing this time.

Regards,
Deborah Achterberg
Achterberg Wealth Partners
0 2 9555 0184
"""
t = write_eml("doc21_adviser_authority_en.eml", [
    ("From", "Deborah Achterberg <d.achterberg@example.org>"), ("To", "Client Services <clientservices@example.com>"),
    ("Date", "Thu, 10 Sep 2026 16:20:00 +1000"), ("Subject", "Adviser authority - account CB-33091277")], body)
reg("doc21_adviser_authority_en.eml", "eml", "en", t,
    ["Hyun-woo Baek", "CB-33091277", "Baek Hyun-woo", "J. Park", "1983-10-19", "0450 663 917", "hwbaek83",
     "hw.baek@example.com", "Deborah Achterberg", "d.achterberg@example.org", "998812", "0 2 9555 0184"],
    ["Crane", "Client Services", "AFSL 401234", "Adviser authority"])
LABELS["doc21_adviser_authority_en.eml"]["pii"].remove("CB-33091277")
LABELS["doc21_adviser_authority_en.eml"]["pii"].append("CB-33091277")  # subject and body both

body = """Hi 客服，

我是Kevin (林凯文)。 I want to change my registered phone number from 0411 552 907 to +61 487 210 395 —
the old one belongs to my ex-girlfriend Sophie Tan now.

My login email is kevin.lin@example.com and 微信 is kevin_lin_0720. 出生日期 20/07/1991, 公民 ID 护照 PE4471820.
Account number CB-99120451. 请在 2026-09-14 前处理，我要买 1000 shares of 00700.

Thx / 谢谢
凯文
"""
t = write_eml("doc22_mixed_phone_change.eml", [
    ("From", "Kevin Lin <kevin.lin@example.com>"), ("To", "客服 <cs@example.com>"),
    ("Date", "Sat, 12 Sep 2026 22:10:00 +1000"), ("Subject", "Change phone number 更改手机号")], body)
reg("doc22_mixed_phone_change.eml", "eml", "mixed", t,
    ["Kevin", "林凯文", "0411 552 907", "+61 487 210 395", "Sophie Tan", "kevin.lin@example.com", "kevin_lin_0720",
     "20/07/1991", "PE4471820", "CB-99120451", "凯文"],
    ["1000 shares of 00700", "2026-09-14", "更改手机号", "Change phone number"])

body = """Hello,

I am writing on behalf of my mother-in-law, Mrs Chrysoula Andreadis (aged 84, account CB-24119045). She passed the
account to a Power of Attorney, my husband Nikolaos Andreadis, in June. I'm sending his details below.

Nikolaos Andreadis, born 17 April 1957, TFN 611 222 903, Greek passport AN4021167.
Phone: +30 691 234 5678 (Greece) or 0424 556 011 (Australia).
Address: Rural Mail Box 214, Mount Barker SA 5251

Please do not call Chrysoula on her home phone (08) 8391 0022 - she's confused by it.

Thanks, Efthymia Andreadis (Effie)
"""
t = write_eml("doc23_poa_family_en.eml", [
    ("From", "Effie <effie.andreadis@example.net>"), ("To", "Crane <support@example.com>"),
    ("Cc", "Nikolaos Andreadis <nick.andreadis@example.net>"), ("Date", "Sun, 13 Sep 2026 09:00:00 +0930"),
    ("Subject", "Power of Attorney - Andreadis account")], body)
reg("doc23_poa_family_en.eml", "eml", "en", t,
    ["Mrs Chrysoula Andreadis", "Chrysoula Andreadis", "CB-24119045", "Nikolaos Andreadis", "17 April 1957",
     "611 222 903", "AN4021167", "+30 691 234 5678", "0424 556 011", "Rural Mail Box 214, Mount Barker SA 5251",
     "Chrysoula", "(08) 8391 0022", "Efthymia Andreadis", "Effie", "effie.andreadis@example.net",
     "nick.andreadis@example.net", "Andreadis"],
    ["Crane", "Power of Attorney", "support@example.com"])
LABELS["doc23_poa_family_en.eml"]["pii"].remove("Andreadis")
LABELS["doc23_poa_family_en.eml"]["pii"].remove("Mrs Chrysoula Andreadis")
LABELS["doc23_poa_family_en.eml"]["pii"].append("Mrs Chrysoula")  if False else None
LABELS["doc23_poa_family_en.eml"]["must_keep"].remove("support@example.com")

body = """白鹤证券 合规部：

关于客户 郑海燕 (Zheng Haiyan, Helen) 的账户 CB-81902275 涉及大额出金一事，请按公司反洗钱程序复核。

客户信息：
  身份证：__ID__
  香港身份证：__HK__
  手机：+852 6234 8890 / +86 159 8800 3321
  地址：香港九龙尖沙咀弥敦道100号 1503室
  生日：1969年12月30日

案件号 AML-2026-0512，经办人 CO-03，复核日期 2026-09-11，涉及金额 HKD 1,200,000。
Helen 郑女士 声称款项来自出售物业。请附上买卖合约。

祝好
CO-03
"""
cn_c = cn_id("330106", "19691230", "045")
hk2 = hk_id("R", "550123")
body = body.replace("__ID__", cn_c).replace("__HK__", hk2)
t = write_eml("doc24_compliance_zh.eml", [
    ("From", "CO-03 <co03@example.com>"), ("To", "合规部 <compliance@example.com>"),
    ("Date", "Fri, 11 Sep 2026 17:40:00 +1000"), ("Subject", "AML-2026-0512 大额出金复核")], body)
reg("doc24_compliance_zh.eml", "eml", "zh", t,
    ["郑海燕", "Zheng Haiyan", "Helen", "CB-81902275", cn_c, hk2, "+852 6234 8890", "+86 159 8800 3321",
     "香港九龙尖沙咀弥敦道100号 1503室", "1969年12月30日", "郑女士"],
    ["白鹤证券", "AML-2026-0512", "CO-03", "2026-09-11", "HKD 1,200,000"])

# =====================================================================
# DOCX
# =====================================================================
cnid3 = cn_id("420106", "19750914", "037")
t = write_docx("doc25_kyc_form_en.docx", [
    ("h", "Crane Brokers - Individual Account Application"),
    ("p", "Section A: Applicant details. Application ref AP-91032."),
    ("t", [["Field", "Response"],
           ["Title / Given names", "Mr Thanh Dat"],
           ["Surname", "Le"],
           ["Preferred name", "Danny"],
           ["Date of birth", "07 / 11 / 1990"],
           ["Residential address", "Unit 9, 44 Hamilton Rd, Fairfield NSW 2165"],
           ["Postal address", "PO Box 302, Cabramatta NSW 2166"],
           ["Mobile", "0421 300 588"],
           ["Email", "danny.le@example.com"],
           ["Tax File Number", "234 567 891"],
           ["Driver licence", "NSW 20387711"],
           ["Passport", "PB0192283 (Australia)"],
           ["Bank (settlement)", "BSB 032-001, Acc 445 908"],
           ["Employer", "Fairfield Plumbing Supplies, ABN 44 129 750 311"],
           ["Source of wealth", "Salary and savings"]]),
    ("p", "Section B: Trading experience - shares 3 years; options none. Products: Crane Trader, US equities."),
    ("p", "Declaration signed: Thanh Dat Le, 2026-08-30. Witness: AU-0155."),
], header="Crane Brokers | CONFIDENTIAL", footer="Applicant: Le Thanh Dat - page 1 of 1")
reg("doc25_kyc_form_en.docx", "docx", "en", t,
    ["Mr Thanh Dat", "Le", "Danny", "07 / 11 / 1990", "Unit 9, 44 Hamilton Rd, Fairfield NSW 2165",
     "PO Box 302, Cabramatta NSW 2166", "0421 300 588", "danny.le@example.com", "234 567 891", "20387711",
     "PB0192283", "032-001", "445 908", "44 129 750 311", "Thanh Dat Le", "Le Thanh Dat"],
    ["Crane Brokers", "AP-91032", "Crane Trader", "US equities", "AU-0155", "Individual Account Application",
     "Fairfield Plumbing Supplies", "2026-08-30"])
LABELS["doc25_kyc_form_en.docx"]["pii"].remove("Le")  # too ambiguous as a bare token

t = write_docx("doc26_beneficiary_nomination_en.docx", [
    ("h", "Beneficiary Nomination Form"),
    ("p", "Account CB-50127384. Nominated by: Bronwyn Ellery-Ashcroft (DOB 4 December 1958)."),
    ("t", [["Beneficiary", "Relationship", "Share", "Contact"],
           ["Callum Ashcroft", "Son", "50%", "0417 998 202, callum.a@example.net"],
           ["Isla Ashcroft-Reyes", "Granddaughter (minor, DOB 19/08/2014)", "25%", "c/o Callum Ashcroft"],
           ["Nam Sook Kang", "Friend / carer", "25%", "nsk1949@example.org; 3 Lorikeet Ave, Robina QLD 4226"]]),
    ("p", "Bronwyn asks that Crane Brokers freeze trading on her death. Contact Bron on 0409 114 780."),
    ("p", "Witnessed by Dr Anil Sharma, Justice of the Peace no. JP 220481, 14 September 2026."),
], footer="Form NOM-01 v3 - Crane Brokers")
reg("doc26_beneficiary_nomination_en.docx", "docx", "en", t,
    ["CB-50127384", "Bronwyn Ellery-Ashcroft", "4 December 1958", "Callum Ashcroft", "0417 998 202", "callum.a@example.net",
     "Isla Ashcroft-Reyes", "19/08/2014", "Nam Sook Kang", "nsk1949@example.org", "3 Lorikeet Ave, Robina QLD 4226",
     "Bronwyn", "Bron", "0409 114 780", "Dr Anil Sharma", "JP 220481"],
    ["Beneficiary Nomination Form", "NOM-01 v3", "Crane Brokers", "50%", "25%", "Son", "Friend / carer"])

t = write_docx("doc27_hardship_note_en.docx", [
    ("h", "Financial Hardship - Case Note"),
    ("p", "Case HS-2026-118. Handler AU-0400. Date 2026-09-08."),
    ("p", "Ms Bui Kim Ngan (goes by Kim) rang on 0432 118 007 explaining that she lost her job at the end of July and her partner,"
          " Ravi Chandrasekhar, has been ill. She is 3 weeks behind on a margin call of AUD 2,300 on her Crane Trader margin account."),
    ("p", "Verified by DOB 28 February 1976 and account no. CB-19022084. Email on file: kim.bui@example.com."),
    ("p", "Hardship options explained: reduce position by 60 shares of QAN, 30-day payment plan. Kim said 'don't call my"
          " employer' - previous employer Southern Cross Freight, contact Ms Talia Moore on (03) 9555 0333."),
    ("p", "Home: 6/210 Sydney Road, Brunswick VIC 3056. Centrelink CRN 123 456 789A. Carer for her mother, Tuyet Bui."),
], header="INTERNAL - Hardship team")
reg("doc27_hardship_note_en.docx", "docx", "en", t,
    ["Ms Bui Kim Ngan", "Kim", "0432 118 007", "Ravi Chandrasekhar", "28 February 1976", "CB-19022084",
     "kim.bui@example.com", "Ms Talia Moore", "(03) 9555 0333", "6/210 Sydney Road, Brunswick VIC 3056",
     "123 456 789A", "Tuyet Bui"],
    ["HS-2026-118", "AU-0400", "AUD 2,300", "Crane Trader", "60 shares of QAN", "margin call", "Southern Cross Freight",
     "Hardship team", "2026-09-08"])

t = write_docx("doc28_deceased_estate_en.docx", [
    ("h", "Deceased Estate Notification"),
    ("t", [["Deceased", "Harold Ignatius Fenwick"],
           ["Date of death", "2026-08-16"],
           ["Date of birth", "3rd of June 1939"],
           ["Account", "CB-11987345 (joint with Doris Fenwick)"],
           ["Executor", "Lachlan Fenwick, 41 Ironbark Drive, Kiama NSW 2533"],
           ["Executor phone", "0412 776 340"],
           ["Executor email", "l.fenwick@example.org"],
           ["Solicitor", "Nasrin Haddad, Haddad Hughes Solicitors, GPO Box 1442 Sydney NSW 2001"]]),
    ("p", "Holdings at date of death: 1,500 shares of TLS, 400 of WBC, cash AUD 18,442.10. Crane has restricted the account."),
    ("p", "Staff note (AU-0331): executor Lachlan called twice; Harry's widow Doris is in aged care and cannot sign."),
], footer="Estates team - Crane Brokers")
reg("doc28_deceased_estate_en.docx", "docx", "en", t,
    ["Harold Ignatius Fenwick", "3rd of June 1939", "CB-11987345", "Doris Fenwick", "Lachlan Fenwick",
     "41 Ironbark Drive, Kiama NSW 2533", "0412 776 340", "l.fenwick@example.org", "Nasrin Haddad",
     "Haddad Hughes Solicitors", "GPO Box 1442 Sydney NSW 2001", "Lachlan", "Harry", "Doris"],
    ["Deceased Estate Notification", "2026-08-16", "1,500 shares of TLS", "AUD 18,442.10", "AU-0331", "Estates team"])
LABELS["doc28_deceased_estate_en.docx"]["pii"].remove("Haddad Hughes Solicitors")

cnid_a = cn_id("110108", "19830521", "016")
t = write_docx("doc29_客户资料确认书_zh.docx", [
    ("h", "客户资料确认书"),
    ("p", "本确认书由白鹤证券（Crane Brokers）出具，请核对以下信息，如有错误请于 2026年9月20日前书面通知。"),
    ("t", [["项目", "内容"],
           ["客户姓名", "赵明轩"],
           ["英文姓名", "Ming Xuan Zhao (Marcus)"],
           ["证件类型", "居民身份证"],
           ["证件号码", cnid_a],
           ["出生日期", "1983年5月21日"],
           ["手机号码", "+86 135 2210 6789"],
           ["电子邮箱", "marcus.zhao@example.com"],
           ["微信号", "zmx_marcus83"],
           ["通讯地址", "北京市海淀区中关村大街27号中关村大厦9层903室，100080"],
           ["银行账户", "招商银行 6214 8300 1188 7745"],
           ["证券账户号", "CB-66120384"],
           ["联系人（紧急）", "赵敏（妹妹） 手机 138 0011 2233"]]),
    ("p", "赵先生确认以上信息真实、准确。产品：美股、港股、期权。工单：TKT-810004。经办：AU-0410。"),
], header="白鹤证券 | 客户资料确认书", footer="客户：赵明轩  第1页")
reg("doc29_客户资料确认书_zh.docx", "docx", "zh", t,
    ["赵明轩", "Ming Xuan Zhao", "Marcus", cnid_a, "1983年5月21日", "+86 135 2210 6789", "marcus.zhao@example.com",
     "zmx_marcus83", "北京市海淀区中关村大街27号中关村大厦9层903室，100080", "6214 8300 1188 7745",
     "CB-66120384", "赵敏", "138 0011 2233", "赵先生"],
    ["白鹤证券", "Crane Brokers", "客户资料确认书", "TKT-810004", "AU-0410", "美股、港股、期权", "2026年9月20日", "招商银行"])

t = write_docx("doc30_chat_transcript_en.docx", [
    ("h", "Chat Transcript 2026-09-15 - Crane Brokers"),
    ("p", "[Customer]<the_real_hamish> 2026-09-15 12:01:44 Hamish MacLeod here. why was my stop loss on 300 GOLD triggered?"),
    ("p", "[agent]AU-0088 2026-09-15 12:02:20 Hi Hamish, can I verify your DOB and mobile?"),
    ("p", "[Customer]<the_real_hamish> 2026-09-15 12:02:59 11 Aug 1972, 0498 771 322"),
    ("p", "[Bot]CraneAssist 2026-09-15 12:03:00 Reminder: never share passwords."),
    ("p", "[agent]AU-0088 2026-09-15 12:03:40 Thanks Mr MacLeod. Stop was at 4.15, gold dropped to 4.11 at 09:31:05, order ORD-220918."),
    ("p", "[Customer]<the_real_hamish> 2026-09-15 12:04:10 my accountant Ines Ferreira-Duarte says I should escalate. ines@example.net"),
    ("p", "[System] Escalation created TKT-820771 at 12:04:30"),
    ("p", "[Customer]<the_real_hamish> 2026-09-15 12:04:55 ta. ring me back on 07 3555 0163 quote CLI 5590 2213 8847"),
], footer="Exported by AU-0088")
reg("doc30_chat_transcript_en.docx", "docx", "en", t,
    ["the_real_hamish", "Hamish MacLeod", "Hamish", "11 Aug 1972", "0498 771 322", "Mr MacLeod",
     "Ines Ferreira-Duarte", "ines@example.net", "07 3555 0163"],
    ["Crane Brokers", "AU-0088", "300 GOLD", "4.15", "4.11", "09:31:05", "ORD-220918", "TKT-820771", "CraneAssist"])
LABELS["doc30_chat_transcript_en.docx"]["pii"].append("5590 2213 8847")

t = write_docx("doc31_中文聊天记录.docx", [
    ("h", "在线客服聊天记录"),
    ("p", "[Customer]<lisa_cn_1990> 2026-09-16 21:14:02 你好，我是李思雨，想问一下入金 HKD 的手续费"),
    ("p", "[agent]AU-0501 2026-09-16 21:14:40 您好李女士，Crane 的 HKD 入金免手续费。请问账户号是多少？"),
    ("p", "[Customer]<lisa_cn_1990> 2026-09-16 21:15:20 CB-43021998。我的英文名叫 Lisa Li，微信 lisa_siyu90，QQ 号 2280913377"),
    ("p", "[Bot]CraneAssist 2026-09-16 21:15:21 请勿透露密码。"),
    ("p", "[agent]AU-0501 2026-09-16 21:16:05 好的，思雨，需要验证您的证件后6位。"),
    ("p", "[Customer]<lisa_cn_1990> 2026-09-16 21:16:33 我身份证 __ID__，手机 187 0123 4567，住在杭州市西湖区文三路90号东部软件园"),
    ("p", "[agent]AU-0501 2026-09-16 21:17:10 收到，李女士。工单 TKT-821330。"),
])
cn_d = cn_id("330106", "19900315", "062")
# patch the id placeholder by regenerating
t = write_docx("doc31_中文聊天记录.docx", [
    ("h", "在线客服聊天记录"),
    ("p", "[Customer]<lisa_cn_1990> 2026-09-16 21:14:02 你好，我是李思雨，想问一下入金 HKD 的手续费"),
    ("p", "[agent]AU-0501 2026-09-16 21:14:40 您好李女士，Crane 的 HKD 入金免手续费。请问账户号是多少？"),
    ("p", "[Customer]<lisa_cn_1990> 2026-09-16 21:15:20 CB-43021998。我的英文名叫 Lisa Li，微信 lisa_siyu90，QQ 号 2280913377"),
    ("p", "[Bot]CraneAssist 2026-09-16 21:15:21 请勿透露密码。"),
    ("p", "[agent]AU-0501 2026-09-16 21:16:05 好的，思雨，需要验证您的证件后6位。"),
    ("p", f"[Customer]<lisa_cn_1990> 2026-09-16 21:16:33 我身份证 {cn_d}，手机 187 0123 4567，住在杭州市西湖区文三路90号东部软件园"),
    ("p", "[agent]AU-0501 2026-09-16 21:17:10 收到，李女士。工单 TKT-821330。"),
])
reg("doc31_中文聊天记录.docx", "docx", "zh", t,
    ["lisa_cn_1990", "李思雨", "李女士", "CB-43021998", "Lisa Li", "lisa_siyu90", "2280913377", "思雨", cn_d,
     "187 0123 4567", "杭州市西湖区文三路90号东部软件园"],
    ["AU-0501", "CraneAssist", "HKD", "TKT-821330", "在线客服聊天记录", "2026-09-16 21:14:02", "入金"])

t = write_docx("doc32_compliance_case_mixed.docx", [
    ("h", "Compliance Case CC-2026-233 / 合规个案"),
    ("p", "Subject 客户: Xu Jiayi 徐佳怡 (nickname 'Jojo'), holder of CB-92014476 and HK account 0129-880-45."),
    ("t", [["Identifier", "Value"], ["CN ID 身份证", cn_id("320102", "19870109", "088")],
           ["HKID", hk_id("C", "301234")], ["AU visa / Passport", "E88213047 (PRC), visa ref 2026 0912 3311"],
           ["Phones", "+61 430 118 222; +86 138 1234 8899; +852 5566 3300"],
           ["Emails", "jojo.xu@example.com; xujiayi88@example.net"],
           ["Address 地址", "1/30 Elizabeth St, Sydney NSW 2000; 江苏省南京市玄武区中山路1号"]]),
    ("p", "Summary: 客户 reported the account was accessed from an unrecognised IP on 2026-09-03 03:14. Two sells of"
          " 100 shares of BABA and 60 NIO were executed at USD 118.20 and USD 5.02. Jojo denies placing them."),
    ("p", "徐女士 was spoken to by AU-0270 on 2026-09-04. Her husband Alan Xu (徐志强) is a joint signatory with DOB 1985/03/19."),
    ("p", "Outcome: reversed trades, password reset, 2FA enabled. SMR not required. Reviewer CO-11."),
], header="CONFIDENTIAL - Compliance", footer="Case CC-2026-233 - Crane Brokers")
cn_e, hk3 = cn_id("320102", "19870109", "088"), hk_id("C", "301234")
reg("doc32_compliance_case_mixed.docx", "docx", "mixed", t,
    ["Xu Jiayi", "徐佳怡", "Jojo", "CB-92014476", "0129-880-45", cn_e, hk3, "E88213047", "+61 430 118 222",
     "+86 138 1234 8899", "+852 5566 3300", "jojo.xu@example.com", "xujiayi88@example.net",
     "1/30 Elizabeth St, Sydney NSW 2000", "江苏省南京市玄武区中山路1号", "徐女士", "Alan Xu", "徐志强", "1985/03/19"],
    ["Compliance Case CC-2026-233", "AU-0270", "CO-11", "100 shares of BABA", "60 NIO", "USD 118.20", "USD 5.02",
     "2026-09-03 03:14", "Crane Brokers", "2FA"])

t = write_docx("doc33_smsf_trustee_form_en.docx", [
    ("h", "SMSF Trustee Application - Supplementary"),
    ("t", [["Fund name", "Pham Family Superannuation Fund"], ["Fund ABN", "84 507 331 902"],
           ["Trustee 1", "Duc Pham (also 'David Pham')"], ["Trustee 1 DOB", "1966-06-06"],
           ["Trustee 2", "Hanh Pham nee Vu"], ["Trustee 2 DOB", "12 January 1968"],
           ["Trustee TFN", "376 112 908"], ["Fund bank", "BSB 484-799 Acc 2233 4455"],
           ["Contact", "duc.pham@example.com or 0411 909 877"],
           ["Fund address", "Lot 3 Nowra Rd, Worrigee NSW 2540"]]),
    ("p", "Investment strategy: ASX ETFs (VAS, IVV) and US large caps, max 20 percent in a single stock. Ticket TKT-830077."),
    ("p", "Mr Pham confirmed by phone with AU-0163 on 2026-09-18 at 11:30 that Hanh will co-sign."),
], header="Crane Brokers - SMSF desk")
reg("doc33_smsf_trustee_form_en.docx", "docx", "en", t,
    ["Pham Family Superannuation Fund", "84 507 331 902", "Duc Pham", "David Pham", "1966-06-06", "Hanh Pham nee Vu",
     "12 January 1968", "376 112 908", "484-799", "2233 4455", "duc.pham@example.com", "0411 909 877",
     "Lot 3 Nowra Rd, Worrigee NSW 2540", "Mr Pham", "Hanh"],
    ["SMSF Trustee Application", "VAS", "IVV", "TKT-830077", "AU-0163", "2026-09-18 at 11:30", "20 percent"])
LABELS["doc33_smsf_trustee_form_en.docx"]["pii"].remove("Pham Family Superannuation Fund")
LABELS["doc33_smsf_trustee_form_en.docx"]["pii"].append("Pham Family Superannuation Fund")

t = write_docx("doc34_call_note_hk_en.docx", [
    ("h", "Call Note - Hong Kong Client"),
    ("p", "Handler AU-0144. Call 2026-09-21 10:05 for 12 min. Client: Mr Lau Kwok-wai (Ricky), HKID "
          + hk_id("G", "654321") + "."),
    ("p", "Ricky wants to move 800 units of 2800.HK from Crane HK to the Australian entity. Bank: HSBC HK 004-123-456789-833."),
    ("p", "Contact: +852 9765 4321, ricky.lau@example.org, WeChat rickylau_hk, Line ID rickylau.line."),
    ("p", "Address: Room 1203, 12/F, Wing On Centre, 111 Connaught Road Central, Hong Kong."),
    ("p", "Colleague reference: Ms Cheung Mei-ling, Compliance HK (ext 2201)."),
])
hk4 = hk_id("G", "654321")
reg("doc34_call_note_hk_en.docx", "docx", "en", t,
    ["Mr Lau Kwok-wai", "Lau Kwok-wai", "Ricky", hk4, "004-123-456789-833", "+852 9765 4321", "ricky.lau@example.org",
     "rickylau_hk", "rickylau.line", "Room 1203, 12/F, Wing On Centre, 111 Connaught Road Central, Hong Kong",
     "Ms Cheung Mei-ling"],
    ["AU-0144", "800 units of 2800.HK", "Crane HK", "HSBC HK", "Call Note", "2026-09-21 10:05", "Compliance HK"])
LABELS["doc34_call_note_hk_en.docx"]["pii"].remove("Mr Lau Kwok-wai")

t = write_docx("doc35_account_closure_mixed.docx", [
    ("h", "Account Closure Request / 销户申请"),
    ("t", [["Field / 项目", "Value / 内容"],
           ["Name 姓名", "Wu Zhihao 吴志豪"],
           ["Nickname", "Hao"],
           ["Account", "CB-77031258"],
           ["Reason 原因", "Relocating to Singapore 搬迁至新加坡"],
           ["Phone", "0468 200 331"],
           ["Email", "zhihao.wu@example.net"],
           ["Alt contact 备用联系人", "吴太太 (Mrs Wu, mother) +86 137 7788 9900"]]),
    ("p", "Remaining balance USD 2,140.66 to be sent to bank account: DBS Singapore 012-345-6789, name on account Wu Zhihao."),
    ("p", "Closure processed by AU-0092 on 2026-09-22. Ticket TKT-831456."),
])
reg("doc35_account_closure_mixed.docx", "docx", "mixed", t,
    ["Wu Zhihao", "吴志豪", "Hao", "CB-77031258", "0468 200 331", "zhihao.wu@example.net", "吴太太", "Mrs Wu",
     "+86 137 7788 9900", "012-345-6789"],
    ["Account Closure Request", "销户申请", "USD 2,140.66", "AU-0092", "TKT-831456", "DBS Singapore", "2026-09-22"])

t = write_docx("doc36_gift_transfer_en.docx", [
    ("h", "Off-market Transfer Request"),
    ("p", "From: Seo-yeon Jung (account CB-30921177), DOB 09/05/1979. To: her son Min-jun Jung, DOB 2001-02-14."),
    ("p", "Transfer 250 shares of TSLA as a gift. Cost base per share AUD 92.10. Both signed at 3/12 Kent St, Epping NSW 2121."),
    ("p", "Sharing my contact: 0455 320 118. Minjun's is 0466 320 118. Emails: sy.jung@example.com and minjun.jung@example.com."),
    ("p", "Passports: KR M12345678; KR M87654321. Please process quickly. - S.Y. Jung"),
], footer="Crane Brokers Transfers")
reg("doc36_gift_transfer_en.docx", "docx", "en", t,
    ["Seo-yeon Jung", "CB-30921177", "09/05/1979", "Min-jun Jung", "2001-02-14", "3/12 Kent St, Epping NSW 2121",
     "0455 320 118", "Minjun", "0466 320 118", "sy.jung@example.com", "minjun.jung@example.com", "M12345678",
     "M87654321", "S.Y. Jung"],
    ["Off-market Transfer Request", "250 shares of TSLA", "AUD 92.10", "Crane Brokers Transfers", "gift"])

# =====================================================================
assert 30 <= len(LABELS) <= 40, len(LABELS)
for n in LABELS:
    assert (OUT / n).exists(), n
    assert re.fullmatch(r"[A-Za-z0-9_一-鿿.]+", n)
(ROOT / "labels.json").write_text(json.dumps(LABELS, ensure_ascii=False, indent=1), encoding="utf-8")
print(len(LABELS), "docs;", sum(len(v["pii"]) for v in LABELS.values()), "pii;",
      sum(len(v["must_keep"]) for v in LABELS.values()), "must_keep")
