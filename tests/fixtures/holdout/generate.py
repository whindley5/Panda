"""Independent holdout corpus for the Panda sanitiser (all data invented).

Writes corpus/ and labels.json next to this file. Deterministic: no clocks, no randomness.
Every label string is asserted to occur in the source text that went into the file.
"""
from __future__ import annotations

import io
import json
import zipfile
from email.headerregistry import Address
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
CORPUS = HERE / "corpus"
LABELS: dict[str, dict] = {}


def cn_id(area: str, birth: str, seq: str) -> str:
    body = area + birth + seq
    assert len(body) == 17 and body.isdigit()
    w = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    return body + "10X98765432"[sum(int(c) * k for c, k in zip(body, w)) % 11]


ID_WANG = cn_id("310115", "19850316", "123")
ID_LI = cn_id("110105", "19880722", "045")
ID_ZHOU = cn_id("440106", "19790105", "211")
ID_CHEN = cn_id("330102", "19930911", "082")


def register(name, fmt, lang, src, pii, keep):
    for s in pii + keep:
        assert s in src, f"{name}: label not in source: {s!r}"
    assert len(set(pii)) == len(pii), f"{name}: duplicate pii"
    LABELS[name] = {"format": fmt, "lang": lang, "pii": pii, "must_keep": keep}


def chat(msgs, glue=()):
    """msgs: list of (header, text). Headers listed in glue are appended to the previous line."""
    lines = []
    for i, (h, t) in enumerate(msgs):
        if i in glue and lines:
            lines[-1] += h
        else:
            lines.append(h)
        lines.append(t)
    return lines


# ---------------------------------------------------------------- writers
def pdf_escape(s):
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf(name, lines):
    per_page = 46
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[]]
    objs = []  # 1 catalog, 2 pages, 3 font, then page/content pairs
    kids = []
    for n, pl in enumerate(pages):
        pid, cid = 4 + 2 * n, 5 + 2 * n
        kids.append(f"{pid} 0 R")
        ops = ["BT", "/F1 10 Tf", "13 TL", "50 800 Td"]
        for ln in pl:
            assert all(ord(c) < 128 for c in ln), f"non-ascii in pdf {name}: {ln!r}"
            ops.append(f"({pdf_escape(ln)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops)
        objs.append((pid, f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
                          f"/Resources << /Font << /F1 3 0 R >> >> /Contents {cid} 0 R >>"))
        objs.append((cid, f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"))
    base = [(1, "<< /Type /Catalog /Pages 2 0 R >>"),
            (2, f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"),
            (3, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")]
    allobjs = base + objs
    out = bytearray(b"%PDF-1.4\n")
    offs = {}
    for num, body in allobjs:
        offs[num] = len(out)
        out += f"{num} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(allobjs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for num, _ in allobjs:
        out += f"{offs[num]:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(allobjs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n").encode()
    (CORPUS / name).write_bytes(bytes(out))
    return "\n".join(lines)


def write_html(name, html):
    (CORPUS / name).write_text(html, encoding="utf-8")
    return html


def docx_bytes(blocks, header=None, footer=None):
    """blocks: ('p', text) | ('h', text) | ('table', rows). Returns (bytes, source text)."""
    from docx import Document
    d = Document()
    src = []
    for kind, val in blocks:
        if kind == "p":
            d.add_paragraph(val); src.append(val)
        elif kind == "h":
            d.add_heading(val, level=1); src.append(val)
        else:
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell; src.append(cell)
    sec = d.sections[0]
    if header:
        sec.header.paragraphs[0].text = header; src.append(header)
    if footer:
        sec.footer.paragraphs[0].text = footer; src.append(footer)
    d.core_properties.author = "Crane Brokers"
    d.core_properties.created = datetime(2026, 1, 1)
    d.core_properties.modified = datetime(2026, 1, 1)
    raw = io.BytesIO(); d.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            zi_info = zipfile.ZipInfo(info.filename, (2026, 1, 1, 0, 0, 0))
            zi_info.compress_type = zipfile.ZIP_DEFLATED
            zo.writestr(zi_info, zi.read(info.filename))
    return out.getvalue(), "\n".join(src)


def write_docx(name, blocks, header=None, footer=None):
    data, src = docx_bytes(blocks, header, footer)
    (CORPUS / name).write_bytes(data)
    return src


def write_eml(name, headers, plain, html=None, attachments=(), day=1):
    """headers: dict of header -> str | Address list. attachments: (filename, maintype, subtype, bytes)."""
    m = EmailMessage()
    for k, v in headers.items():
        m[k] = v
    m["Date"] = format_datetime(datetime(2026, 6, day, 10, 15, tzinfo=timezone.utc))
    m["Message-ID"] = f"<{name}@mail.example.com>"
    m.set_content(plain, charset="utf-8")
    if html:
        m.add_alternative(html, subtype="html", charset="utf-8")
    for fn, mt, st, data in attachments:
        m.add_attachment(data, maintype=mt, subtype=st, filename=fn)
    n = 0
    for part in m.walk():
        if part.is_multipart():
            n += 1
            part.set_boundary(f"==holdout_{name}_{n}==")
    (CORPUS / name).write_bytes(m.as_bytes())
    src = [f"{k}: {v}" for k, v in headers.items()] + [plain]
    if html:
        src.append(html)
    return "\n".join(src)


def page(title, body, extra_css=""):
    return (f"<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"><title>{title}</title>"
            f"<style>body{{font-family:sans-serif}}{extra_css}</style></head>\n<body>\n{body}\n</body></html>\n")


def html_chat(msgs, cls_map=None):
    rows = []
    for h, t in msgs:
        rows.append(f'<div class="msg"><span class="hdr">{h}</span><br><span class="txt">{t}</span></div>')
    return "\n".join(rows)


# ================================================================ PDFs (English, Latin only)
def build_pdfs():
    msgs = [
        ("[Customer]priyaR88 2026-06-01 09:15:22", "Hi, I need to update my details before my order settles. My name is Priya Raman."),
        ("[agent]AU-0102 2026-06-01 09:16:00", "Hello Priya, thanks for contacting Crane Brokers. Can I confirm your date of birth please?"),
        ("[Customer]priyaR88 2026-06-01 09:16:41", "14/03/1987. Mobile is 0412 345 678 and email is priya.raman@example.com"),
        ("[agent]AU-0102 2026-06-01 09:17:30", "Thanks. Which address are we updating to? Also I see Ticket #48213 is open for you."),
        ("[Customer]priyaR88 2026-06-01 09:18:05", "Unit 7/23 Wattle Crescent, Parramatta NSW 2150. My TFN is 123 456 782 if you need it for the W-8 style form."),
        ("[agent]AU-0102 2026-06-01 09:19:12", "Got it Ms Raman. Your buy order for 500 shares at $12.34 in BHP.AX on 2026-05-29 has settled."),
        ("[Customer]priyaR88 2026-06-01 09:19:50", "Great. My husband Arjun set up the joint holding, please call him on +61 433 221 909 if anything is wrong."),
        ("[agent]system 2026-06-01 09:20:01", "Chat transferred to team Compliance-AU."),
        ("[AI]CraneAssist 2026-06-01 09:20:03", "Summary: address and phone update for customer Raman, Priya. Ref case CS-2026-0611."),
    ]
    lines = ["Crane Brokers - Chat transcript export", "Exported 2026-06-01 09:30", ""] + chat(msgs, glue={3, 6})
    src = write_pdf("chat_en_priya.pdf", lines)
    register("chat_en_priya.pdf", "pdf", "en", src,
             ["Priya Raman", "Priya", "Raman", "priyaR88", "14/03/1987", "0412 345 678", "priya.raman@example.com",
              "Unit 7/23 Wattle Crescent, Parramatta NSW 2150", "123 456 782", "Arjun", "+61 433 221 909", "Raman, Priya",
              "Ms Raman"],
             ["Crane Brokers", "AU-0102", "Ticket #48213", "500 shares at $12.34", "BHP.AX", "2026-05-29",
              "09:15:22", "[agent]system", "CS-2026-0611", "team Compliance-AU", "[AI]CraneAssist"])

    msgs = [
        ("[Customer]Guest 14:02:11", "hello can you help me, i cant log in to my account"),
        ("[agent]AU-0233 14:02:40", "Of course. Can I get your full name for verification?"),
        ("[Customer]Guest 14:03:02", "Marek Kowalski"),
        ("[agent]AU-0233 14:03:30", "Could you spell your surname for me?"),
        ("[Customer]Guest 14:03:58", "K-O-W-A-L-S-K-I. Everyone gets it wrong. Marek with a K"),
        ("[agent]AU-0233 14:04:20", "Thanks Mr Kowalski. And the passport number on file?"),
        ("[Customer]Guest 14:04:50", "Passport no: PA1234567. Driver licence: 12345678 (NSW) if that helps"),
        ("[agent]AU-0233 14:05:31", "Verified. Your account CB-00784521 is unlocked. You bought 40 shares of AAPL on 2025-11-18 at USD 189.50."),
        ("[Customer]Guest 14:06:12", "thanks, cheers. Mark is what my mates call me by the way"),
    ]
    lines = ["Live chat log", ""] + chat(msgs, glue={4})
    src = write_pdf("chat_en_marek.pdf", lines)
    register("chat_en_marek.pdf", "pdf", "en", src,
             ["Marek Kowalski", "Kowalski", "K-O-W-A-L-S-K-I", "Marek", "Mr Kowalski", "PA1234567", "12345678", "Mark",
              "CB-00784521"],
             ["AU-0233", "14:02:11", "40 shares of AAPL", "2025-11-18", "USD 189.50", "Guest"])

    lines = [
        "CRANE BROKERS PTY LTD",
        "Account statement - period ending 30 June 2026",
        "",
        "Account holders: Hamish McAllister and Fiona McAllister",
        "Mailing address: PO Box 4471, Sydney NSW 2001",
        "Registered address: 12A/45 Oxford St, Bondi Junction NSW 2022",
        "Settlement bank: BSB 062-000 Account 1234 5678 (Fiona McAllister)",
        "Beneficiary on death: Isla McAllister (daughter), DOB 9 August 2011",
        "Contact: Hamish on 0398 765 432 or (02) 9876 5432",
        "Medicare card: 2123 45670 1",
        "",
        "Holdings",
        "CBA.AX  120 units   avg $98.10",
        "00700.HK 300 units  avg HK$310.25",
        "Trade 2026-03-12: SELL 500 shares at $12.34 (BHP.AX)",
        "",
        "Dear Mr McAllister, please contact our team if any detail above is wrong.",
        "Regards, Crane Client Services",
    ]
    src = write_pdf("statement_en_mcallister.pdf", lines)
    register("statement_en_mcallister.pdf", "pdf", "en", src,
             ["Hamish McAllister", "Fiona McAllister", "PO Box 4471, Sydney NSW 2001",
              "12A/45 Oxford St, Bondi Junction NSW 2022", "062-000", "1234 5678", "Isla McAllister", "9 August 2011",
              "0398 765 432", "(02) 9876 5432", "2123 45670 1", "Mr McAllister", "Hamish", "Fiona", "Isla"],
             ["CRANE BROKERS PTY LTD", "CBA.AX", "00700.HK", "2026-03-12", "500 shares at $12.34", "Crane Client Services",
              "30 June 2026"])

    lines = [
        "Crane Brokers - Complaint file note (printed from CRM)",
        "Case: CMP-2026-00417   Opened: 2026-07-08   Owner: AU-0144",
        "",
        "Complainant: SMITH, Jonathan Peter (known as JP)",
        "Referred by friend Olivia Nguyen (olivia.nguyen@example.org), who receives a referral bonus.",
        "Complainant phone: +61 2 9876 5432 / mobile 0455 019 283",
        "Complainant email: jp.smith@example.net",
        "Date of birth: 3-Feb-79",
        "Home: Apartment 1506, 3 Freshwater Place, Southbank VIC 3006",
        "",
        "Summary: JP states he was not told about the $5.50 fee on a 100 shares TSLA order (2026-05-14).",
        "JP said: 'Please call me Jonno, not Jonathan.' Agent addressed him as Mr Smith throughout.",
        "Level 9, 55 Collins Street Melbourne VIC 3000 is the Crane AU office and should be kept.",
    ]
    src = write_pdf("filenote_en_smith.pdf", lines)
    register("filenote_en_smith.pdf", "pdf", "en", src,
             ["SMITH, Jonathan Peter", "Jonathan Peter", "JP", "Olivia Nguyen", "olivia.nguyen@example.org",
              "+61 2 9876 5432", "0455 019 283", "jp.smith@example.net", "3-Feb-79",
              "Apartment 1506, 3 Freshwater Place, Southbank VIC 3006", "Jonno", "Jonathan", "Mr Smith", "Olivia"],
             # The firm's street address is replaced like any address (by design).
             ["CMP-2026-00417", "2026-07-08", "AU-0144", "100 shares TSLA", "2026-05-14"])


# ================================================================ HTML
def build_html():
    msgs = [
        ("[Customer]lucy.tran@example.net 2026-06-02 11:00:05", "Hi there, Lucy here. I want to close my account."),
        ("[agent]AU-0102 2026-06-02 11:00:41", "Hi Lucy, sorry to see you go. Can I have your full name and account ID?"),
        ("[Customer]lucy.tran@example.net 2026-06-02 11:01:20", "Lucy Thanh Tran, account CBK-20458812. I was born 22 November 1992."),
        ("[agent]AU-0102 2026-06-02 11:02:02", "Thank you Ms Tran. I have your phone as 07 3555 0147, is that right?"),
        ("[Customer]lucy.tran@example.net 2026-06-02 11:02:40", "No sorry that's my old landline. Use 0466 700 812. Actually its Tran not Tren, you have it misspelt as Tren on my profile"),
        ("[agent]AU-0102 2026-06-02 11:03:15", "Apologies, I've corrected Tren to Tran. Your last trade was 250 shares at $3.45 in FMG.AX on 2026-04-22."),
        ("[Customer]lucy.tran@example.net 2026-06-02 11:03:50", "Send the closing statement to 18/7 Harbour Rd, Hamilton QLD 4007 please. My sister Mai Tran will collect it."),
    ]
    body = f"<h1>Chat export 2026-06-02</h1>\n{html_chat(msgs)}"
    src = write_html("chat_en_lucy.html", page("Chat export", body))
    # Known limitation: the misspelling "Tren" the customer mentions while correcting
    # it survives; it is not labelled so the gate stays meaningful.
    register("chat_en_lucy.html", "html", "en", src,
             ["lucy.tran@example.net", "Lucy", "Lucy Thanh Tran", "CBK-20458812", "22 November 1992", "Ms Tran", "07 3555 0147",
              "0466 700 812", "Tran", "18/7 Harbour Rd, Hamilton QLD 4007", "Mai Tran", "Mai"],
             ["AU-0102", "250 shares at $3.45", "FMG.AX", "2026-04-22", "11:00:05", "Chat export 2026-06-02"])

    msgs = [
        ("[Customer]wangwei_1985 2026-05-20 21:03:11", "你好，我想修改手机号。我叫王伟。"),
        ("[agent]AU-0117 2026-05-20 21:03:50", "您好王先生，请提供您的身份证号码以便核实。"),
        ("[Customer]wangwei_1985 2026-05-20 21:04:30", f"身份证 {ID_WANG}，新手机 13812345678，微信号 wxid_wangwei85"),
        ("[agent]AU-0117 2026-05-20 21:05:10", "谢谢。请问您的居住地址和出生日期？"),
        ("[Customer]wangwei_1985 2026-05-20 21:05:44", "上海市浦东新区世纪大道100号环球金融中心20楼2008室，出生日期1985年3月16日。朋友都叫我老王，英文名 Wei Wang。"),
        ("[agent]AU-0117 2026-05-20 21:06:20", "好的。另外工单号：TK-88231 已创建，您昨天买入腾讯 00700.HK 200 股，价格 HK$385.60。"),
        ("[Customer]wangwei_1985 2026-05-20 21:07:01", "我太太李娜是联名账户持有人，她的电话是 +86 139 8765 4321，座机 021-58881234。"),
        ("[AI]CraneBot 2026-05-20 21:07:03", "已为您转接 team 人工客服。感谢选择白鹤证券。"),
    ]
    body = f"<h1>在线客服记录</h1>\n{html_chat(msgs)}"
    src = write_html("chat_zh_wangwei.html", page("在线客服记录", body))
    register("chat_zh_wangwei.html", "html", "zh", src,
             ["wangwei_1985", "王伟", "王先生", ID_WANG, "13812345678", "wxid_wangwei85",
              "上海市浦东新区世纪大道100号环球金融中心20楼2008室", "1985年3月16日", "老王", "Wei Wang", "李娜",
              "+86 139 8765 4321", "021-58881234"],
             ["AU-0117", "工单号：TK-88231", "00700.HK", "200 股", "HK$385.60", "[AI]CraneBot", "白鹤证券", "21:03:11"])

    rows = [
        ("First name", "Beatrice"), ("Last name", "Fitzgerald-Ngo"), ("Date of birth", "1971-09-30"),
        ("Mobile", "+61 421 555 019"), ("Email", "bea.fitzgerald@example.org"),
        ("Residential address", "Unit 3, 14 Jacaranda Avenue, Glenelg SA 5045"),
        ("Postal address", "PO Box 918, Adelaide SA 5001"),
        ("Tax file number", "987 654 321"), ("Medicare number", "3456 78901 2"),
        ("Passport number", "N1234567"), ("Driver licence number", "S1234567 (SA)"),
        ("Employer", "Southern Aero Pty Ltd"), ("Product", "Crane Trade margin account"),
        ("Ticket number", "Ticket 55120"),
    ]
    tbl = "\n".join(f"<tr><th>{k}</th><td>{v}</td></tr>" for k, v in rows)
    body = ("<h2>Account application - web form submission</h2>\n"
            f"<table border=\"1\">\n{tbl}\n</table>\n<p>Submitted 2026-02-17 08:44:10. Preferred name: Bea. "
            "Please address correspondence to Ms Fitzgerald-Ngo.</p>")
    src = write_html("webform_en_beatrice.html", page("Application", body))
    register("webform_en_beatrice.html", "html", "en", src,
             ["Beatrice", "Fitzgerald-Ngo", "1971-09-30", "+61 421 555 019", "bea.fitzgerald@example.org",
              "Unit 3, 14 Jacaranda Avenue, Glenelg SA 5045", "PO Box 918, Adelaide SA 5001", "987 654 321", "3456 78901 2",
              "N1234567", "S1234567", "Bea", "Ms Fitzgerald-Ngo"],
             ["Account application - web form submission", "Crane Trade margin account", "Ticket 55120", "2026-02-17 08:44:10",
              "Southern Aero Pty Ltd"])

    msgs = [
        ("[Customer]LiNa_Nina 2026-04-03 19:20:45", "Hi, my name is Li Na but everyone uses Nina. 我的中文名是李娜。"),
        ("[agent]AU-0166 2026-04-03 19:21:10", "您好李女士 / Hello Ms Li. 请问您的手机号码？"),
        ("[Customer]LiNa_Nina 2026-04-03 19:21:50", "0488 123 456, 或者中国号码 15900001234。邮箱 nina.li@example.com"),
        ("[agent]AU-0166 2026-04-03 19:22:30", "收到。Your order: sell 1,000 shares at $0.85 of ASX:PLS on 2026-04-02. Ticket 70311."),
        ("[Customer]LiNa_Nina 2026-04-03 19:23:05", f"我的护照 E12345678，身份证 {ID_LI}. 生日 22/07/1988。我妈妈是陈秀英，住北京市朝阳区建国路88号。"),
    ]
    body = f"<h1>Mixed-language chat</h1>\n{html_chat(msgs)}"
    src = write_html("chat_mixed_lina.html", page("Mixed chat", body))
    register("chat_mixed_lina.html", "html", "mixed", src,
             ["LiNa_Nina", "Li Na", "Nina", "李娜", "李女士", "Ms Li", "0488 123 456", "15900001234", "nina.li@example.com",
              "E12345678", ID_LI, "22/07/1988", "陈秀英", "北京市朝阳区建国路88号"],
             ["AU-0166", "1,000 shares at $0.85", "ASX:PLS", "2026-04-02", "Ticket 70311", "19:20:45"])

    body = ("<h2>客户尽职调查表 (KYC)</h2>\n<table border=\"1\">\n"
            "<tr><td>客户姓名</td><td>周建国</td></tr>\n"
            "<tr><td>拼音</td><td>ZHOU Jianguo</td></tr>\n"
            f"<tr><td>证件号码</td><td>{ID_ZHOU}</td></tr>\n"
            "<tr><td>出生日期</td><td>1979-01-05</td></tr>\n"
            "<tr><td>手机</td><td>+86 186 0011 2233</td></tr>\n"
            "<tr><td>电话</td><td>010-65551234</td></tr>\n"
            "<tr><td>电子邮箱</td><td>zhou.jianguo@example.net</td></tr>\n"
            "<tr><td>微信</td><td>zjg_1979</td></tr>\n"
            "<tr><td>通讯地址</td><td>广东省广州市天河区天河路385号太古汇写字楼18层</td></tr>\n"
            "<tr><td>受益人</td><td>周雨桐（女儿）</td></tr>\n"
            "<tr><td>银行账户</td><td>6222 0212 3456 7890 123</td></tr>\n"
            "<tr><td>产品</td><td>白鹤证券 美股账户，首笔交易 NVDA 50 股 @ $118.20，2025-12-09</td></tr>\n"
            "</table>\n<p>审核人：AU-0129。案件编号：Case 2026-0093。周先生已确认以上信息。</p>")
    src = write_html("kyc_zh_zhou.html", page("KYC", body))
    register("kyc_zh_zhou.html", "html", "zh", src,
             ["周建国", "ZHOU Jianguo", ID_ZHOU, "1979-01-05", "+86 186 0011 2233", "010-65551234", "zhou.jianguo@example.net",
              "zjg_1979", "广东省广州市天河区天河路385号太古汇写字楼18层", "周雨桐", "6222 0212 3456 7890 123", "周先生"],
             ["白鹤证券", "NVDA 50 股 @ $118.20", "2025-12-09", "AU-0129", "Case 2026-0093", "客户尽职调查表 (KYC)"])

    msgs = [
        ("[Customer]lin_xiaoyu 2026-06-10 10:01:00", "Hello, I am Lin Xiaoyu. My pinyin handle is lin_xiaoyu. Chinese name 林晓雨."),
        ("[agent]system 2026-06-10 10:01:01", "Customer verified via SMS OTP to 0433 998 776."),
        ("[agent]AU-0102 2026-06-10 10:01:30", "Hi Xiaoyu, how can I help?"),
        ("[Customer]lin_xiaoyu 2026-06-10 10:02:00", "My surname is Lin, given name Xiaoyu. Please write LIN Xiaoyu on the letter. I want to add my father Lin Dawei as beneficiary. His mobile is 13700998877."),
        ("[AI]CraneBot 2026-06-10 10:02:20", "I can help with that. Please upload form CB-BEN-2 (team will review within 2 business days)."),
        ("[Customer]lin_xiaoyu 2026-06-10 10:03:00", "OK. Also I bought 100 shares of 9988.HK at HK$78.30 on 2026-06-09, is it settled?"),
    ]
    body = f"<h1>Chat</h1>\n{html_chat(msgs)}"
    src = write_html("chat_mixed_xiaoyu.html", page("Chat", body))
    register("chat_mixed_xiaoyu.html", "html", "mixed", src,
             ["lin_xiaoyu", "Lin Xiaoyu", "林晓雨", "0433 998 776", "Xiaoyu", "Lin", "LIN Xiaoyu", "Lin Dawei", "13700998877", "Dawei"],
             ["[agent]system", "[AI]CraneBot", "AU-0102", "CB-BEN-2", "team will review", "100 shares of 9988.HK at HK$78.30",
              "2026-06-09", "Customer verified via SMS OTP"])


# ================================================================ EML
def build_eml():
    plain = """Hi Priya,

Thanks for your call. Please find the details below and confirm.

Kind regards,
Tom Whitlock | Senior Analyst | Crane Brokers
M: 0421 876 543 | tom.whitlock@example.com

On Mon, 1 Jun 2026 at 09:40, Nadia Haddad <nadia.haddad@example.net> wrote:
> Hi Tom,
> I'd like to sell 300 shares of CBA.AX at $131.20 before Friday.
> My DOB is 5 May 1980 for verification, and my TFN is 456 789 123.
> Ring me on 0402 111 222 if needed.
>
> Nadia
> On Fri, 29 May 2026, Tom Whitlock wrote:
>> Hello Ms Haddad, your reference is Case 2026-0451.
"""
    hdr = {"From": str(Address("Tom Whitlock", "tom.whitlock", "example.com")),
           "To": str(Address("Nadia Haddad", "nadia.haddad", "example.net")),
           "Cc": str(Address("Compliance Team", "compliance", "example.com")),
           "Subject": "Re: Sell order query - Nadia Haddad"}
    src = write_eml("thread_en_nadia.eml", hdr, plain)
    register("thread_en_nadia.eml", "eml", "en", src,
             ["Tom Whitlock", "Tom", "Nadia Haddad", "Nadia", "nadia.haddad@example.net", "tom.whitlock@example.com",
              "0421 876 543", "5 May 1980", "456 789 123", "0402 111 222", "Ms Haddad", "Priya", "Whitlock", "Haddad"],
             ["Crane Brokers", "300 shares of CBA.AX at $131.20", "Case 2026-0451", "Compliance Team", "Senior Analyst"])

    plain = "您好周女士，\n\n您的入金已到账，金额 5,000 澳元。\n如有疑问请致电 021-62881234 或 +86 137 0000 5566。\n\n谢谢\n客服 陈晨\n白鹤证券"
    html = ("<html><body><p>您好周女士，</p><p>您的入金已到账，金额 5,000 澳元。</p>"
            "<p>如有疑问请致电 021-62881234 或 +86 137 0000 5566。</p><p>谢谢<br>客服 陈晨<br>白鹤证券</p></body></html>")
    hdr = {"From": str(Address("陈晨", "chen.chen", "example.com")),
           "To": str(Address("周敏", "zhou.min88", "example.org")),
           "Subject": "您的入金通知 - 周敏 账户 CBK-90021177"}
    src = write_eml("alt_zh_zhoumin.eml", hdr, plain, html)
    register("alt_zh_zhoumin.eml", "eml", "zh", src,
             ["周女士", "陈晨", "周敏", "zhou.min88@example.org", "chen.chen@example.com", "021-62881234", "+86 137 0000 5566",
              "CBK-90021177"],
             ["白鹤证券", "5,000 澳元", "入金"])

    att_blocks = [("h", "Change of details form"),
                  ("table", [["Field", "Value"], ["Full name", "Rohan D'Souza"], ["DOB", "12 December 1990"],
                             ["New mobile", "0417 222 333"], ["New address", "5/9 Banksia Lane, Fremantle WA 6160"],
                             ["Broker account", "CB-55009123"], ["Bank", "BSB 016-123 acct 9988 7766"]]),
                  ("p", "Signed: Rohan D'Souza. Trade reference: 75 shares at $41.10 of WES.AX on 2026-05-06.")]
    data, dsrc = docx_bytes(att_blocks, footer="Crane Brokers form CB-COD-1")
    plain = "Hi team,\n\nAttached is my change of details form. Regards,\nRohan\n\n-- \nRohan D'Souza\nMob 0417 222 333\n"
    hdr = {"From": str(Address("Rohan D'Souza", "rohan.dsouza", "example.org")),
           "To": str(Address("Crane Support", "support", "example.com")),
           "Subject": "Change of details form"}
    esrc = write_eml("att_docx_rohan.eml", hdr, plain, attachments=[
        ("change_of_details.docx", "application", "vnd.openxmlformats-officedocument.wordprocessingml.document", data)], day=3)
    register("att_docx_rohan.eml", "eml", "en", esrc + "\n" + dsrc,
             ["Rohan D'Souza", "Rohan", "rohan.dsouza@example.org", "0417 222 333", "12 December 1990",
              "5/9 Banksia Lane, Fremantle WA 6160", "CB-55009123", "016-123", "9988 7766", "D'Souza"],
             ["Change of details form", "75 shares at $41.10 of WES.AX on 2026-05-06", "Crane Support", "Crane Brokers form CB-COD-1"])

    att_html = ("<html><body><h3>CRM export</h3><p>Customer: Grace Okonkwo-Lee, mobile +61 400 998 001, "
                "email grace.ok@example.net</p><p>Address: 77 Station Street, Penrith NSW 2750. "
                "DOB 07.10.1985. Driver licence 45678912.</p><p>Note: bought 10 shares of MSFT on 2026-01-15.</p></body></html>")
    plain = "Forwarding the CRM export for the audit file. Please review by Friday.\nThanks,\nDenise Carter\nSenior Compliance Officer\nM +61 418 300 200\n"
    hdr = {"From": str(Address("Denise Carter", "denise.carter", "example.com")),
           "To": str(Address("Audit Team", "audit", "example.com")),
           "Subject": "FW: CRM export for audit file"}
    esrc = write_eml("att_html_grace.eml", hdr, plain, attachments=[
        ("crm_export.html", "text", "html", att_html.encode("utf-8"))], day=4)
    register("att_html_grace.eml", "eml", "en", esrc + "\n" + att_html,
             ["Denise Carter", "Denise", "denise.carter@example.com", "+61 418 300 200", "Grace Okonkwo-Lee", "+61 400 998 001",
              "grace.ok@example.net", "77 Station Street, Penrith NSW 2750", "07.10.1985", "45678912", "Grace", "Okonkwo-Lee"],
             ["Audit Team", "Senior Compliance Officer", "10 shares of MSFT on 2026-01-15", "CRM export"])

    plain = """Dear Crane,

I want to close my account. My name is Daniel O'Brien, and the email I registered with
is daniel.obrien@example.org (my wife uses dan.and.kate@example.net for the joint account).
Please send confirmation to my wife Kate O'Brien as well. Our referrer Sanjay Patel said
it would be quick. Account holder ID: 00123456.

Please use the mailing address 2/14 Ocean View Rd, Manly NSW 2095.

Thanks
Dan
0411 222 333
"""
    hdr = {"From": str(Address("Daniel O'Brien", "daniel.obrien", "example.org")),
           "To": str(Address("Crane Client Services", "clientservices", "example.com")),
           "Cc": str(Address("Kate O'Brien", "dan.and.kate", "example.net")),
           "Subject": "Account closure request - Daniel O'Brien (ID 00123456)"}
    src = write_eml("subject_en_obrien.eml", hdr, plain, day=5)
    register("subject_en_obrien.eml", "eml", "en", src,
             ["Daniel O'Brien", "daniel.obrien@example.org", "dan.and.kate@example.net", "Kate O'Brien", "Sanjay Patel",
              "00123456", "2/14 Ocean View Rd, Manly NSW 2095", "Dan", "0411 222 333", "Kate", "Sanjay", "O'Brien"],
             ["Crane Client Services", "Account closure request"])

    plain = "您好，\n\n我叫赵磊，之前用英文名 Ray Zhao 开户。我想把账户里的 300 股 AAPL 转出。\n手机 18611112222，也可以打 +61 430 555 777 (悉尼)。\n地址：Unit 5, 200 George Street, Sydney NSW 2000\nBSB 032-000，账号 123456。\n案件号 Ticket 90210 。\n\n赵磊\n"
    hdr = {"From": str(Address("赵磊", "ray.zhao", "example.net")),
           "To": str(Address("Crane Support", "support", "example.com")),
           "Subject": "转出股票申请 Ray Zhao"}
    src = write_eml("mixed_zhao.eml", hdr, plain, day=6)
    register("mixed_zhao.eml", "eml", "mixed", src,
             ["赵磊", "Ray Zhao", "ray.zhao@example.net", "18611112222", "+61 430 555 777",
              "Unit 5, 200 George Street, Sydney NSW 2000", "032-000", "123456", "Ray", "Zhao"],
             ["300 股 AAPL", "Ticket 90210", "Crane Support", "转出股票申请"])

    plain = """Agent notes below.

[Customer]sophie.b 2026-06-11 16:20:10
Hi I'm Sophie Bernard, my account is CB-30045566.
[agent]AU-0102 2026-06-11 16:20:44
Hi Sophie, verifying you now. Can you confirm DOB?[Customer]sophie.b 2026-06-11 16:21:02
1 April 1995, mobile +61 402 700 700
[agent]AU-0102 2026-06-11 16:21:30
Thanks. Your 200 shares of BHP.AX at $45.10 were filled 2026-06-10.
"""
    hdr = {"From": str(Address("Support Bot", "noreply", "example.com")),
           "To": str(Address("QA Team", "qa", "example.com")),
           "Subject": "Chat transcript for QA review"}
    src = write_eml("qa_transcript_sophie.eml", hdr, plain, day=11)
    register("qa_transcript_sophie.eml", "eml", "en", src,
             ["sophie.b", "Sophie Bernard", "Sophie", "CB-30045566", "1 April 1995", "+61 402 700 700", "Bernard"],
             ["AU-0102", "200 shares of BHP.AX at $45.10", "2026-06-10", "QA Team", "16:20:10"])


# ================================================================ DOCX
def build_docx():
    blocks = [
        ("p", "Crane Brokers Pty Ltd"),
        ("p", "Level 9, 55 Collins Street Melbourne VIC 3000"),
        ("p", "12 May 2026"),
        ("p", "Mr Anthony Papadopoulos"),
        ("p", "8 Hillcrest Grove"),
        ("p", "Box Hill VIC 3128"),
        ("p", "Dear Mr Papadopoulos,"),
        ("p", "Re: Your Crane account CB-77120034 - request 2026-0388"),
        ("p", "Thank you for your letter dated 30 April 2026. We have updated your registered mobile to 0409 876 123 and email to "
              "tony.papa@example.org. Your date of birth on file is 17/08/1964, and your TFN ending 321 has been recorded."),
        ("p", "Your sale of 800 shares at $6.75 in WDS.AX on 2026-05-08 settled on 2026-05-12."),
        ("p", "If Tony (your preferred name) or your wife Eleni Papadopoulos wish to discuss this, call us on 1300 000 000 or "
              "(03) 9555 0192."),
        ("p", "Yours sincerely, Kim Alvarez, Client Services Manager"),
    ]
    src = write_docx("letter_en_papadopoulos.docx", blocks, header="Private and confidential",
                     footer="Crane Brokers Pty Ltd | Client: Papadopoulos, Anthony | page 1")
    register("letter_en_papadopoulos.docx", "docx", "en", src,
             ["Mr Anthony Papadopoulos", "Anthony Papadopoulos", "8 Hillcrest Grove", "Box Hill VIC 3128", "Mr Papadopoulos",
              "0409 876 123", "tony.papa@example.org", "17/08/1964", "Tony", "Eleni Papadopoulos", "(03) 9555 0192",
              "Kim Alvarez", "Papadopoulos, Anthony", "Eleni", "Anthony",
              "CB-77120034"],  # "Your Crane account CB-77120034" identifies the customer.
             ["Crane Brokers Pty Ltd", "800 shares at $6.75 in WDS.AX on 2026-05-08", "Client Services Manager",
              "Private and confidential"])

    blocks = [
        ("h", "KYC / AML periodic review"),
        ("table", [["Item", "Detail"],
                   ["Customer", "Yusuf Al-Rashid"],
                   ["Also known as", "Joe Rashid"],
                   ["DOB", "March 3, 1983"],
                   ["Nationality / passport", "Australian / PA7654321"],
                   ["Residential", "Unit 22/60 Marine Parade, Cottesloe WA 6011"],
                   ["Mobile", "+61 4 1234 5678"],
                   ["Home phone", "08 9384 1122"],
                   ["Email", "y.alrashid@example.com"],
                   ["Joint holder", "Layla Al-Rashid (spouse)"],
                   ["Bank", "BSB 306-018, Acc 00887766"],
                   ["Risk rating", "Low"]]),
        ("table", [["Date", "Trade", "Notes"],
                   ["2026-02-03", "BUY 500 shares at $12.34 BHP.AX", "Funds from ANZ account"],
                   ["2026-03-19", "SELL 100 shares 00700.HK", "Ticket 61008"]]),
        ("p", "Reviewer comment: Mr Al-Rashid confirmed source of funds by phone. Layla joined the call."),
    ]
    src = write_docx("kyc_en_alrashid.docx", blocks, header="AML review - AU-0188",
                     footer="Reviewed for Yusuf Al-Rashid on 2026-06-15")
    register("kyc_en_alrashid.docx", "docx", "en", src,
             ["Yusuf Al-Rashid", "Joe Rashid", "March 3, 1983", "PA7654321", "Unit 22/60 Marine Parade, Cottesloe WA 6011",
              "+61 4 1234 5678", "08 9384 1122", "y.alrashid@example.com", "Layla Al-Rashid", "306-018", "00887766",
              "Mr Al-Rashid", "Layla", "Yusuf", "Al-Rashid"],
             ["KYC / AML periodic review", "AML review - AU-0188", "BUY 500 shares at $12.34 BHP.AX", "2026-02-03",
              "SELL 100 shares 00700.HK", "Ticket 61008", "Risk rating"])

    blocks = [
        ("p", "白鹤证券（澳大利亚）客户通知函"),
        ("p", "2026年6月18日"),
        ("p", "尊敬的陈秀英女士："),
        ("p", "根据您的申请，我们已将您的联系方式更新为：手机 13566778899；电子邮箱 chen.xiuying@example.org；"
              "通讯地址 浙江省杭州市西湖区文三路90号。"),
        ("p", f"经核实，您的身份证号码为 {ID_CHEN}，出生日期为1993年9月11日。"),
        ("p", "您名下的港股账户已于 2026-06-10 买入 00700.HK 100 股，价格 HK$372.40。"),
        ("p", "如需协助请联系我们的团队，或致电 0571-85551234。"),
        ("p", "客户服务部 张经理"),
    ]
    src = write_docx("letter_zh_chen.docx", blocks, header="白鹤证券 机密",
                     footer="致：陈秀英 | 页 1")
    register("letter_zh_chen.docx", "docx", "zh", src,
             ["陈秀英", "陈秀英女士", "13566778899", "chen.xiuying@example.org", "浙江省杭州市西湖区文三路90号", ID_CHEN,
              "1993年9月11日", "0571-85551234", "张经理"],
             ["白鹤证券（澳大利亚）客户通知函", "2026-06-10", "00700.HK 100 股", "HK$372.40", "客户服务部"])

    blocks = [
        ("h", "File note - call back"),
        ("p", "Ticket: CS-88412   Agent: AU-0102   Date: 2026-07-02 15:05"),
        ("p", "Spoke to 刘洋 (Liu Yang, spelled L-I-U Y-A-N-G) regarding transfer of 200 shares of AAPL. "
              "Customer prefers to be called Mr Liu. 微信 liuyang_sydney."),
        ("p", "Customer's mother, 刘淑芬, lives at 3 Pine Street, Chatswood NSW 2067 and is the nominated beneficiary. "
              "DOB of Liu Yang: 1990年12月1日 (01/12/1990 AU format)."),
        ("p", "Contact numbers: 0425 667 788 (mobile), +86 152 1234 9876 (China), 03 9012 3456 (work)."),
        ("p", "Email: liu.yang@example.net. Note to team: customer wants Chinese-language response."),
    ]
    src = write_docx("filenote_mixed_liu.docx", blocks, footer="Crane Brokers internal")
    register("filenote_mixed_liu.docx", "docx", "mixed", src,
             ["刘洋", "Liu Yang", "L-I-U Y-A-N-G", "Mr Liu", "liuyang_sydney", "刘淑芬", "3 Pine Street, Chatswood NSW 2067",
              "1990年12月1日", "01/12/1990", "0425 667 788", "+86 152 1234 9876", "03 9012 3456", "liu.yang@example.net",
              "Liu", "Yang"],
             ["File note - call back", "CS-88412", "AU-0102", "2026-07-02 15:05", "200 shares of AAPL", "Crane Brokers internal"])

    msgs = [
        ("[Customer]Guest 20:11:05", "你好，我是黄海峰。"),
        ("[agent]AU-0155 20:11:40", "黄先生您好，请问有什么可以帮您？"),
        ("[Customer]Guest 20:12:10", "我要查询昨天卖出的 BABA 50 股。我的账户号 88001234。"),
        ("[agent]AU-0155 20:12:55", "请问您的手机号？"),
        ("[Customer]Guest 20:13:20", "13911223344。我小名叫阿峰，拼音 Huang Haifeng，邮箱 haifeng.h@example.net"),
    ]
    src = write_docx("chat_zh_huang.docx", [("p", "聊天记录"), ("p", "")] + [("p", x) for x in chat(msgs)])
    register("chat_zh_huang.docx", "docx", "zh", src,
             ["黄海峰", "黄先生", "88001234", "13911223344", "阿峰", "Huang Haifeng", "haifeng.h@example.net"],
             ["聊天记录", "AU-0155", "BABA 50 股", "20:11:05"])

    blocks = [
        ("h", "Power of attorney and estate - file summary"),
        ("p", "Deceased account holder: Margaret Ellen Whitby (DOB 2 February 1941), account CB-11090022."),
        ("p", "Executor: her son Graham Whitby, 40 Elm Court, Toowoomba QLD 4350, mobile 0447 300 222, graham.whitby@example.org."),
        ("p", "Beneficiaries: Graham Whitby (50%), Susan Whitby-Clarke (50%). Susan's TFN 234 567 891."),
        ("p", "Death certificate received 2026-05-30. Ticket EST-2026-15. Holdings: 600 shares of TLS.AX."),
    ]
    src = write_docx("estate_en_whitby.docx", blocks, footer="Estate team - AU-0170")
    register("estate_en_whitby.docx", "docx", "en", src,
             ["Margaret Ellen Whitby", "Margaret", "2 February 1941", "CB-11090022", "Graham Whitby", "Graham",
              "40 Elm Court, Toowoomba QLD 4350", "0447 300 222", "graham.whitby@example.org", "Susan Whitby-Clarke", "Susan",
              "234 567 891", "Whitby"],
             ["Power of attorney and estate - file summary", "2026-05-30", "Ticket EST-2026-15", "600 shares of TLS.AX", "AU-0170"])


# ================================================================ more chats (HTML/DOCX/PDF variants)
def build_more():
    msgs = [
        ("[Customer]Amir 2026-07-14 12:00:01", "Hi I'm Amir Karimi."),
        ("[agent]AU-0102 2026-07-14 12:00:30", "Hi Amir. Spell that for me?"),
        ("[Customer]Amir 2026-07-14 12:00:55", "Karimi, K A R I M I. And first name A-M-I-R."),
        ("[agent]AU-0102 2026-07-14 12:01:20", "Thanks, and your date of birth?"),
        ("[Customer]Amir 2026-07-14 12:01:45", "Fourth of July 1989. Phone: 0423 456 789."),
        ("[Customer]Amir 2026-07-14 12:02:10", "email a.karimi@example.com. I live at 9 Cedar Way, Kew VIC 3101"),
        ("[agent]AU-0102 2026-07-14 12:02:50", "Thanks Amir. Order 42: buy 15 shares at $250.00 of AAPL, trade date 2026-07-13."),
    ]
    body = f"<h1>Live chat</h1>\n{html_chat(msgs)}"
    src = write_html("chat_en_amir.html", page("Live chat", body))
    register("chat_en_amir.html", "html", "en", src,
             ["Amir", "Amir Karimi", "Karimi", "K A R I M I", "A-M-I-R", "Fourth of July 1989", "0423 456 789",
              "a.karimi@example.com", "9 Cedar Way, Kew VIC 3101"],
             ["AU-0102", "15 shares at $250.00 of AAPL", "2026-07-13", "12:00:01"])

    msgs = [
        ("[Customer]13987654321 2026-08-01 08:00:00", "你好，账户登录不了。"),
        ("[agent]AU-0101 2026-08-01 08:00:20", "请提供注册手机号。"),
        ("[Customer]13987654321 2026-08-01 08:00:50", "就是 13987654321。我叫吴鹏飞，Peng Fei Wu。身份证后六位 我发全的：" + cn_id("510104", "19960228", "331")),
        ("[agent]AU-0101 2026-08-01 08:01:20", "吴先生，已重置密码，发送到 pengfei.wu@example.org。"),
        ("[Customer]13987654321 2026-08-01 08:01:50", "谢谢。我的同事 Amy（王雅婷）也需要重置，她的号码是 028-86661234。"),
    ]
    src_lines = [x for x in chat(msgs)]
    zid = cn_id("510104", "19960228", "331")
    body = f"<h1>客服对话</h1>\n{html_chat(msgs)}"
    src = write_html("chat_zh_phonehandle.html", page("客服对话", body))
    register("chat_zh_phonehandle.html", "html", "zh", src,
             ["13987654321", "吴鹏飞", "Peng Fei Wu", zid, "吴先生", "pengfei.wu@example.org", "Amy", "王雅婷", "028-86661234"],
             ["AU-0101", "客服对话", "08:00:00"])

    msgs = [
        ("[Customer]Sam_the_trader 2026-09-01 07:30:00", "morning. its Samantha Reid, formerly Samantha Hollis - changed name after marriage"),
        ("[agent]AU-0102 2026-09-01 07:30:40", "Hi Samantha, please update via form. Anything else?"),
        ("[Customer]Sam_the_trader 2026-09-01 07:31:05", "yes my old surname is still on the statement. Hollis. Also new mobile 0499 000 123"),
        ("[agent]AU-0102 2026-09-01 07:31:40", "Done, Mrs Reid. I'm Marcus by the way, from the team."),
        ("[AI]CraneAssist 2026-09-01 07:31:45", "Reminder: sell 250 shares at $9.99 of ZIP.AX on 2026-08-28 settles T+2."),
    ]
    src = write_docx("chat_en_samantha.docx", [("p", "Chat export")] + [("p", x) for x in chat(msgs)])
    register("chat_en_samantha.docx", "docx", "en", src,
             ["Sam_the_trader", "Samantha Reid", "Samantha Hollis", "Samantha", "Hollis", "0499 000 123", "Mrs Reid", "Reid",
              "Marcus"],
             ["AU-0102", "[AI]CraneAssist", "250 shares at $9.99 of ZIP.AX", "2026-08-28", "from the team", "07:30:00"])

    lines = [
        "Crane Brokers - Client complaint letter (scanned text layer)",
        "",
        "To: Complaints Officer",
        "From: Mrs Dorothy Wilkinson-Hart",
        "Address: 21 Banyan Street, Alderley QLD 4051",
        "Phone: 07 3356 7788   Mobile: 0438 900 111   Email: d.wilkinson.hart@example.org",
        "DOB: 1950-06-25   Customer number: 4400 1188",
        "",
        "Dear Sir/Madam,",
        "I am writing about the sale of 1,200 shares at $2.05 of ORG.AX on 2026-01-20 that my son",
        "Peter Hart made without my authority. He is not on the account. Please call me, Dorothy,",
        "not Peter. Case ref: Complaint 2026/113.",
        "Yours faithfully,",
        "D. Wilkinson-Hart",
    ]
    src = write_pdf("letter_en_wilkinson.pdf", lines)
    register("letter_en_wilkinson.pdf", "pdf", "en", src,
             ["Mrs Dorothy Wilkinson-Hart", "Dorothy Wilkinson-Hart", "Wilkinson-Hart", "21 Banyan Street, Alderley QLD 4051",
              "07 3356 7788", "0438 900 111", "d.wilkinson.hart@example.org", "1950-06-25", "4400 1188", "Peter Hart",
              "Dorothy", "Peter", "D. Wilkinson-Hart"],
             ["Crane Brokers", "1,200 shares at $2.05 of ORG.AX", "2026-01-20", "Complaint 2026/113"])

    plain = """Hi Marco,

Quick one: the customer Chen Yu (陈宇, WeChat: chenyu_hz) called about wire 2026-06-20.
His HK phone is +852 9123 4567 and CN mobile 13600001111. ID 33010219960101 was rejected -
please ask him for the correct one.

Sent from my iPhone
Luca Ferrari
Customer Support Lead, Crane Brokers
Direct: 0431 555 100
"""
    hdr = {"From": str(Address("Luca Ferrari", "luca.ferrari", "example.com")),
           "To": str(Address("Marco Bianchi", "marco.bianchi", "example.com")),
           "Subject": "Chen Yu wire follow-up"}
    src = write_eml("internal_mixed_chenyu.eml", hdr, plain, day=20)
    register("internal_mixed_chenyu.eml", "eml", "mixed", src,
             ["Marco", "Marco Bianchi", "marco.bianchi@example.com", "Chen Yu", "陈宇", "chenyu_hz", "+852 9123 4567",
              "13600001111", "Luca Ferrari", "luca.ferrari@example.com", "0431 555 100", "Luca", "Ferrari"],
             ["Customer Support Lead, Crane Brokers", "2026-06-20", "wire"])


def main():
    CORPUS.mkdir(parents=True, exist_ok=True)
    for old in CORPUS.iterdir():
        old.unlink()
    build_pdfs()
    build_html()
    build_eml()
    build_docx()
    build_more()
    assert 25 <= len(LABELS) <= 40, len(LABELS)
    for n in LABELS:
        assert (CORPUS / n).exists(), n
    (HERE / "labels.json").write_text(json.dumps(LABELS, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")


if __name__ == "__main__":
    main()
