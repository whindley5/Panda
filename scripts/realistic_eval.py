"""End-to-end PII-survival harness over the synthetic corpus in tests/fixtures/realistic/.

Usage: python scripts/realistic_eval.py [--keep] [--summary] [--fixture realistic|holdout|blind]
A fixture directory holds generate.py, which writes corpus/ and labels.json.
Exits non-zero if any labelled PII string survives sanitisation. All data is invented,
so surviving values are printed verbatim.
"""
from __future__ import annotations

import argparse
import email
import json
import re
import shutil
import subprocess
import sys
import tempfile
from email import policy
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "tests" / "fixtures" / "realistic"
CORPUS = FIX / "corpus"
# The corpora are written for a fictional firm, Crane Brokers (白鹤证券).
KEEP_ARGS = [
    "--keep-file", str(ROOT / "tests" / "fixtures" / "keep.txt"), "--account-prefix", "CB",
]


def docx_text(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    out: list[str] = []

    def table(t):
        for row in t.rows:
            for c in row.cells:
                out.extend(p.text for p in c.paragraphs)
                for nt in c.tables:
                    table(nt)

    out.extend(p.text for p in d.paragraphs)
    for t in d.tables:
        table(t)
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                     s.even_page_header, s.even_page_footer):
            out.extend(p.text for p in part.paragraphs)
            for t in part.tables:
                table(t)
    return "\n".join(out)


def eml_texts(path: Path) -> list[str]:
    raw = path.read_bytes()
    texts = [raw.decode("utf-8", "replace")]
    try:
        m = email.message_from_bytes(raw, policy=policy.default)
        parts = [f"{k}: {v}" for k, v in m.items()]
        for part in m.walk():
            if part.get_content_maintype() == "text":
                parts.append(part.get_content())
        texts.append("\n".join(parts))
    except Exception as exc:  # unparsable output is itself a finding, raw text still checked
        texts.append(f"<<unparsable eml: {exc}>>")
    return texts


def read_output(path: Path) -> list[str]:
    if path.name.endswith(".docx"):
        return [docx_text(path)]
    if path.name.endswith(".eml"):
        # Include text inside retained attachments (a sanitised .docx is base64).
        from panda.handlers.eml import extract_eml_units

        return eml_texts(path) + [unit.text for unit in extract_eml_units(path)]
    return [path.read_text(encoding="utf-8", errors="replace")]


def find_outputs(out: Path) -> dict[int, Path]:
    found: dict[int, Path] = {}
    files = sorted(out.rglob("file-*"), key=lambda p: (len(p.relative_to(out).parts), str(p)))
    for p in files:
        if p.is_file():
            try:
                n = int(p.name.split("-")[1].split(".")[0])
            except ValueError:
                continue
            found.setdefault(n, p)
    return found


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--keep", action="store_true", help="keep the temp output/report dirs")
    ap.add_argument("--fixture", default="realistic", help="fixture directory under tests/fixtures")
    ap.add_argument("--summary", action="store_true", help="omit the per-file table")
    args = ap.parse_args()
    global FIX, CORPUS
    FIX = ROOT / "tests" / "fixtures" / args.fixture
    CORPUS = FIX / "corpus"

    subprocess.run([sys.executable, str(FIX / "generate.py")], check=True)
    labels = json.loads((FIX / "labels.json").read_text(encoding="utf-8"))
    sources = sorted(labels, key=lambda n: str(CORPUS / n))

    tmp = Path(tempfile.mkdtemp(prefix="realistic_eval_"))
    out, rep = tmp / "out", tmp / "rep"
    proc = subprocess.run(
        [sys.executable, "-m", "panda", "run", "--opaque-names", *KEEP_ARGS,
         str(CORPUS), "--output", str(out), "--report", str(rep)],
        cwd=ROOT, capture_output=True, text=True)
    print(f"CLI exit code: {proc.returncode}   (tmp: {tmp})")
    if proc.returncode != 0:
        print("stderr tail:", "\n".join(proc.stderr.strip().splitlines()[-8:]))
    found = find_outputs(out)
    outputs = {name: found[i] for i, name in enumerate(sources, 1) if i in found}

    tot_pii = tot_surv = tot_keep = tot_lost = missing = 0
    rows, survivors, losses = [], [], []
    for name in sources:
        lab = labels[name]
        o = outputs.get(name)
        if o is None:
            missing += 1
            rows.append((name, "NO OUTPUT", "-", "-"))
            continue
        want = ".pdf.txt" if name.endswith(".pdf") else "." + name.rsplit(".", 1)[1]
        if not o.name.endswith(want):
            rows.append((name, f"MAPPING MISMATCH {o.name}", "-", "-"))
            missing += 1
            continue
        texts = read_output(o)
        surv = []
        for s in lab["pii"]:
            # Word-bounded for word-like labels, so "Sam" does not match the dummy
            # "16 Sample Street" and "Mark" does not match inside "Marketing".
            pattern = re.compile(
                (r"(?<![A-Za-z])" if s[:1].isalpha() else "") + re.escape(s)
                + (r"(?![a-z])" if s[-1:].isalpha() else "")
            )
            c = max(len(pattern.findall(t)) for t in texts)
            if c:
                surv.append((s, c))
        lost = [s for s in lab["must_keep"] if not any(s in t for t in texts)]
        tot_pii += len(lab["pii"])
        tot_surv += len(surv)
        tot_keep += len(lab["must_keep"])
        tot_lost += len(lost)
        survivors.append((name, surv))
        losses.append((name, lost))
        rows.append((name, o.name, f"{len(surv)}/{len(lab['pii'])}", f"{len(lost)}/{len(lab['must_keep'])}"))

    if not args.summary:
        w = max(len(r[0]) for r in rows)
        print(f"\n{'file':<{w}}  {'output':<22} {'PII left':>9} {'keep lost':>10}")
        for r in rows:
            print(f"{r[0]:<{w}}  {r[1]:<22} {r[2]:>9} {r[3]:>10}")
    print(f"\nTOTAL: {tot_surv}/{tot_pii} labelled PII strings survived; "
          f"{tot_lost}/{tot_keep} must_keep strings lost; {missing} files without usable output")

    print("\nSURVIVING PII (string x occurrences):")
    for name, surv in survivors:
        for s, c in surv:
            print(f"  {name}: {s!r} x{c}")
    print("\nLOST must_keep:")
    for name, lost in losses:
        for s in lost:
            print(f"  {name}: {s!r}")
    if not args.keep:
        shutil.rmtree(tmp, ignore_errors=True)
    return 1 if (tot_surv or missing) else 0


if __name__ == "__main__":
    sys.exit(main())
